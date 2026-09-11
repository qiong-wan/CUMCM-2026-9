"""Appendix-3 coupled radial drying, SDIRK2 finite volumes and event refinement.

Only problem-3 output is written. Run --audit, --case NAME --n N --factor F,
then deliver_problem3.py after the convergence cases have completed.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import time
from pathlib import Path

import numpy as np
import openpyxl
from scipy.linalg import solve_banded

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "problem-3"
sys.path.insert(0, str(ROOT / "src"))
from common.fvm import build_geometry  # noqa: E402

R, L, H, HM = 0.02, 0.25, 25.0, 8e-7
GAMMA = 1.0 - 1.0 / np.sqrt(2.0)
ITER_TOL, RES_TOL = 2e-11, 2e-11
CHECK_TIMES = np.array([0.01, 0.1, 1., 10., 60., 600., 1800., 3600.,
                        5400., 7200., 9000., 10800., 14400.])


def dump_json(path, obj):
    """Write readable, finite JSON to an authorized derived file."""
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2,
                               allow_nan=False), encoding="utf-8")


def hashes():
    """Fingerprint protected task inputs, code, reports and earlier outputs."""
    paths = [ROOT / "problem.pdf", ROOT / "AGENTS.md"]
    for folder in ["data", "src/common", "src/problem-1", "src/problem-2",
                   "report", "output/problem-1", "output/problem-2"]:
        paths.extend(p for p in (ROOT / folder).rglob("*") if p.is_file()
                     and not p.name.startswith("~$") and "__pycache__" not in p.parts
                     and not ("output" in p.parts and p.suffix.lower() == ".png"))
    return {str(p.relative_to(ROOT)): hashlib.sha256(p.read_bytes()).hexdigest()
            for p in sorted(paths)}


def audit():
    """Inspect every environmental record without correcting or omitting rows."""
    OUT.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(ROOT / "data" / "附件1.xlsx", read_only=True,
                               data_only=True)
    ws = wb["Sheet1"]
    rows = list(ws.values)
    wb.close()
    if rows[0] != ("时间", "温度", "水分浓度"):
        raise ValueError(f"Unexpected fields: {rows[0]}")
    try:
        values = np.asarray(rows[1:], dtype=float)
    except (TypeError, ValueError) as exc:
        raise ValueError("Anomaly: nonnumeric observation; user decision needed") from exc
    checks = {"shape_241_by_3": values.shape == (241, 3),
              "finite": bool(np.isfinite(values).all()),
              "missing": int(sum(v is None for row in rows[1:] for v in row)),
              "duplicate_times": int(len(values) - len(np.unique(values[:, 0]))),
              "duplicate_records": int(len(values) - len(np.unique(values, axis=0))),
              "strict_order": bool(np.all(np.diff(values[:, 0]) > 0)),
              "all_intervals_60_s": bool(np.all(np.diff(values[:, 0]) == 60)),
              "coverage": values[[0, -1], 0].tolist(),
              "last_observation": values[-1].tolist(),
              "minimum": values.min(axis=0).tolist(),
              "maximum": values.max(axis=0).tolist()}
    valid = (checks["shape_241_by_3"] and checks["finite"] and not checks["missing"]
             and not checks["duplicate_times"] and not checks["duplicate_records"]
             and checks["strict_order"] and checks["all_intervals_60_s"]
             and checks["coverage"] == [0, 14400]
             and np.allclose(values[-1, 1:], [50.165, .04986], rtol=0, atol=1e-12)
             and np.all(values[:, 1] > -273.15) and np.all(values[:, 2] > 0))
    wb = openpyxl.load_workbook(ROOT / "data" / "附件3" / "result3.xlsx",
                               read_only=True, data_only=True)
    template = {"sheets": wb.sheetnames, "rows": list(wb["Sheet1"].values)}
    wb.close()
    valid = valid and template["sheets"] == ["Sheet1"]
    valid = valid and template["rows"][0][0] == "时间\\到药材中心的距离"
    valid = valid and [row[0] for row in template["rows"][1:4]] == [60, 120, 180]
    tail = values[values[:, 0] >= 12600]
    report = {"status": "PASS" if valid else "FAIL", "environment": checks,
              "units": ["s", "degC", "kg/kg"], "template": template,
              "tail_window_s": [12600, 14400], "tail_count": len(tail),
              "tail_arithmetic_means": tail[:, 1:].mean(axis=0).tolist(),
              "treatment": "No corrections, filtering, sorting, smoothing or clipping."}
    dump_json(OUT / "input_audit.json", report)
    if not valid:
        raise ValueError("Input audit FAIL; do not proceed without user decision")
    if not (OUT / "protected_hashes_before.json").exists():
        dump_json(OUT / "protected_hashes_before.json", hashes())
    np.savetxt(OUT / "observations_readonly_copy.csv", values, delimiter=",",
               header="time_s,T_degC,C_kg_kg", comments="", fmt="%.17g")
    return values, report


class Environment:
    """Exact piecewise-linear observations plus an explicit tail hypothesis."""

    def __init__(self, values, mode="last"):
        """Select final values or the inclusive last-30-minute arithmetic mean."""
        self.values = values
        self.mode = mode
        self.tail = (values[-1, 1:] if mode == "last" else
                     values[values[:, 0] >= 12600, 1:].mean(axis=0))

    def __call__(self, t):
        """Query the actual boundary at each internal implicit stage time."""
        if t > 14400:
            return self.tail
        return np.array([np.interp(t, self.values[:, 0], self.values[:, j])
                         for j in [1, 2]])


def properties(T, C):
    """Return positive local volumetric heat capacity, conductivity and D."""
    if not np.isfinite(T).all() or not np.isfinite(C).all():
        raise FloatingPointError("Nonfinite state")
    if np.min(C) <= 0 or np.min(T + 273.15) <= 0:
        raise FloatingPointError("Nonpositive C or absolute temperature")
    a = (650 + 128 * C) * (1450 + 2736 * C / (1 + C))
    k = .21 + .38 * C / (1 + C)
    D = 2.4e-3 * np.exp(-.45 / C - 3850 / (T + 273.15))
    if np.any(D <= 0) or not np.isfinite(D).all():
        raise FloatingPointError("Nonpositive or nonfinite diffusivity")
    return a, k, D


class Solver:
    """Shared-flux finite volumes with converged nonlinear implicit stages."""

    def __init__(self, n, env):
        """Reuse only the audited common geometry, with N intervals/N+1 nodes."""
        self.n, self.env = n, env
        self.dr, self.r, self.v, self.ge, _, self.ar = build_geometry(n, R)
        self.max_iter = 0
        self.max_change = 0.
        self.max_residual = 0.

    def conductance(self, coefficient):
        """Evaluate stable positive harmonic interface properties, without clips."""
        face = 2 / (1 / coefficient[:-1] + 1 / coefficient[1:])
        if np.any(face <= 0) or not np.isfinite(face).all():
            raise FloatingPointError("Nonpositive or nonfinite harmonic face property")
        return self.ge * face / self.dr

    def net(self, u, G, h, ambient):
        """Assemble independent inward net flux, cancelling shared inner faces."""
        flux = G * np.diff(u)
        out = np.zeros_like(u)
        out[:-1] += flux
        out[1:] -= flux
        out[-1] += self.ar * h * (ambient - u[-1])
        return out

    def linear(self, base, capacity, G, h, ambient, dt):
        """Solve for an increment to avoid subtracting large absolute temperatures."""
        mass = capacity * self.v / dt
        diag = mass.copy()
        diag[:-1] += G
        diag[1:] += G
        diag[-1] += self.ar * h
        ab = np.zeros((3, len(base)))
        ab[0, 1:] = -G
        ab[1] = diag
        ab[2, :-1] = -G
        delta = solve_banded((1, 1), ab, self.net(base, G, h, ambient),
                             check_finite=False)
        return base + delta

    def implicit(self, base, initial, t, dt, max_iter=60):
        """Converge both scaled changes and the original nonlinear stage defects."""
        T, C = initial[0].copy(), initial[1].copy()
        ambient = self.env(t)
        for iteration in range(1, max_iter + 1):
            a, k, _ = properties(T, C)
            T1 = self.linear(base[0], a, self.conductance(k), H, ambient[0], dt)
            _, _, D = properties(T1, C)
            C1 = self.linear(base[1], np.ones_like(C), self.conductance(D),
                             HM, ambient[1], dt)
            change = max(np.max(abs(T1 - T)) / 50, np.max(abs(C1 - C)) / 2.55)
            T, C = T1, C1
            if change <= ITER_TOL:
                a, k, D = properties(T, C)
                GT, GC = self.conductance(k), self.conductance(D)
                FT, FC = self.net(T, GT, H, ambient[0]), self.net(C, GC, HM, ambient[1])
                defects = []
                for u, b, cap, G, h, flux, scale in [
                    (T, base[0], a, GT, H, FT, 50.),
                    (C, base[1], np.ones_like(C), GC, HM, FC, 2.55)]:
                    diagonal = cap * self.v / dt
                    diagonal[:-1] += G
                    diagonal[1:] += G
                    diagonal[-1] += self.ar * h
                    residual = cap * self.v * (u - b) / dt - flux
                    defects.append(float(np.max(abs(residual) / (diagonal * scale))))
                equation = max(defects)
                if equation <= RES_TOL:
                    self.max_iter = max(self.max_iter, iteration)
                    self.max_change = max(self.max_change, float(change))
                    self.max_residual = max(self.max_residual, equation)
                    rate = np.stack([(T - base[0]) / dt, (C - base[1]) / dt])
                    heat_storage = float(np.sum(self.v * a * rate[0]))
                    boundary = self.ar * np.array([H, HM]) * (ambient - [T[-1], C[-1]])
                    return np.stack([T, C]), rate, np.r_[heat_storage, boundary]
        raise RuntimeError(f"Unconverged stage at {t:.12g}, dt={dt:.6g}")

    def step(self, old, t, dt):
        """L-stable second-order SDIRK; diagnostics use the same stage weights."""
        stage1, rate1, balance1 = self.implicit(old, old, t + GAMMA * dt, GAMMA * dt)
        base2 = old + dt * (1 - GAMMA) * rate1
        stage2, _, balance2 = self.implicit(base2, stage1, t + dt, GAMMA * dt)
        return stage2, dt * ((1 - GAMMA) * balance1 + GAMMA * balance2)

    def boundaries(self, state, t):
        """Independently reconstruct Robin and symmetry gradients from node values."""
        T, C = state
        _, k, D = properties(T, C)
        surface_grad = (3 * state[:, -1] - 4 * state[:, -2] + state[:, -3]) / (2 * self.dr)
        center_grad = (-3 * state[:, 0] + 4 * state[:, 1] - state[:, 2]) / (2 * self.dr)
        robin = -np.array([k[-1], D[-1]]) * surface_grad - np.array([H, HM]) * (
            state[:, -1] - self.env(t))
        return robin, center_grad


def refine_event(solver, old, t0, width, tolerance):
    """Bisect g using new implicit solves from the same last non-dry state."""
    lower, upper = 0., width
    low_state = old.copy()
    upper_state, _ = solver.step(old, t0, upper)
    trace = [[t0, float(old[1].max() - .15)],
             [t0 + upper, float(upper_state[1].max() - .15)]]
    if trace[0][1] < 0 or trace[1][1] >= 0:
        raise RuntimeError("Invalid event bracket")
    while upper - lower > tolerance:
        middle = (lower + upper) / 2
        state, _ = solver.step(old, t0, middle)
        g = float(state[1].max() - .15)
        trace.append([t0 + middle, g])
        if g < 0:
            upper, upper_state = middle, state
        else:
            lower, low_state = middle, state
    # A short forward margin resolves strict inequality beyond the bracket roundoff.
    end_offset = upper + .1
    end_state, balance = solver.step(old, t0, end_offset)
    if not np.all(end_state[1] < .15):
        raise RuntimeError("Strict full-field endpoint check failed")
    return {"lower_s": t0 + lower, "upper_s": t0 + upper,
            "estimate_s": t0 + (lower + upper) / 2,
            "end_s": t0 + end_offset, "g_lower": float(low_state[1].max() - .15),
            "g_upper": float(upper_state[1].max() - .15),
            "g_end": float(end_state[1].max() - .15),
            "width_s": upper - lower, "trace": trace}, end_state, balance


def simulate(n, factor, mode, name, event_tol=.001, stop_s=None):
    """Integrate from the uniform initial state until every fine-grid node is dry."""
    started = time.perf_counter()
    values, audit_report = audit()
    env = Environment(values, mode)
    solver = Solver(n, env)
    state = np.stack([np.full(n + 1, 28.), np.full(n + 1, 2.55)])
    initial = state.copy()
    times, outputs, full_times, full_states = [0.], [state[:, ::n // 20].copy()], [0.], [state.copy()]
    history = []
    balance_total = np.zeros(3)
    accepted, rejected, t, step_min, step_max = 0, 0, 0., 1e30, 0.
    radial_inversion = 0.
    temporal_increase = 0.
    temperature_decrease = 0.
    max_center_gap = 0.
    min_T, max_T, min_C, max_C = 28., 28., 2.55, 2.55
    min_D, max_D = 1., 0.
    next_output, next_six = 60., 21600.
    check_index = 0
    event = None
    limit = 72 * 3600.
    horizon_log = []
    while event is None:
        cap = (2. if t < 14400 else 30.) * factor
        dt = min(cap, max(.0005, .03 * (t + .01)) * factor)
        target = min(next_output, next_six, limit)
        if check_index < len(CHECK_TIMES):
            target = min(target, CHECK_TIMES[check_index])
        if stop_s is not None:
            target = min(target, stop_s)
        dt = min(dt, target - t)
        if dt <= 0:
            raise RuntimeError(f"Nonpositive step at {t}, target {target}")
        while True:
            try:
                new, balance = solver.step(state, t, dt)
                properties(*new)
                break
            except (RuntimeError, FloatingPointError, np.linalg.LinAlgError):
                rejected += 1
                dt /= 2
                if dt < 1e-7:
                    raise
        if np.max(new[1]) < .15 and stop_s is None:
            event, new, balance = refine_event(solver, state, t, dt, event_tol)
            np.savez_compressed(OUT / f"{name}_event_seed.npz", state=state, t=t, width=dt)
            dt = event["end_s"] - t
        accepted += 1
        balance_total += balance
        radial_inversion = max(radial_inversion, float(np.max(np.diff(new[1]))))
        temporal_increase = max(temporal_increase, float(np.max(new[1] - state[1])))
        temperature_decrease = min(temperature_decrease, float(np.min(new[0] - state[0])))
        max_center_gap = max(max_center_gap, float(new[1].max() - new[1, 0]))
        min_T, max_T = min(min_T, float(new[0].min())), max(max_T, float(new[0].max()))
        min_C, max_C = min(min_C, float(new[1].min())), max(max_C, float(new[1].max()))
        _, _, D = properties(*new)
        min_D, max_D = min(min_D, float(D.min())), max(max_D, float(D.max()))
        step_min, step_max = min(step_min, dt), max(step_max, dt)
        state, t = new, t + dt
        if abs(t - target) < 1e-8:
            t = target
        full = False
        if check_index < len(CHECK_TIMES) and abs(t - CHECK_TIMES[check_index]) < 1e-8:
            full = True
            check_index += 1
        if abs(t - next_six) < 1e-8:
            full = True
            next_six += 21600
        if abs(t - next_output) < 1e-8 or event is not None or (stop_s and t == stop_s):
            times.append(t)
            outputs.append(state[:, ::n // 20].copy())
            robin, symmetry = solver.boundaries(state, t)
            mass_change = float(np.dot(solver.v, state[1] - initial[1]))
            history.append([t, float(state[1].max()), float(solver.r[np.argmax(state[1])]),
                            float(np.dot(solver.v, state[1]) / R**2),
                            mass_change, *balance_total.tolist(), *robin.tolist(), *symmetry.tolist()])
            if event is None:
                next_output += 60
        if full or event is not None:
            full_times.append(t)
            full_states.append(state.copy())
        if t >= limit - 1e-8:
            horizon_log.append({"time_s": t, "max_C": float(state[1].max())})
            limit += 72 * 3600
            if limit > 10000 * 3600:
                raise RuntimeError("10000 h resource guard; not certified dry")
        if accepted % 4000 == 0:
            print(f"{name}: {t / 3600:.3f} h, Cmax={state[1].max():.9f}, "
                  f"steps={accepted}, elapsed={time.perf_counter()-started:.1f}s", flush=True)
        if stop_s is not None and t >= stop_s:
            break
    mass_change = float(np.dot(solver.v, state[1] - initial[1]))
    mass_rel = abs(mass_change - balance_total[2]) / max(abs(mass_change), abs(balance_total[2]), R**2)
    heat_rel = abs(balance_total[0] - balance_total[1]) / max(abs(balance_total[0]), abs(balance_total[1]), 1.)
    metadata = {"name": name, "N": n, "factor": factor, "mode": mode,
                "environment_tail": env.tail.tolist(), "event": event,
                "steps": accepted, "rejected_steps": rejected,
                "step_range_s": [step_min, step_max], "max_iterations": solver.max_iter,
                "max_scaled_change": solver.max_change, "max_scaled_residual": solver.max_residual,
                "mass_relative_defect": float(mass_rel), "heat_relative_defect": float(heat_rel),
                "mass_change_scaled": mass_change, "cumulative_heat_storage_scaled": balance_total[0],
                "cumulative_heat_boundary_scaled": balance_total[1],
                "cumulative_moisture_boundary_scaled": balance_total[2],
                "radial_inversion_max": radial_inversion,
                "temporal_C_increase_max": temporal_increase,
                "temperature_step_decrease_min": temperature_decrease,
                "max_C_above_center": max_center_gap,
                "temperature_range": [min_T, max_T], "moisture_range": [min_C, max_C],
                "diffusivity_range": [min_D, max_D], "horizon_extensions": horizon_log,
                "end_argmax_r_m": float(solver.r[np.argmax(state[1])]),
                "runtime_s": time.perf_counter() - started,
                "all_necessary_numerical_checks": "PENDING_CONVERGENCE"}
    if mass_rel > 1e-7 or heat_rel > 1e-7 or radial_inversion > 1e-9 or max_center_gap > 1e-9:
        metadata["all_necessary_numerical_checks"] = "FAIL"
    np.savez_compressed(OUT / f"{name}.npz", times=times, outputs=outputs,
                        full_times=full_times, full_states=full_states, r=solver.r,
                        history=history, final_state=state)
    dump_json(OUT / f"{name}.json", metadata)
    print(json.dumps(metadata, ensure_ascii=False), flush=True)
    if metadata["all_necessary_numerical_checks"] == "FAIL":
        raise RuntimeError("Numerical invariant check failed")
    return metadata


def main():
    """Run an audited isolated case without any previous-question entry points."""
    parser = argparse.ArgumentParser()
    parser.add_argument("--audit", action="store_true")
    parser.add_argument("--case", default="space800")
    parser.add_argument("--n", type=int, default=800)
    parser.add_argument("--factor", type=float, default=1.)
    parser.add_argument("--mode", choices=["last", "mean30"], default="last")
    parser.add_argument("--event-tol", type=float, default=.001)
    parser.add_argument("--stop", type=float)
    args = parser.parse_args()
    if args.audit:
        _, report = audit()
        print(json.dumps(report, ensure_ascii=False, indent=2))
    else:
        if args.n < 20 or args.n % 20 or args.factor <= 0:
            raise ValueError("N must be a positive multiple of 20; factor must be positive")
        simulate(args.n, args.factor, args.mode, args.case, args.event_tol, args.stop)


if __name__ == "__main__":
    main()
