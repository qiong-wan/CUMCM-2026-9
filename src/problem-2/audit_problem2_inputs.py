"""Read Problem 2 sources and report evidence without modifying input files."""

from __future__ import annotations

import ast
import hashlib
import json
import math
from pathlib import Path

import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[2]
ORIGINAL = ROOT.parent / "A题"
OUTPUT = ROOT / "output" / "problem-2" / "input_audit.json"


def sha256(path: Path) -> str:
    """Return the SHA-256 digest of a source file."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def audit_ambient(path: Path) -> dict:
    """Inspect every observation in original row order without repairing data."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.values)
    result = {
        "sheet": sheet.title,
        "dimensions": [sheet.max_row, sheet.max_column],
        "headers": list(rows[0]),
        "anomalies": [],
    }
    workbook.close()
    for index, row in enumerate(rows[1:], start=2):
        if len(row) != 3 or any(
            isinstance(value, bool)
            or not isinstance(value, (int, float))
            or not math.isfinite(value)
            for value in row
        ):
            result["anomalies"].append({"row": index, "values": list(row)})
    if result["anomalies"]:
        return result

    values = np.asarray(rows[1:], dtype=float)
    times = values[:, 0]
    differences = np.diff(times)
    for index in np.flatnonzero(differences <= 0):
        result["anomalies"].append(
            {"row": int(index + 3), "issue": "non-increasing time"}
        )
    for index in np.flatnonzero(
        (times < 0) | (values[:, 1] <= -273.15) | (values[:, 2] < 0)
    ):
        result["anomalies"].append(
            {"row": int(index + 2), "issue": "invalid physical range"}
        )
    selected = (times >= 0) & (times <= 10800)
    result.update(
        observation_count=len(times),
        time_range_s=[float(times.min()), float(times.max())],
        time_steps_s=np.unique(differences).tolist(),
        whole_file_min=values.min(axis=0).tolist(),
        whole_file_max=values.max(axis=0).tolist(),
        problem2_observation_count=int(selected.sum()),
        problem2_excel_rows=[2, int(np.flatnonzero(selected)[-1] + 2)],
        problem2_min=values[selected].min(axis=0).tolist(),
        problem2_max=values[selected].max(axis=0).tolist(),
        temperature_decreasing_intervals=int(np.sum(np.diff(values[:, 1]) < 0)),
        moisture_decreasing_intervals=int(np.sum(np.diff(values[:, 2]) < 0)),
        selected_observations=[
            {"excel_row": i + 2, "values": row.tolist()}
            for i, row in enumerate(values)
            if row[0] in (0, 1800, 3600, 5400, 7200, 9000, 10800, 14400)
        ],
    )
    if not (times[0] == 0 and times[-1] >= 10800):
        result["anomalies"].append({"issue": "insufficient 3-hour coverage"})
    if not result["anomalies"]:
        interpolated = np.column_stack(
            [np.interp(times, times, values[:, column]) for column in (1, 2)]
        )
        result["interpolation_at_observations_max_error"] = float(
            np.max(np.abs(interpolated - values[:, 1:]))
        )
    return result


def inspect_workbook(path: Path, is_template: bool = False) -> dict:
    """Report labels and representative values from a template or prior result."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    result = {}
    for sheet in workbook:
        rows = list(sheet.values)
        item = {
            "dimensions": [sheet.max_row, sheet.max_column],
            "header": list(rows[0]),
            "first_row": list(rows[1]),
            "last_row": list(rows[-1]),
            "B2_number_format": sheet["B2"].number_format,
        }
        if is_template:
            item["role"] = "Blank result cells and ellipses are template placeholders."
        else:
            matrix = np.asarray(rows[1:], dtype=float)
            item["all_values_finite"] = bool(np.isfinite(matrix).all())
            item["field_range"] = [
                float(matrix[:, 1:].min()), float(matrix[:, 1:].max())
            ]
            item["unique_time_steps_s"] = np.unique(np.diff(matrix[:, 0])).tolist()
        result[sheet.title] = item
    workbook.close()
    return result


def inspect_problem1_code(path: Path) -> dict:
    """Check syntax without importing or executing the Problem 1 solver."""
    source = path.read_text(encoding="utf-8-sig")
    result = {
        "conflict_start_lines": [
            index for index, line in enumerate(source.splitlines(), start=1)
            if line.startswith("<<<<<<<")
        ]
    }
    try:
        ast.parse(source, filename=str(path))
        result["syntax_valid"] = True
    except SyntaxError as error:
        result.update(syntax_valid=False, syntax_error_line=error.lineno,
                      syntax_error=error.msg)
    return result


def property_snapshots() -> list[dict]:
    """Evaluate Appendix 3 at declared states, not simulated observations."""
    result = []
    for temperature, moisture in ((28.0, 2.55), (50.0, 2.55), (50.0, 0.15)):
        kelvin = temperature + 273.15
        density = 650.0 + 128.0 * moisture
        heat_capacity = 1450.0 + 2736.0 * moisture / (moisture + 1.0)
        conductivity = 0.21 + 0.38 * moisture / (moisture + 1.0)
        diffusivity = 2.4e-3 * math.exp(-0.45 / moisture - 3850.0 / kelvin)
        result.append({
            "temperature_C": temperature, "moisture_kg_kg": moisture,
            "rho_kg_m3": density, "cp_J_kg_K": heat_capacity,
            "k_W_m_K": conductivity, "D_m2_s": diffusivity,
            "alpha_m2_s": conductivity / (density * heat_capacity),
            "Bi_heat": 25.0 * 0.02 / conductivity,
            "Bi_moisture": 8e-7 * 0.02 / diffusivity,
            "dry_density_if_rho_is_wet": density / (1.0 + moisture),
        })
    return result


def main() -> None:
    """Write a derived audit report and verify all inspected inputs are intact."""
    pairs = [
        (ROOT / "problem.pdf", ORIGINAL / "A题.pdf"),
        (ROOT / "data" / "附件1.xlsx", ORIGINAL / "附件" / "附件1.xlsx"),
        (ROOT / "data" / "附件3" / "result2.xlsx",
         ORIGINAL / "附件" / "附件3" / "result2.xlsx"),
    ]
    solver = ROOT / "src" / "problem-1" / "solve_problem1.py"
    prior_result = ROOT / "output" / "problem-1" / "result1.xlsx"
    paths = [path for pair in pairs for path in pair] + [solver, prior_result]
    before = {str(path): sha256(path) for path in paths}
    report = {
        "source_sha256": before,
        "source_copy_matches": [
            {"project_file": str(first.relative_to(ROOT)),
             "matches_original": before[str(first)] == before[str(second)]}
            for first, second in pairs
        ],
        "ambient": audit_ambient(pairs[1][0]),
        "result2_template": inspect_workbook(pairs[2][0], is_template=True),
        "problem1_code": inspect_problem1_code(solver),
        "problem1_workbook": inspect_workbook(prior_result),
        "appendix3_property_snapshots_not_predictions": property_snapshots(),
        "scope": "Input audit and framework evidence only; no Problem 2 PDE solve.",
    }
    report["all_inspected_files_unchanged"] = all(
        sha256(path) == before[str(path)] for path in paths
    )
    assert report["all_inspected_files_unchanged"]
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n",
                      encoding="utf-8")
    print(json.dumps({
        "report": str(OUTPUT),
        "ambient_anomalies": report["ambient"]["anomalies"],
        "raw_unchanged": report["all_inspected_files_unchanged"],
        "problem1_syntax_valid": report["problem1_code"]["syntax_valid"],
        "property_snapshots": report["appendix3_property_snapshots_not_predictions"],
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
