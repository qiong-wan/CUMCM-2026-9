"""Generate six paper figures from existing Problem 2 results, without solving.

Run from any working directory:
    python src/problem-2/generate_figures.py

PNG and SVG figures go to output/problem-2/image. Render previews are written
directly from the figure canvas to tmp/problem2_figure_previews; output PNGs
are never read back. Source workbooks and verification records are read-only.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
import warnings

import matplotlib

matplotlib.use("Agg")

from matplotlib import font_manager
from matplotlib import pyplot as plt
from matplotlib.ticker import ScalarFormatter
import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[2]
AMBIENT_FILE = ROOT / "data" / "附件1.xlsx"
RESULT_FILE = ROOT / "output" / "problem-2" / "result2.xlsx"
VERIFICATION_FILE = ROOT / "output" / "problem-2" / "verification.txt"
DEFAULT_IMAGE_DIR = ROOT / "output" / "problem-2" / "image"
DEFAULT_PREVIEW_DIR = ROOT.parent / "tmp" / "problem2_figure_previews"
T_END = 10800
T_INIT = 28.0
C_INIT = 2.55
SELECTED_TIMES = (1800, 3600, 5400, 7200, 9000, 10800)
SELECTED_RADII_CM = (0.0, 0.5, 1.0, 1.5, 2.0)
TARGET = 5.0e-5
COLORS = ("#0072B2", "#009E73", "#E69F00", "#D55E00", "#CC79A7", "#56B4E9")
TEMP_COLOR = "#C44E52"
MOISTURE_COLOR = "#277DA1"
NEUTRAL_COLOR = "#3D4652"
GRID_COLOR = "#D8DEE6"
NUMBER = r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"


@dataclass(frozen=True)
class Fields:
    """Validated workbook fields on a common time and radial grid."""

    times: np.ndarray
    radii_cm: np.ndarray
    temperature: np.ndarray
    moisture: np.ndarray
    initial_state_added: bool


def sha256(path: Path) -> str:
    """Return a file fingerprint for input-preservation and provenance checks."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def configure_matplotlib() -> str:
    """Use the established Chinese publication style with a verified font."""
    installed = {item.name for item in font_manager.fontManager.ttflist}
    candidates = ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC")
    font = next((name for name in candidates if name in installed), None)
    if font is None:
        raise RuntimeError("需要安装支持中文的字体：Microsoft YaHei、SimHei 或 Noto Sans CJK SC。")
    plt.rcParams.update({
        "font.family": "sans-serif", "font.sans-serif": [font, "DejaVu Sans"],
        "axes.unicode_minus": False, "mathtext.fontset": "stix",
        "font.size": 10, "axes.titlesize": 11, "axes.labelsize": 10,
        "legend.fontsize": 9, "xtick.labelsize": 9, "ytick.labelsize": 9,
        "axes.edgecolor": NEUTRAL_COLOR, "axes.labelcolor": "#20242A",
        "text.color": "#20242A", "xtick.color": "#343A40", "ytick.color": "#343A40",
        "savefig.facecolor": "white", "svg.fonttype": "path",
        "axes.formatter.use_mathtext": True,
    })
    warnings.filterwarnings("error", message="Glyph .* missing from font.*")
    return font


def load_ambient() -> np.ndarray:
    """Audit all observations and select the supported 0--3 h interval."""
    workbook = openpyxl.load_workbook(AMBIENT_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        if sheet.max_column != 3:
            raise ValueError("附件1必须包含时间、温度、水分浓度三列。")
        records = list(sheet.iter_rows(min_row=2, values_only=True))
        for row_number, row in enumerate(records, start=2):
            if any(value is None or isinstance(value, (str, bool)) for value in row):
                raise ValueError(f"附件1第{row_number}行有缺失或非数值，未作修改。")
        values = np.asarray(records, dtype=float)
    finally:
        workbook.close()
    if not np.all(np.isfinite(values)) or np.any(np.diff(values[:, 0]) <= 0):
        raise ValueError("附件1有非有限值、重复或逆序时间，未作修改。")
    if np.any(values[:, 0] < 0) or np.any(values[:, 1] <= -273.15) or np.any(values[:, 2] < 0):
        raise ValueError("附件1有超出物理定义域的数值，未作修改。")
    selected = values[(values[:, 0] >= 0) & (values[:, 0] <= T_END)]
    if selected.shape != (181, 3) or not np.array_equal(selected[:, 0], np.arange(0, T_END + 1, 60)):
        raise ValueError("附件1前3小时必须完整覆盖0--10800 s，间隔为60 s。")
    return selected


def load_results() -> Fields:
    """Read both result sheets without changing values, order, or formatting."""
    workbook = openpyxl.load_workbook(RESULT_FILE, read_only=True, data_only=True)
    sheets = []
    try:
        for name in ("温度", "水分浓度"):
            if name not in workbook.sheetnames:
                raise ValueError(f"result2.xlsx 缺少工作表：{name}")
            rows = list(workbook[name].iter_rows(values_only=True))
            if len(rows) not in (T_END + 1, T_END + 2) or any(len(row) != 22 for row in rows):
                raise ValueError(f"工作表{name}的行列数与3小时、21个径向点不符。")
            if any(value is None or isinstance(value, (str, bool)) for row in rows[1:] for value in row):
                raise ValueError(f"工作表{name}存在缺失或非数值结果，未作修正。")
            radii = np.asarray(rows[0][1:], dtype=float)
            values = np.asarray(rows[1:], dtype=float)
            times, field = values[:, 0], values[:, 1:]
            if not np.all(np.isfinite(values)) or not np.allclose(radii, np.arange(21) / 10, rtol=0, atol=1e-10):
                raise ValueError(f"工作表{name}存在非有限值或错误的径向网格。")
            if times[0] not in (0, 1) or not np.array_equal(times, np.arange(times[0], T_END + 1)):
                raise ValueError(f"工作表{name}必须从0或1 s开始，逐秒覆盖到10800 s。")
            sheets.append((times, radii, field))
    finally:
        workbook.close()
    if not np.array_equal(sheets[0][0], sheets[1][0]) or not np.array_equal(sheets[0][1], sheets[1][1]):
        raise ValueError("温度与水分浓度工作表的时间或空间坐标不一致。")
    times, radii, temperature = sheets[0]
    moisture = sheets[1][2]
    if np.any(temperature <= -273.15) or np.any(moisture < 0):
        raise ValueError("工作簿存在超出物理定义域的结果，未作修正。")
    add_initial = bool(times[0] == 1)
    if add_initial:
        times = np.insert(times, 0, 0.0)
        temperature = np.vstack((np.full(21, T_INIT), temperature))
        moisture = np.vstack((np.full(21, C_INIT), moisture))
    return Fields(times, radii, temperature, moisture, add_initial)


def load_convergence() -> dict:
    """Parse the current solver's text report without rerunning convergence tests."""
    content = VERIFICATION_FILE.read_text(encoding="utf-8-sig")
    spatial_pattern = re.compile(
        rf"^\s*N ->\s*(\d+):\s*whole max\|dT\|=({NUMBER}), max\|dC\|=({NUMBER})"
        rf"\s*\| table max\|dT\|=({NUMBER}), max\|dC\|=({NUMBER}).*$", re.MULTILINE)
    temporal_pattern = re.compile(
        rf"^\s*factor (\d+)x:\s*whole max\|dT\|=({NUMBER}), max\|dC\|=({NUMBER})"
        rf"\s*\| table max\|dT\|=({NUMBER}), max\|dC\|=({NUMBER}).*$", re.MULTILINE)
    startup_pattern = re.compile(
        rf"^\s*dt = ({NUMBER}) s:\s*C\(1 s, R\) = ({NUMBER})\s+diff =\s*({NUMBER}|--)\s*$", re.MULTILINE)
    spatial = [[int(match[0]), *map(float, match[1:])] for match in spatial_pattern.findall(content)]
    temporal = [[int(match[0]), *map(float, match[1:])] for match in temporal_pattern.findall(content)]
    startup = [{"dt_s": float(dt), "surface_C_1s": float(value),
                "adjacent_difference": None if diff == "--" else float(diff)}
               for dt, value, diff in startup_pattern.findall(content)]
    grid_condition = re.search(r"Grid convergence \(uniform nsub = (\d+),", content)
    time_condition = re.search(r"Temporal convergence, ramped schedule \(N = (\d+),", content)
    startup_condition = re.search(r"Start-up convergence \(N = (\d+),", content)
    if not spatial or not temporal or not startup or not all((grid_condition, time_condition, startup_condition)):
        raise ValueError("验证记录缺少当前格式的空间、时间或启动收敛数据；可用 --skip-convergence 只生成分布图。")
    for label, rows in (("空间", spatial), ("时间", temporal)):
        array = np.asarray(rows, dtype=float)
        if not np.all(np.isfinite(array)) or np.any(array[:, 1:] <= 0) or np.any(np.diff(array[:, 0]) <= 0):
            raise ValueError(f"{label}收敛记录无效，未删改记录。")
    finite_startup = np.array([[row["dt_s"], row["surface_C_1s"]] for row in startup])
    if not np.all(np.isfinite(finite_startup)) or np.any(finite_startup <= 0):
        raise ValueError("启动收敛记录包含非法数值。")
    difference_rows = [row for row in startup if row["adjacent_difference"] is not None]
    if not difference_rows or any(not np.isfinite(row["adjacent_difference"]) or row["adjacent_difference"] <= 0 for row in difference_rows):
        raise ValueError("启动收敛记录缺少有效的相邻步长差值。")
    status = re.search(r"^STATUS: (.+)$", content, re.MULTILINE)
    return {"spatial": spatial, "temporal": temporal, "startup_all_records": startup,
            "spatial_nsub": int(grid_condition[1]), "temporal_N": int(time_condition[1]),
            "startup_N": int(startup_condition[1]),
            "reported_status": status[1] if status else "Not specified", "recomputed": False}


def style_axis(axis: plt.Axes) -> None:
    """Use unobtrusive grids and remove redundant frame edges."""
    axis.spines[["top", "right"]].set_visible(False)
    axis.grid(True, color=GRID_COLOR, linewidth=0.65, alpha=0.8)
    axis.set_axisbelow(True)


def time_axis(axis: plt.Axes) -> None:
    """Format a shared three-hour time axis."""
    axis.set(xlim=(0, 3), xlabel="时间 / h", xticks=np.arange(0, 3.01, 0.5))
    style_axis(axis)


def make_ambient_figure(ambient: np.ndarray) -> plt.Figure:
    """Show every original ambient observation and its piecewise-linear boundary."""
    figure, axes = plt.subplots(2, 1, figsize=(10.6, 6.0), sharex=True, layout="constrained")
    figure.suptitle("前3小时烘房环境变化", fontsize=13)
    seconds = np.arange(T_END + 1)
    for axis, column, color, title, ylabel in zip(
        axes, (1, 2), (TEMP_COLOR, MOISTURE_COLOR),
        ("(a) 烘房温度", "(b) 烘房水分浓度"),
        ("温度 / °C", "水分浓度 / (kg/kg)"), strict=True,
    ):
        axis.plot(seconds / 3600, np.interp(seconds, ambient[:, 0], ambient[:, column]),
                  color=color, linewidth=1.7, label="分段线性边界")
        axis.scatter(ambient[:, 0] / 3600, ambient[:, column], s=10,
                     facecolors="white", edgecolors=color, linewidths=0.7, zorder=3, label="附件1观测")
        axis.set(title=title, ylabel=ylabel)
        time_axis(axis)
    axes[0].set_xlabel("")
    axes[0].legend(frameon=False, ncol=2, loc="lower right")
    return figure


def make_history_figure(fields: Fields, ambient: np.ndarray) -> plt.Figure:
    """Show temperature and moisture histories at the five paper radii."""
    figure, axes = plt.subplots(1, 2, figsize=(11.4, 4.7), layout="constrained")
    figure.suptitle("典型径向位置的温度与含水率变化", fontsize=13)
    markers = ("o", "s", "^", "D", "v")
    for radius, color, marker in zip(SELECTED_RADII_CM, COLORS[:5], markers, strict=True):
        column = int(round(radius * 10))
        for axis, data in zip(axes, (fields.temperature, fields.moisture), strict=True):
            axis.plot(fields.times / 3600, data[:, column], color=color, linewidth=1.5,
                      marker=marker, markersize=3.8, markerfacecolor="white",
                      markevery=900, label=f"r = {radius:g} cm")
    axes[0].plot(ambient[:, 0] / 3600, ambient[:, 1], color=NEUTRAL_COLOR,
                 linewidth=1.25, linestyle="--", label="烘房温度")
    axes[0].set(title="(a) 温度时间历程", ylabel="温度 / °C")
    axes[1].set(title="(b) 干基含水率时间历程", ylabel="干基含水率 / (kg/kg)")
    for axis in axes:
        time_axis(axis)
    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(handles, labels, loc="outside lower center", ncol=6, frameon=False)
    return figure


def make_profile_figure(fields: Fields) -> plt.Figure:
    """Connect the actual 21 radial nodes at all six required paper times."""
    figure, axes = plt.subplots(1, 2, figsize=(11.4, 4.7), layout="constrained")
    figure.suptitle("典型时刻的径向温度与含水率分布", fontsize=13)
    for second, color, marker in zip(SELECTED_TIMES, COLORS, ("o", "s", "^", "D", "v", "P"), strict=True):
        for axis, data in zip(axes, (fields.temperature, fields.moisture), strict=True):
            axis.plot(fields.radii_cm, data[second], color=color, linewidth=1.6,
                      marker=marker, markersize=4, markerfacecolor="white", markevery=5,
                      label=f"{second / 3600:.1f} h")
    axes[0].set(title="(a) 温度径向剖面", ylabel="温度 / °C")
    axes[1].set(title="(b) 干基含水率径向剖面", ylabel="干基含水率 / (kg/kg)")
    for axis in axes:
        axis.set(xlim=(0, 2), xticks=np.arange(0, 2.01, 0.5), xlabel="到药材中心的距离 / cm")
        style_axis(axis)
    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(handles, labels, loc="outside lower center", ncol=6, frameon=False)
    return figure


def clipped_edges(nodes: np.ndarray) -> np.ndarray:
    """Bound nearest-node display intervals at the physical domain endpoints."""
    return np.concatenate(([nodes[0]], (nodes[:-1] + nodes[1:]) / 2, [nodes[-1]]))


def make_field_figure(fields: Fields) -> plt.Figure:
    """Show unsmoothed nearest-node fields, without invented contour detail."""
    figure, axes = plt.subplots(1, 2, figsize=(11.4, 4.8), layout="constrained")
    figure.suptitle("前3小时温度与含水率的时空分布", fontsize=13)
    x_edges = clipped_edges(fields.times / 3600)
    y_edges = clipped_edges(fields.radii_cm)
    for axis, data, cmap, title, unit in zip(
        axes, (fields.temperature, fields.moisture), ("inferno", "viridis"),
        ("(a) 温度场", "(b) 干基含水率场"), ("温度 / °C", "干基含水率 / (kg/kg)"), strict=True,
    ):
        mesh = axis.pcolormesh(x_edges, y_edges, data.T, shading="flat", cmap=cmap,
                              vmin=float(data.min()), vmax=float(data.max()), rasterized=True)
        axis.set(title=title, xlabel="时间 / h", ylabel="到药材中心的距离 / cm",
                 xlim=(0, 3), ylim=(0, 2), xticks=np.arange(0, 3.01, 0.5), yticks=np.arange(0, 2.01, 0.5))
        figure.colorbar(mesh, ax=axis, pad=0.025, label=unit)
    return figure


def make_spatial_figure(convergence: dict) -> plt.Figure:
    """Plot recorded adjacent-grid differences, with fixed time-step context."""
    data = np.asarray(convergence["spatial"], dtype=float)
    figure, axes = plt.subplots(1, 2, figsize=(11.4, 4.7), layout="constrained")
    dt = 1 / convergence["spatial_nsub"]
    figure.suptitle(f"空间网格加密差（固定时间步 {dt:g} s）", fontsize=13)
    for axis, full_column, table_column, title, unit in zip(
        axes, (1, 2), (3, 4), ("(a) 温度", "(b) 干基含水率"),
        ("相邻网格最大差 / °C", "相邻网格最大差 / (kg/kg)"), strict=True,
    ):
        axis.plot(data[:, 0], data[:, full_column], "o-", color=COLORS[0], linewidth=1.7, label="完整交付网格")
        axis.plot(data[:, 0], data[:, table_column], "s--", color=COLORS[1], linewidth=1.7, label="论文表指定点")
        axis.axhline(TARGET, color=NEUTRAL_COLOR, linestyle=":", linewidth=1.4, label=r"参考阈值 $5\times10^{-5}$")
        axis.set_xscale("log", base=2)
        axis.set_yscale("log")
        axis.set(title=title, ylabel=unit, xlabel="径向区间数 N（与 N/2 比较）", xticks=data[:, 0])
        axis.xaxis.set_major_formatter(ScalarFormatter())
        style_axis(axis)
    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(handles, labels, loc="outside lower center", ncol=3, frameon=False)
    return figure


def make_temporal_figure(convergence: dict) -> plt.Figure:
    """Show actual schedule comparisons and startup differences without fitting."""
    data = np.asarray(convergence["temporal"], dtype=float)
    startup = [row for row in convergence["startup_all_records"] if row["adjacent_difference"] is not None]
    figure, axes = plt.subplots(1, 3, figsize=(13.6, 4.8), layout="constrained")
    figure.suptitle("时间步加密与启动阶段收敛", fontsize=13)
    positions = np.arange(len(data))
    previous_factors = np.concatenate(([1], data[:-1, 0]))
    labels = [f"{first:g}x → {second:g}x" for first, second in zip(previous_factors, data[:, 0], strict=True)]
    for axis, full_column, table_column, title, ylabel in zip(
        axes[:2], (1, 2), (3, 4), ("(a) 温度时间步加密差", "(b) 含水率时间步加密差"),
        ("最大差 / °C", "最大差 / (kg/kg)"), strict=True,
    ):
        for offset, column, color, label in ((-0.18, full_column, COLORS[0], "完整交付网格"),
                                            (0.18, table_column, COLORS[1], "论文表指定点")):
            bars = axis.bar(positions + offset, data[:, column], width=0.32, color=color, label=label)
            axis.bar_label(bars, labels=[f"{value:.2e}" for value in data[:, column]], padding=4, fontsize=8)
        axis.axhline(TARGET, color=NEUTRAL_COLOR, linestyle=":", linewidth=1.4, label=r"参考阈值 $5\times10^{-5}$")
        axis.set(title=title, ylabel=ylabel, xticks=positions, xticklabels=labels,
                 xlabel=f"子步数放大倍数（固定 N={convergence['temporal_N']}）")
        axis.set_ylim(0, max(TARGET, float(data[:, [full_column, table_column]].max())) * 1.32)
        axis.ticklabel_format(axis="y", style="sci", scilimits=(0, 0))
        style_axis(axis)
    axis = axes[2]
    dt = [row["dt_s"] for row in startup]
    errors = [row["adjacent_difference"] for row in startup]
    axis.plot(dt, errors, "o-", color=COLORS[3], linewidth=1.7)
    axis.set_xscale("log", base=2)
    axis.set_yscale("log")
    axis.set_xticks(dt, [f"{value:.6f}" for value in dt])
    axis.tick_params(axis="x", labelsize=8)
    axis.set(title=f"(c) 启动加密差（N={convergence['startup_N']}）",
             xlabel="加密后时间步 / s", ylabel="1 s表面含水率相邻差 / (kg/kg)")
    style_axis(axis)
    handles, legend_labels = axes[0].get_legend_handles_labels()
    figure.legend(handles, legend_labels, loc="outside lower center", ncol=3, frameon=False)
    return figure


def save_figure(figure: plt.Figure, name: str, output_dir: Path, preview_dir: Path, dpi: int) -> dict:
    """Save paper assets and inspect the live canvas without reading output PNGs."""
    figure.canvas.draw()
    renderer = figure.canvas.get_renderer()
    outside_text = []
    undrawn_tick_labels = set()
    # Matplotlib retains tick artists beyond the displayed axis limits.
    for axis in figure.axes:
        for coordinate in (axis.xaxis, axis.yaxis):
            lower, upper = sorted(coordinate.get_view_interval())
            ticks = coordinate.get_major_ticks() + coordinate.get_minor_ticks()
            for tick in ticks:
                if not lower <= tick.get_loc() <= upper:
                    undrawn_tick_labels.update((tick.label1, tick.label2))
    for item in figure.findobj(matplotlib.text.Text):
        if (
            item in undrawn_tick_labels
            or not item.get_visible()
            or not item.get_text().strip()
        ):
            continue
        box = item.get_window_extent(renderer)
        if box.width and box.height and (box.x0 < -2 or box.y0 < -2 or box.x1 > figure.bbox.width + 2 or box.y1 > figure.bbox.height + 2):
            outside_text.append(item.get_text())
    canvas = np.asarray(figure.canvas.buffer_rgba())[:, :, :3]
    fraction = float(np.mean(np.any(canvas < 245, axis=2)))
    if fraction < 0.01:
        raise RuntimeError(f"{name}画布近乎空白，停止输出。")
    for extension in ("png", "svg"):
        figure.savefig(output_dir / f"{name}.{extension}", dpi=dpi, bbox_inches="tight", pad_inches=0.08)
    figure.savefig(preview_dir / f"{name}.jpg", dpi=115, bbox_inches="tight", pad_inches=0.08)
    info = {"name": name, "files": [f"{name}.png", f"{name}.svg"],
            "axes_count": len(figure.axes), "live_canvas_nonwhite_fraction": fraction,
            "text_outside_original_canvas": outside_text}
    plt.close(figure)
    return info


def write_notes(output_dir: Path, manifest: dict) -> None:
    """Document figure captions, transformations, and limits of the evidence."""
    notes = """# 问题2图像说明

全部物理分布图仅展示前3小时。PNG为300 dpi默认输出，SVG中的文字转为路径以避免字体替换；时空热图栅格化，其余曲线保留矢量。

| 文件 | 建议图题及论文用途 |
|---|---|
| fig1_ambient_conditions | 前3小时烘房温度与水分浓度变化。点为附件观测，实线为分段线性边界；不标注未经题面确定的阶段分界。 |
| fig2_time_histories | 典型径向位置的温度与含水率时间历程。两子图相同颜色、标记对应相同半径；温度图中的虚线为烘房温度。 |
| fig3_radial_profiles | 0.5至3 h内每隔0.5 h的温度与含水率径向剖面。标记位于题目指定的五个位置，曲线经过全部21个实际输出节点。 |
| fig4_spatiotemporal_fields | 前3小时温度与含水率时空分布。颜色分别表示°C和kg/kg，按实际网格作分块显示，不进行空间平滑或额外插值。 |
| fig5_spatial_convergence | 固定时间步下，相邻空间网格在完整交付网格及论文指定点上的最大差。空间误差不能替代时间误差。 |
| fig6_temporal_convergence | 固定空间网格下，生产子步数整体放大后的时间加密差，以及启动阶段1 s表面含水率的相邻步长差。 |

## 数据处理

1. 附件1全表先作数值与时间顺序审计，然后按本问范围选取0至10800 s的181个观测，不删改异常记录。
2. 工作簿的两个结果表按原始顺序读取，检查完整逐秒时间轴、21个径向节点及有限值。若从1 s开始，则仅在绘图数组前添加题给的0 s均匀初值28°C、2.55 kg/kg，这不是缺失值填补。
3. 时间由秒除以3600转换为小时；半径直接使用工作簿的cm坐标。除环境的分段线性边界外，不对结果作平滑、拟合或重采样，曲线保留所有逐秒值。
4. 典型位置选择0、0.5、1、1.5、2 cm；典型时刻选择0.5、1、1.5、2、2.5、3 h。所有抽样位置和时刻均实际存在。
5. 热图在相邻节点中点划分显示区间，并把端点截在0至3 h、0至2 cm内，避免将绘图区延伸到药材外部。颜色范围取各场实际最小值和最大值。
6. 收敛数据读取当前verification.txt，不调用求解器，也不使用旧review目录的对照结果。完整记录保存在figure_manifest.json中。启动记录中没有相邻差的基准与生产方案确认行保留在清单内，不加入相邻差曲线。

## 收敛量的定义

对变量u，相邻计算方案在相同输出点的最大差定义为：

$$
E_u=\\max_{(t,r)\\in\\mathcal G}|u_{\\mathrm{fine}}(r,t)-u_{\\mathrm{coarse}}(r,t)|.\\tag{1}
$$

公式（1）的温度单位为°C，含水率单位为kg/kg；完整交付网格与论文指定点分别计算。图中5×10⁻⁵为相应变量单位下的参考目标。它是相邻方案的自收敛比较，不是已证明的连续方程真实误差上界。

启动图采用验证记录中已经计算的差：

$$
E_{C,1}=|C_{\\Delta t_{\\mathrm{fine}}}(R,1)-C_{\\Delta t_{\\mathrm{coarse}}}(R,1)|.\\tag{2}
$$

公式（2）仅反映1 s表面点，不能替代全场时间收敛。横坐标引用验证记录中打印的步长，可能已有显示舍入。若报告只有一次时间加密对照，绘制分组柱图，不虚构多级收敛曲线。

## 使用边界

这些图展示的是现有工作簿和验证记录。作图成功不等于重新完成数值验收；验证文本的状态仅记录在清单内。没有运行72 h过程，没有制作潜热能量收支图，没有把温度作为严格物理上界。

复现命令：`python src/problem-2/generate_figures.py`。可用`--dpi`指定PNG分辨率，或用`--skip-convergence`只生成前四张分布图。实际输出清单与输入SHA-256见figure_manifest.json。
"""
    notes = notes.replace("PNG为300 dpi默认输出", f"本次PNG为{manifest['dpi']} dpi输出")
    (output_dir / "图像说明.md").write_text(notes, encoding="utf-8")


def main() -> None:
    """Audit inputs, generate figures, and record unmodified-source fingerprints."""
    parser = argparse.ArgumentParser(description="Generate Problem 2 paper figures from existing outputs.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_IMAGE_DIR)
    parser.add_argument("--preview-dir", type=Path, default=DEFAULT_PREVIEW_DIR)
    parser.add_argument("--dpi", type=int, default=300)
    parser.add_argument("--skip-convergence", action="store_true")
    args = parser.parse_args()
    if not 100 <= args.dpi <= 600:
        parser.error("--dpi must be between 100 and 600")
    paths = [AMBIENT_FILE, RESULT_FILE] + ([] if args.skip_convergence else [VERIFICATION_FILE])
    before = {str(path): sha256(path) for path in paths}
    ambient = load_ambient()
    fields = load_results()
    convergence = None if args.skip_convergence else load_convergence()
    font = configure_matplotlib()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    args.preview_dir.mkdir(parents=True, exist_ok=True)
    generators = [
        ("fig1_ambient_conditions", lambda: make_ambient_figure(ambient)),
        ("fig2_time_histories", lambda: make_history_figure(fields, ambient)),
        ("fig3_radial_profiles", lambda: make_profile_figure(fields)),
        ("fig4_spatiotemporal_fields", lambda: make_field_figure(fields)),
    ]
    if convergence is not None:
        generators += [("fig5_spatial_convergence", lambda: make_spatial_figure(convergence)),
                       ("fig6_temporal_convergence", lambda: make_temporal_figure(convergence))]
    outputs = []
    for name, generate in generators:
        info = save_figure(generate(), name, args.output_dir, args.preview_dir, args.dpi)
        outputs.append(info)
        print(f"Generated {name}.png and .svg", flush=True)
    unchanged = all(sha256(path) == before[str(path)] for path in paths)
    if not unchanged:
        raise RuntimeError("输入在绘图期间发生变化，请按最新输入重新生成。")
    manifest = {
        "font": font, "dpi": args.dpi, "source_sha256": before,
        "source_files_unchanged": unchanged, "initial_state_added_for_plotting": fields.initial_state_added,
        "time_range_s": [0, T_END], "ambient_observations": len(ambient),
        "field_shape_with_initial_state": list(fields.temperature.shape),
        "temperature_range_C": [float(fields.temperature.min()), float(fields.temperature.max())],
        "moisture_range_kg_kg": [float(fields.moisture.min()), float(fields.moisture.max())],
        "figures": outputs, "convergence": convergence,
        "preview_directory": str(args.preview_dir.resolve()), "output_pngs_read_back": False,
    }
    write_notes(args.output_dir, manifest)
    (args.output_dir / "figure_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Completed {len(outputs)} figures; all inputs unchanged.", flush=True)


if __name__ == "__main__":
    main()
