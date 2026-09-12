"""Problem 4: audited moving-cylinder finite volumes and reproducible delivery.

All non-graphical implementation is in this file. Importing is side-effect free.
Use --audit, --self-test, --case NAME, --verify, or --all from any directory.
"""

from __future__ import annotations

import argparse
import ast
import csv
import hashlib
import json
import math
import platform
import sys
import time
from concurrent.futures import ProcessPoolExecutor
from dataclasses import asdict, dataclass
from pathlib import Path

import numpy as np
import openpyxl
import scipy
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from scipy.integrate import solve_ivp
from scipy.linalg import solveh_banded
from scipy.signal import lfilter
from scipy.sparse import diags
from threadpoolctl import threadpool_info, threadpool_limits


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output" / "problem-4"
SOURCE = Path(__file__).resolve()
SCHEMA = 1
GAMMA = 1.0 - 1.0 / math.sqrt(2.0)
ITER_TOL = 1e-11
_BLAS_CONTROLLER = None
RES_TOL = 1e-11
H = 25.0
HM = 8e-7
R0 = 0.02
LENGTH = 0.25
C0 = 2.55
THRESHOLD = 0.15
START_TIMES = np.array([0.01, 0.1, 1.0, 10.0])
TARGETS = {
    "field_C_kg_kg": 5e-5,
    "field_T_C": 5e-5,
    "fluct_field_T_C": 1e-3,
    "event_change_s": 5.0,
    "balance_relative": 1e-7,
    "robin_heat_W_m2": 1e-3,
    "robin_moisture_kg_kg_m_s": 1e-9,
    "axis_T_C_m": 1e-5,
    "axis_C_kg_kg_m": 1e-5,
    "event_local_change_s": 0.01,
    "independent_T_C": 2e-5,
    "independent_C_kg_kg": 2e-5,
}
HISTORY_COLUMNS = [
    "time_s",
    "radius_m",
    "T_center_C",
    "T_surface_C",
    "C_center",
    "C_surface",
    "C_max",
    "argmax_x",
    "C_mean",
    "heat_storage_J",
    "heat_boundary_J",
    "water_boundary_kg_kg",
    "water_defect_kg_kg",
    "robin_T_W_m2",
    "robin_C_kg_kg_m_s",
    "axis_T_C_m",
    "axis_C_kg_kg_m",
    "T_ambient_C",
    "C_ambient",
]


def sha256(path: Path) -> str:
    """Hash a file by bounded reads, including large numerical chunks."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def json_write(path: Path, value: object) -> None:
    """Atomically write finite, human-readable JSON in a derived directory."""
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".pending")
    temporary.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )
    temporary.replace(path)


def json_read(path: Path) -> dict:
    """Read a UTF-8 JSON document."""
    return json.loads(path.read_text(encoding="utf-8"))


def dependencies() -> dict:
    """Record actual numerical and workbook dependency versions."""
    return {
        "python": platform.python_version(),
        "numpy": np.__version__,
        "scipy": scipy.__version__,
        "openpyxl": openpyxl.__version__,
        "blas": sorted(
            [
                {
                    "library": item["prefix"],
                    "version": item["version"],
                    "threads": item["num_threads"],
                }
                for item in threadpool_info()
            ],
            key=lambda item: (item["library"], item["version"]),
        ),
    }


def input_hashes() -> dict:
    """Hash the immutable problem, observations, and original output template."""
    return {
        str(p.relative_to(ROOT)).replace("\\", "/"): sha256(p)
        for p in [
            ROOT / "problem.pdf",
            ROOT / "data/附件1.xlsx",
            ROOT / "data/附件2.xlsx",
            ROOT / "data/附件3/result4.xlsx",
        ]
    }


def protected_hashes() -> dict:
    """Record existing questions and shared modules without reading PNG files."""
    result = input_hashes()
    pending = [
        ROOT / "src/common",
        *(ROOT / f"src/problem-{i}" for i in (1, 2, 3)),
        *(ROOT / f"output/problem-{i}" for i in (1, 2, 3)),
    ]
    while pending:
        base = pending.pop()
        for path in sorted(base.iterdir()):
            if path.name in {
                "__pycache__",
                "node_modules",
            } or not path.resolve().is_relative_to(ROOT):
                continue
            if path.is_dir():
                pending.append(path)
            elif path.suffix.lower() != ".png":
                result[str(path.relative_to(ROOT)).replace("\\", "/")] = sha256(path)
    return result


def audit(save: bool = True) -> tuple[np.ndarray, np.ndarray, dict]:
    """Read every input row; preserve and report anomalies without correcting them."""
    records, sheets, anomalies = [], [], []
    specs = [
        ("附件1.xlsx", ["时间", "温度", "水分浓度"], 241, 60.0),
        ("附件2.xlsx", ["时间", "半径"], 145, 1800.0),
    ]
    for name, headers, expected_count, interval in specs:
        book = openpyxl.load_workbook(
            ROOT / "data" / name, read_only=True, data_only=True
        )
        if book.sheetnames != ["Sheet1"]:
            anomalies.append({"file": name, "sheets": book.sheetnames})
        sheet = book["Sheet1"]
        rows = list(sheet.iter_rows(values_only=True))
        book.close()
        if list(rows[0]) != headers:
            anomalies.append({"file": name, "row": 1, "values": list(rows[0])})
        valid = True
        for row_number, row in enumerate(rows[1:], 2):
            if len(row) != len(headers) or any(
                isinstance(v, bool)
                or not isinstance(v, (int, float))
                or not math.isfinite(v)
                for v in row
            ):
                anomalies.append(
                    {"file": name, "row": row_number, "values": [str(v) for v in row]}
                )
                valid = False
        if not valid:
            continue
        array = np.asarray(rows[1:], dtype=np.float64)
        differences = np.diff(array[:, 0])
        if len(array) != expected_count or array[0, 0] != 0:
            anomalies.append(
                {"file": name, "count": len(array), "first_time": float(array[0, 0])}
            )
        bad_times = np.flatnonzero(differences != interval)
        if len(bad_times):
            anomalies.append(
                {"file": name, "bad_time_excel_rows": (bad_times + 3).tolist()}
            )
        invalid = (
            array[:, 1] <= 0
            if name == "附件2.xlsx"
            else ((array[:, 1] + 273.15 <= 0) | (array[:, 2] <= 0))
        )
        if invalid.any():
            anomalies.append(
                {
                    "file": name,
                    "invalid_state_excel_rows": (np.flatnonzero(invalid) + 2).tolist(),
                }
            )
        if name == "附件2.xlsx" and (
            array[0, 1] != 2 or np.any(np.diff(array[:, 1]) > 0)
        ):
            anomalies.append(
                {"file": name, "reason": "Unexpected radius trend or initial radius"}
            )
        records.append(array)
        sheets.append(
            {
                "file": name,
                "sheet": "Sheet1",
                "headers": headers,
                "count": len(array),
                "time_range_s": array[[0, -1], 0].tolist(),
                "interval_s": interval,
                "minima": array.min(axis=0).tolist(),
                "maxima": array.max(axis=0).tolist(),
                "duplicate_time_count": int(len(array) - len(np.unique(array[:, 0]))),
                "repeated_radius_values_are_valid": name == "附件2.xlsx",
            }
        )
    template = openpyxl.load_workbook(
        ROOT / "data/附件3/result4.xlsx", read_only=True, data_only=True
    )
    template_rows = list(template["Sheet1"].values)
    if (
        template.sheetnames != ["Sheet1"]
        or template_rows[0][0] != "时间\\到药材中心的距离"
        or template_rows[0][-1] != "药材表面"
    ):
        anomalies.append(
            {"file": "result4.xlsx", "reason": "Unexpected template structure"}
        )
    template.close()
    report = {
        "status": "FAIL" if anomalies else "PASS",
        "sha256": input_hashes(),
        "sheets": sheets,
        "template_rows": template_rows,
        "anomalies": anomalies,
        "processing": "只读原始记录；半径 cm 转 m 存入独立对象；不删改、填补、平滑或拟合。",
    }
    if save:
        json_write(OUT / "review/input_audit.json", report)
    if anomalies:
        raise ValueError(
            "Input anomalies require user decision: "
            + json.dumps(anomalies, ensure_ascii=False)
        )
    env, radius_cm = records
    radius = radius_cm.copy()
    radius[:, 1] *= 0.01
    if save:
        (OUT / "data").mkdir(parents=True, exist_ok=True)
        np.savetxt(
            OUT / "data/ambient_observed.csv",
            env,
            delimiter=",",
            header="time_s,temperature_C,moisture_kg_kg",
            comments="",
        )
        np.savetxt(
            OUT / "data/radius_observed_SI.csv",
            radius,
            delimiter=",",
            header="time_s,radius_m",
            comments="",
        )
    return env, radius, report


@dataclass(frozen=True)
class Config:
    """Serializable scenario and numerical parameters, all included in cache identity."""

    n: int = 1600
    factor: float = 1.0
    appendix: int = 4
    shrink: bool = True
    tail: str = "last"
    horizon: float = 259200.0
    event_tol: float = 0.001
    method: str = "sdirk2"
    seed: int | None = None
    tau_s: float | None = None
    sigma_scale: float = 1.0


def plateau_fluctuation_stats(env: np.ndarray, window_start: float = 12600.0) -> dict:
    """Estimate plateau mean, residual std and lag-1 autocorrelation per channel."""
    window = env[env[:, 0] >= window_start]
    t = window[:, 0]
    stats = {}
    for j, key in [(1, "T"), (2, "C")]:
        y = window[:, j]
        slope, intercept = np.polyfit(t, y, 1)
        residual = y - (slope * t + intercept)
        sigma = float(residual.std(ddof=1))
        acf1 = (
            float(np.corrcoef(residual[:-1], residual[1:])[0, 1]) if sigma > 0 else 0.0
        )
        z_sample = (residual - residual.mean()) / sigma if sigma > 0 else residual * 0.0
        stats[key] = {
            "mean": float(y.mean()),
            "sigma": sigma,
            "acf1": acf1,
            "slope_per_s": float(slope),
            "kurtosis": float((z_sample**4).mean()),
            "max_abs_z": float(np.abs(z_sample).max()),
            "z_sample": [float(v) for v in z_sample],
        }
    return stats


class Boundary:
    """Piecewise-linear observed inputs with explicit post-observation scenarios."""

    def __init__(self, env: np.ndarray, radius: np.ndarray, config: Config):
        """Keep independent read-only observations and calculate the declared tail."""
        self.env_data, self.radius_data, self.config = env, radius, config
        self.dt_env = 60.0
        window = env[env[:, 0] >= env[-1, 0] - 1800]
        self.mean_window = {
            "start_s": float(window[0, 0]),
            "end_s": float(window[-1, 0]),
            "count": len(window),
            "mean": window[:, 1:].mean(axis=0).tolist(),
        }
        self.mean_tail = window[:, 1:].mean(axis=0)
        self.stats = None
        self.fluct = None
        if config.tail == "last":
            self.tail = env[-1, 1:].copy()
        elif config.tail == "mean30":
            self.tail = self.mean_tail.copy()
        elif config.tail == "fluct":
            self.tail = self.mean_tail.copy()
            self._build_fluctuation()
        else:
            raise ValueError("Tail mode must be last, mean30 or fluct")

    def _build_fluctuation(self) -> None:
        """Precompute a seeded AR(1) plateau path from the detrended observed window."""
        self.stats = plateau_fluctuation_stats(self.env_data)
        sigma = (
            np.array([self.stats["T"]["sigma"], self.stats["C"]["sigma"]])
            * self.config.sigma_scale
        )
        rng = np.random.default_rng(self.config.seed)
        grid = np.arange(
            self.env_data[-1, 0],
            max(self.config.horizon, self.env_data[-1, 0]) + 3600.0 + self.dt_env,
            self.dt_env,
        )
        z = np.empty((2, grid.size))
        phi = {}
        for j, key in enumerate(("T", "C")):
            if self.config.tau_s is None:
                p = float(np.clip(self.stats[key]["acf1"], 0.0, 0.95))
            else:
                p = float(np.exp(-self.dt_env / self.config.tau_s))
            phi[key] = p
            sample = np.asarray(self.stats[key]["z_sample"], dtype=float)
            if sample.size and sample.std() > 0:
                sample = (sample - sample.mean()) / sample.std()
                eps = sample[rng.integers(0, sample.size, grid.size)]
            else:
                eps = rng.standard_normal(grid.size)
            z[j] = np.clip(lfilter([np.sqrt(1.0 - p**2)], [1.0, -p], eps), -4.0, 4.0)
        self.fluct = {
            "grid": grid,
            "z": z,
            "sigma": sigma,
            "phi": phi,
            "dt_env_s": self.dt_env,
        }

    def ambient(self, t: float) -> np.ndarray:
        """Interpolate observations and apply the declared post-observation scenario."""
        if t < 0:
            raise ValueError("Negative observation time")
        if t <= self.env_data[-1, 0]:
            return np.array(
                [np.interp(t, self.env_data[:, 0], self.env_data[:, j]) for j in (1, 2)]
            )
        if self.fluct is not None:
            grid = self.fluct["grid"]
            if t >= grid[-1]:
                return self.mean_tail.copy()
            return np.array(
                [
                    np.interp(
                        t,
                        grid,
                        self.mean_tail[j] + self.fluct["sigma"][j] * self.fluct["z"][j],
                    )
                    for j in (0, 1)
                ]
            )
        return self.tail.copy()

    def radius(self, t: float) -> float:
        """Reject all extrapolation beyond the supplied radius time coverage."""
        if t < 0 or t > self.radius_data[-1, 0] + 1e-8:
            raise ValueError(f"Radius outside observed coverage: t={t}")
        if not self.config.shrink:
            return R0
        return float(np.interp(t, self.radius_data[:, 0], self.radius_data[:, 1]))


def thermal(C: np.ndarray, appendix: int = 4) -> tuple:
    """Return local positive heat capacity and conductivity from moisture alone."""
    cmin = float(C.min())
    if not math.isfinite(cmin) or cmin <= 0.0:
        raise FloatingPointError("Nonpositive or nonfinite moisture")
    with np.errstate(over="raise", under="raise", divide="raise", invalid="raise"):
        fraction = C / (1.0 + C)
        if appendix == 4:
            a = (760.0 + 90.0 * C) * (1850.0 + 2150.0 * fraction)
            k = 0.12 + 0.20 * fraction
        elif appendix == 3:
            a = (650.0 + 128.0 * C) * (1450.0 + 2736.0 * fraction)
            k = 0.21 + 0.38 * fraction
        else:
            raise ValueError("Appendix must be 3 or 4")
    if not (float(a.min()) > 0.0 and float(k.min()) > 0.0):
        raise FloatingPointError("Nonpositive or nonfinite heat property")
    return a, k


def diffusivity(T: np.ndarray, C: np.ndarray, appendix: int = 4) -> np.ndarray:
    """Return local positive SI moisture diffusivity from both fields."""
    tmin, tmax, cmin = float(T.min()), float(T.max()), float(C.min())
    if not (math.isfinite(tmin) and math.isfinite(tmax)):
        raise FloatingPointError("Nonfinite temperature")
    if tmin + 273.15 <= 0.0 or not math.isfinite(cmin) or cmin <= 0.0:
        raise FloatingPointError("Nonpositive C or absolute temperature")
    with np.errstate(over="raise", under="raise", divide="raise", invalid="raise"):
        if appendix == 4:
            D = 4.2e-4 * np.exp(-0.30 / C) * np.exp(-3850.0 / (T + 273.15))
        elif appendix == 3:
            D = 2.4e-3 * np.exp(-0.45 / C) * np.exp(-3850.0 / (T + 273.15))
        else:
            raise ValueError("Appendix must be 3 or 4")
    if not (float(D.min()) > 0.0):
        raise FloatingPointError("Nonpositive or nonfinite diffusivity")
    return D


def properties(T: np.ndarray, C: np.ndarray, appendix: int = 4) -> tuple:
    """Return local positive heat capacity, conductivity, and SI diffusivity."""
    a, k = thermal(C, appendix)
    return a, k, diffusivity(T, C, appendix)


class Solver:
    """Coupled SDIRK2 with moving geometry and shared harmonic face fluxes."""

    def __init__(
        self, config: Config, boundary: Boundary, h: float = H, hm: float = HM
    ):
        """Construct uniform material nodes and exact cylindrical control weights."""
        if config.n < 4 or config.factor <= 0 or config.event_tol <= 0:
            raise ValueError("Invalid numerical configuration")
        self.config, self.boundary, self.h, self.hm = config, boundary, h, hm
        self.x = np.linspace(0.0, 1.0, config.n + 1)
        self.dx = 1.0 / config.n
        edges = np.r_[0.0, (self.x[:-1] + self.x[1:]) / 2, 1.0]
        self.weights = np.diff(edges**2)
        self.face_factor = 2.0 * edges[1:-1] / self.dx
        self.ones = np.ones(config.n + 1)
        self._bands = np.zeros((2, config.n + 1))
        self._rhs = np.empty(config.n + 1)
        self._diff = np.empty(config.n)
        self._flux = np.empty(config.n)
        self._net_t = np.empty(config.n + 1)
        self._net_c = np.empty(config.n + 1)
        self.last_rate = None
        self.max_iter = 0
        self.max_change = 0.0
        self.max_residual = 0.0

    def conductance(self, coefficient: np.ndarray) -> np.ndarray:
        """Combine two half-cell resistances; explicit radius factors cancel."""
        with np.errstate(over="raise", divide="raise", invalid="raise"):
            face = 2.0 / (1.0 / coefficient[:-1] + 1.0 / coefficient[1:])
            face *= self.face_factor
        if not (float(face.min()) > 0.0):
            raise FloatingPointError("Invalid harmonic face property")
        return face

    def net(
        self,
        u: np.ndarray,
        G: np.ndarray,
        transfer: float,
        ambient: float,
        radius: float,
        out: np.ndarray | None = None,
    ) -> np.ndarray:
        """Assemble each internal face once, and add the actual surface exchange."""
        result = self._net_t if out is None else out
        np.subtract(u[1:], u[:-1], out=self._diff)
        np.multiply(G, self._diff, out=self._flux)
        result[:] = 0.0
        result[:-1] += self._flux
        result[1:] -= self._flux
        result[-1] += 2.0 * radius * transfer * (ambient - u[-1])
        return result

    def linear(
        self,
        base: np.ndarray,
        capacity: np.ndarray,
        G: np.ndarray,
        transfer: float,
        ambient: float,
        radius: float,
        dt: float,
    ) -> np.ndarray:
        """Solve an increment system to reduce temperature subtraction cancellation."""
        ab = self._bands
        diagonal = ab[1]
        np.multiply(self.weights, capacity, out=diagonal)
        diagonal *= radius * radius
        diagonal /= dt
        diagonal[:-1] += G
        diagonal[1:] += G
        diagonal[-1] += 2.0 * radius * transfer
        ab[0, 0] = 0.0
        ab[0, 1:] = -G
        self.net(base, G, transfer, ambient, radius, out=self._rhs)
        increment = solveh_banded(
            ab,
            self._rhs,
            lower=False,
            overwrite_b=True,
            check_finite=False,
        )
        return base + increment

    def implicit(
        self,
        base: np.ndarray,
        initial: np.ndarray,
        t: float,
        dt: float,
        max_iter: int = 60,
    ) -> tuple:
        """Require scaled iterate changes and original nonlinear equation defects."""
        T, C = initial.copy()
        ambient, radius = self.boundary.ambient(t), self.boundary.radius(t)
        volume = radius**2 * self.weights
        for iteration in range(1, max_iter + 1):
            a, k = thermal(C, self.config.appendix)
            T1 = self.linear(
                base[0], a, self.conductance(k), self.h, ambient[0], radius, dt
            )
            D = diffusivity(T1, C, self.config.appendix)
            C1 = self.linear(
                base[1], self.ones, self.conductance(D), self.hm, ambient[1], radius, dt
            )
            change = max(
                float(np.max(np.abs(T1 - T))) / 50.0, float(np.max(np.abs(C1 - C))) / C0
            )
            T, C = T1, C1
            if change > ITER_TOL:
                continue
            a, k = thermal(C, self.config.appendix)
            D = diffusivity(T, C, self.config.appendix)
            GT, GC = self.conductance(k), self.conductance(D)
            defects = []
            for u, old, capacity, G, transfer, amb, scale in (
                (T, base[0], a, GT, self.h, ambient[0], 50.0),
                (C, base[1], self.ones, GC, self.hm, ambient[1], C0),
            ):
                mass = capacity * volume / dt
                diagonal = mass.copy()
                diagonal[:-1] += G
                diagonal[1:] += G
                diagonal[-1] += 2 * radius * transfer
                residual = mass * (u - old) - self.net(
                    u, G, transfer, amb, radius, out=self._rhs
                )
                defects.append(float(np.max(np.abs(residual) / (diagonal * scale))))
            defect = max(defects)
            if defect <= RES_TOL:
                self.max_iter = max(self.max_iter, iteration)
                self.max_change = max(self.max_change, change)
                self.max_residual = max(self.max_residual, defect)
                state = np.stack([T, C])
                rate = (state - base) / dt
                balance = np.array(
                    [
                        math.pi * LENGTH * np.dot(volume * a, rate[0]),
                        2 * math.pi * LENGTH * radius * self.h * (ambient[0] - T[-1]),
                        2 * self.hm / radius * (ambient[1] - C[-1]),
                    ]
                )
                return state, rate, balance
        raise RuntimeError(f"Unconverged implicit stage: t={t:.12g}, dt={dt:.6g}")

    def step(self, old: np.ndarray, t: float, dt: float) -> tuple:
        """Integrate one step and independently accumulate actual stage boundary fluxes."""
        if self.config.method == "be":
            state, _, balance = self.implicit(old, old, t + dt, dt)
            return state, dt * balance
        guess = old
        if self.last_rate is not None:
            guess = old + GAMMA * dt * self.last_rate
            if guess[1].min() <= 0.0 or not math.isfinite(float(guess[0].min())):
                guess = old
        stage1, rate1, balance1 = self.implicit(old, guess, t + GAMMA * dt, GAMMA * dt)
        base2 = old + dt * (1.0 - GAMMA) * rate1
        stage2, _, balance2 = self.implicit(base2, stage1, t + dt, GAMMA * dt)
        self.last_rate = (stage2 - old) / dt
        return stage2, dt * ((1.0 - GAMMA) * balance1 + GAMMA * balance2)

    def rhs(self, t: float, state: np.ndarray) -> np.ndarray:
        """Evaluate the semidiscrete ODE for an independent SciPy BDF time integrator."""
        T, C = state
        a, k, D = properties(T, C, self.config.appendix)
        radius, ambient = self.boundary.radius(t), self.boundary.ambient(t)
        volume = radius**2 * self.weights
        return np.stack(
            [
                self.net(
                    T,
                    self.conductance(k),
                    self.h,
                    ambient[0],
                    radius,
                    out=self._net_t,
                )
                / (a * volume),
                self.net(
                    C,
                    self.conductance(D),
                    self.hm,
                    ambient[1],
                    radius,
                    out=self._net_c,
                )
                / volume,
            ]
        )

    def boundaries(self, state: np.ndarray, t: float) -> tuple:
        """Reconstruct physical gradients independently by second-order one-sided differences."""
        radius = self.boundary.radius(t)
        dr = radius * self.dx
        surface = (3 * state[:, -1] - 4 * state[:, -2] + state[:, -3]) / (2 * dr)
        center = (-3 * state[:, 0] + 4 * state[:, 1] - state[:, 2]) / (2 * dr)
        _, k, D = properties(*state, self.config.appendix)
        robin = -np.array([k[-1], D[-1]]) * surface - np.array([self.h, self.hm]) * (
            state[:, -1] - self.boundary.ambient(t)
        )
        return robin, center


def local_integrate(
    solver: Solver, seed: np.ndarray, t0: float, target: float, subdivisions: int = 1
) -> tuple:
    """Reintegrate from a saved full state, never infer an endpoint field by interpolation."""
    state, total = seed.copy(), np.zeros(3)
    for j in range(subdivisions):
        left = t0 + (target - t0) * j / subdivisions
        right = t0 + (target - t0) * (j + 1) / subdivisions
        state, balance = solver.step(state, left, right - left)
        total += balance
    return state, total


def refine_event(
    solver: Solver,
    seed: np.ndarray,
    t0: float,
    width: float,
    tolerance: float,
    subdivisions: int = 1,
) -> tuple:
    """Bisect the all-node event with full integrations and retain a strictly dry upper state."""
    low, high = 0.0, width
    low_state = seed.copy()
    high_state, high_balance = local_integrate(
        solver, seed, t0, t0 + high, subdivisions
    )
    if seed[1].max() < THRESHOLD or high_state[1].max() >= THRESHOLD:
        raise RuntimeError("Invalid full-domain event bracket")
    trace = [
        [t0, float(seed[1].max() - THRESHOLD)],
        [t0 + high, float(high_state[1].max() - THRESHOLD)],
    ]
    while high - low > tolerance:
        middle = (low + high) / 2
        trial, balance = local_integrate(solver, seed, t0, t0 + middle, subdivisions)
        gap = float(trial[1].max() - THRESHOLD)
        trace.append([t0 + middle, gap])
        if gap < 0:
            high, high_state, high_balance = middle, trial, balance
        else:
            low, low_state = middle, trial
    # Keeping the accepted upper endpoint avoids crossing a later input/output knot.
    event = {
        "lower_s": t0 + low,
        "upper_s": t0 + high,
        "critical_s": t0 + (low + high) / 2,
        "end_s": t0 + high,
        "width_s": high - low,
        "g_lower": float(low_state[1].max() - THRESHOLD),
        "g_upper": float(high_state[1].max() - THRESHOLD),
        "margin_kg_kg": float(THRESHOLD - high_state[1].max()),
        "argmax_x": float(solver.x[np.argmax(high_state[1])]),
        "trace": trace,
    }
    return event, high_state, high_balance


def case_identity(config: Config, audit_report: dict) -> dict:
    """Fingerprint the complete solver file, parameters, sources, dependencies and schema."""
    return {
        "schema": SCHEMA,
        "config": asdict(config),
        "inputs": audit_report["sha256"],
        "solve_code_sha256": sha256(SOURCE),
        "dependencies": dependencies(),
    }


def case_path(name: str, identity: dict) -> Path:
    """Derive an immutable case directory from a descriptive name and exact fingerprint."""
    if not name or any(c not in "abcdefghijklmnopqrstuvwxyz0123456789_-" for c in name):
        raise ValueError("Case name must use lowercase ASCII letters, digits, _ or -")
    fingerprint = hashlib.sha256(
        json.dumps(identity, sort_keys=True).encode()
    ).hexdigest()[:16]
    return OUT / "cases" / f"{name}-{fingerprint}"


def load_case(path: Path, identity: dict | None = None) -> dict:
    """Verify identity and every saved numerical file before any reuse."""
    meta = json_read(path / "metadata.json")
    if meta["status"] != "COMPLETE" or (
        identity is not None and meta["identity"] != identity
    ):
        raise ValueError(f"Case identity or completion mismatch: {path}")
    if meta["identity"]["solve_code_sha256"] != sha256(SOURCE):
        raise ValueError(f"Solver source changed: {path}")
    if meta["identity"]["inputs"] != input_hashes():
        raise ValueError(f"Input hashes changed: {path}")
    for filename, digest in meta["files"].items():
        if sha256(path / filename) != digest:
            raise ValueError(f"Saved data checksum mismatch: {path / filename}")
    with np.load(path / "summary.npz", allow_pickle=False) as data:
        if data["x"].shape != (meta["identity"]["config"]["n"] + 1,):
            raise ValueError("Saved data schema mismatch")
        if data["history"].shape[1] != len(HISTORY_COLUMNS):
            raise ValueError("Saved history schema mismatch")
    return meta


def iter_fields(path: Path, meta: dict):
    """Stream full unrounded states in chronological order with bounded memory."""
    n = meta["identity"]["config"]["n"]
    previous = -1.0
    for chunk in meta["chunks"]:
        with np.load(path / chunk, allow_pickle=False) as data:
            times, fields = data["times"], data["states"]
            if fields.shape != (len(times), 2, n + 1) or not np.isfinite(fields).all():
                raise ValueError("Invalid full-state chunk structure")
            for t, state in zip(times, fields):
                if t <= previous:
                    raise ValueError("Nonascending saved times")
                previous = float(t)
                yield float(t), state


def simulate(
    config: Config,
    name: str,
    env: np.ndarray,
    radius: np.ndarray,
    audit_report: dict,
    resume: bool = True,
) -> Path:
    """Run one case to strict dryness or the observed radius horizon, preserving all evidence."""
    identity = case_identity(config, audit_report)
    path = case_path(name, identity)
    if resume:
        candidates = [path, *sorted(path.parent.glob(path.name + "-retry-*"))]
        for candidate in candidates:
            if (candidate / "metadata.json").exists():
                load_case(candidate, identity)
                print(f"REUSE {candidate.name}", flush=True)
                return candidate
    if path.exists():
        # Incomplete runs remain inspectable; a retry never erases a previous attempt.
        path = path.with_name(path.name + f"-retry-{time.time_ns()}")
    path.mkdir(parents=True)
    if config.horizon > radius[-1, 0] or config.horizon <= 0:
        raise ValueError("Case horizon must stay within the observed radius coverage")
    boundary = Boundary(env, radius, config)
    solver = Solver(config, boundary)
    state = np.stack([np.full(config.n + 1, 28.0), np.full(config.n + 1, C0)])
    started, last_print = time.perf_counter(), time.perf_counter()
    corner_times = np.r_[START_TIMES, env[-1, 0] + START_TIMES]
    output_times = np.unique(
        np.r_[
            0.0,
            corner_times[corner_times < config.horizon],
            np.arange(60.0, config.horizon + 1e-8, 60.0),
            config.horizon,
        ]
    )
    knots = np.unique(
        np.r_[
            output_times,
            env[:, 0][env[:, 0] <= config.horizon],
            radius[:, 0][radius[:, 0] <= config.horizon],
        ]
    )
    history, chunks, buffer_times, buffer_states = [], [], [], []
    checks_times, checks_states = [], []
    total = np.zeros(3)
    accepted = rejected = 0
    t, step_min, step_max = 0.0, math.inf, 0.0
    output_index, knot_index = 1, 1
    event = None
    event_seed = state.copy()
    event_t0, event_width = 0.0, 0.0
    ranges = np.array(
        [[28.0, 28.0], [C0, C0], [math.inf, 0.0], [math.inf, 0.0], [math.inf, 0.0]]
    )
    radial_inversion, temporal_increase, above_center = 0.0, 0.0, 0.0
    max_mass_defect, max_heat_defect = 0.0, 0.0

    def flush_chunk() -> None:
        """Compress a bounded number of full double-precision states."""
        if not buffer_times:
            return
        filename = f"fields_{len(chunks):04d}.npz"
        np.savez_compressed(
            path / filename,
            times=np.asarray(buffer_times),
            states=np.asarray(buffer_states),
        )
        chunks.append(filename)
        buffer_times.clear()
        buffer_states.clear()

    def capture() -> None:
        """Save actual full states and separately computed integral and boundary diagnostics."""
        radius_now = boundary.radius(t)
        robin, axis = solver.boundaries(state, t)
        average = float(np.dot(solver.weights, state[1]))
        history.append(
            [
                t,
                radius_now,
                *state[0, [0, -1]],
                *state[1, [0, -1]],
                float(state[1].max()),
                float(solver.x[np.argmax(state[1])]),
                average,
                *total,
                average - C0 - total[2],
                *robin,
                *axis,
                *boundary.ambient(t),
            ]
        )
        buffer_times.append(t)
        buffer_states.append(state.copy())
        is_check = (
            t == 0
            or t in corner_times
            or t == 60
            or t % 21600 < 1e-8
            or event is not None
            or t == config.horizon
        )
        if is_check:
            checks_times.append(t)
            checks_states.append(state.copy())
        if len(buffer_times) >= 32:
            flush_chunk()

    capture()
    try:
        while t < config.horizon - 1e-8 and event is None:
            target = float(knots[knot_index])
            cap = (4.0 if t < env[-1, 0] else 60.0) * config.factor
            age = t if t < env[-1, 0] else t - env[-1, 0]
            dt = min(
                cap, max(0.0001, 0.025 * (age + 0.002)) * config.factor, target - t
            )
            if dt <= 0:
                raise RuntimeError("Nonpositive time step")
            for retry in range(21):
                try:
                    new, increment = solver.step(state, t, dt)
                    properties(*new, config.appendix)
                    break
                except (RuntimeError, FloatingPointError, np.linalg.LinAlgError):
                    rejected += 1
                    dt *= 0.5
                    if retry == 20 or dt < 1e-10:
                        raise
            if new[1].max() < THRESHOLD:
                event_seed, event_t0, event_width = state.copy(), t, dt
                event, new, increment = refine_event(
                    solver, state, t, dt, config.event_tol
                )
                dt = event["end_s"] - t
            total += increment
            radial_inversion = max(radial_inversion, float(np.max(np.diff(new[1]))))
            temporal_increase = max(temporal_increase, float(np.max(new[1] - state[1])))
            above_center = max(above_center, float(new[1].max() - new[1, 0]))
            a, k, D = properties(*new, config.appendix)
            for j, values in enumerate((new[0], new[1], D, a, k)):
                ranges[j, 0] = min(ranges[j, 0], float(values.min()))
                ranges[j, 1] = max(ranges[j, 1], float(values.max()))
            max_mass_defect = max(
                max_mass_defect,
                abs(float(np.dot(solver.weights, new[1] - C0)) - total[2]),
            )
            max_heat_defect = max(max_heat_defect, abs(total[0] - total[1]))
            state, t = new, t + dt
            accepted += 1
            step_min, step_max = min(step_min, dt), max(step_max, dt)
            if abs(t - target) < 1e-8:
                t = target
                knot_index += 1
            if event is not None or abs(t - output_times[output_index]) < 1e-8:
                capture()
                output_index += 1
            if time.perf_counter() - last_print > 25:
                print(
                    f"RUN {name} N={config.n} f={config.factor:g} t={t/3600:.3f}h "
                    f"maxC={state[1].max():.7f} steps={accepted} "
                    f"elapsed={time.perf_counter()-started:.1f}s",
                    flush=True,
                )
                last_print = time.perf_counter()
        flush_chunk()
        payload = {
            "x": solver.x,
            "weights": solver.weights,
            "history": np.asarray(history),
            "check_times": np.asarray(checks_times),
            "check_states": np.asarray(checks_states),
            "final_state": state,
            "event_seed": event_seed,
            "event_t0": event_t0,
            "event_width": event_width,
        }
        if boundary.fluct is not None:
            payload["ambient_grid"] = boundary.fluct["grid"]
            payload["ambient_path"] = (
                boundary.mean_tail[:, None]
                + boundary.fluct["sigma"][:, None] * boundary.fluct["z"]
            )
        np.savez_compressed(path / "summary.npz", **payload)
        mass_relative = max_mass_defect / C0
        heat_relative = max_heat_defect / max(abs(total[0]), abs(total[1]), 1.0)
        meta = {
            "status": "COMPLETE",
            "validation": "PENDING",
            "name": name,
            "identity": identity,
            "chunks": chunks,
            "history_columns": HISTORY_COLUMNS,
            "event": event,
            "end_s": t,
            "horizon_reached_without_dryness": event is None,
            "end_max_C": float(state[1].max()),
            "end_argmax_x": float(solver.x[np.argmax(state[1])]),
            "environment_tail": boundary.tail.tolist(),
            "environment_mode": config.tail,
            "environment_seed": config.seed,
            "environment_tau_s": config.tau_s,
            "environment_sigma_scale": config.sigma_scale,
            "environment_ar1_phi": (
                boundary.fluct["phi"] if boundary.fluct is not None else None
            ),
            "fluctuation_stats": boundary.stats,
            "mean30_window": boundary.mean_window,
            "steps": accepted,
            "rejected_steps": rejected,
            "step_range_s": [step_min, step_max],
            "max_iterations": solver.max_iter,
            "max_scaled_change": solver.max_change,
            "max_scaled_residual": solver.max_residual,
            "property_ranges_T_C_D_a_k": ranges.tolist(),
            "mass_max_abs_defect_kg_kg": max_mass_defect,
            "mass_relative_defect": mass_relative,
            "heat_max_abs_defect_J": max_heat_defect,
            "heat_relative_defect": heat_relative,
            "radial_inversion_kg_kg": radial_inversion,
            "temporal_increase_kg_kg": temporal_increase,
            "max_C_above_center": above_center,
            "runtime_s": time.perf_counter() - started,
            "files": {p.name: sha256(p) for p in sorted(path.glob("*.npz"))},
        }
        if (
            input_hashes() != identity["inputs"]
            or sha256(SOURCE) != identity["solve_code_sha256"]
        ):
            raise RuntimeError("Inputs or solver source changed during integration")
        json_write(path / "metadata.json", meta)
        print(
            f"DONE {name} N={config.n} t={t/3600:.8f}h Cmax={state[1].max():.15g} "
            f"mass={mass_relative:.2e} heat={heat_relative:.2e} "
            f"runtime={meta['runtime_s']:.1f}s",
            flush=True,
        )
        return path
    except (Exception, KeyboardInterrupt) as exc:
        flush_chunk()
        np.savez_compressed(
            path / "failure_state.npz", state=state, t=t, cumulative=total
        )
        json_write(
            path / "failure.json",
            {
                "status": "FAIL",
                "identity": identity,
                "time_s": t,
                "error": repr(exc),
                "steps": accepted,
                "rejected": rejected,
            },
        )
        raise


def compare_cases(coarse: Path, fine: Path) -> dict:
    """Compare every common saved full field, surface layer and true-distance output point."""
    ma, mb = load_case(coarse), load_case(fine)
    ca, cb = ma["identity"]["config"], mb["identity"]["config"]
    for key in ("appendix", "shrink", "tail", "horizon", "seed", "tau_s", "sigma_scale"):
        if ca[key] != cb[key]:
            raise ValueError("Convergence cases have different physical inputs")
    xa, xb = np.linspace(0, 1, ca["n"] + 1), np.linspace(0, 1, cb["n"] + 1)
    env, radius, _ = audit(save=False)
    boundary = Boundary(env, radius, Config(**cb))
    ita, itb = iter_fields(coarse, ma), iter_fields(fine, mb)
    a, b = next(ita, None), next(itb, None)
    metrics = {
        key: {"T": 0.0, "C": 0.0, "worst_T_s": 0.0, "worst_C_s": 0.0, "worst_C_x": 0.0}
        for key in ("full", "startup", "delivery", "surface_layer", "late")
    }
    count, last_common = 0, 0.0
    while a is not None and b is not None:
        ta, ua = a
        tb, ub = b
        if abs(ta - tb) > 1e-8:
            if ta < tb:
                a = next(ita, None)
            else:
                b = next(itb, None)
            continue
        count += 1
        last_common = ta
        # Evaluate the coarse piecewise-linear field at every fine node, including midpoints.
        interpolation = np.stack([np.interp(xb, xa, u) for u in ua])
        difference = np.abs(interpolation - ub)
        selections = {"surface_layer": xb >= 0.98}
        if ta <= 60:
            selections["startup"] = np.ones(len(xb), dtype=bool)
        else:
            selections["full"] = np.ones(len(xb), dtype=bool)
        if ta >= min(ma["end_s"], mb["end_s"]) - 3600:
            selections["late"] = np.ones(len(xb), dtype=bool)
        for key, mask in selections.items():
            for j, label in enumerate(("T", "C")):
                error = float(difference[j, mask].max())
                if error > metrics[key][label]:
                    metrics[key][label] = error
                    metrics[key][f"worst_{label}_s"] = ta
                    if label == "C":
                        metrics[key]["worst_C_x"] = float(
                            xb[mask][np.argmax(difference[j, mask])]
                        )
        if ta >= 60 and abs(ta / 60 - round(ta / 60)) < 1e-8:
            R = boundary.radius(ta)
            regular = np.arange(0, R0, 0.001)
            positions = np.r_[regular[regular < R - 1e-12] / R, 1.0]
            for j, label in enumerate(("T", "C")):
                error = float(
                    np.max(
                        np.abs(
                            np.interp(positions, xa, ua[j])
                            - np.interp(positions, xb, ub[j])
                        )
                    )
                )
                if error > metrics["delivery"][label]:
                    metrics["delivery"][label] = error
                    metrics["delivery"][f"worst_{label}_s"] = ta
        a, b = next(ita, None), next(itb, None)
    event_delta = None
    if ma["event"] is not None and mb["event"] is not None:
        event_delta = mb["event"]["critical_s"] - ma["event"]["critical_s"]
    field_T_target = (
        TARGETS["fluct_field_T_C"] if ca["tail"] == "fluct" else TARGETS["field_T_C"]
    )
    passed = (
        metrics["full"]["C"] < TARGETS["field_C_kg_kg"]
        and metrics["full"]["T"] < field_T_target
        and (
            abs(event_delta) < TARGETS["event_change_s"]
            if event_delta is not None
            else ma["event"] is None and mb["event"] is None
        )
        and count > 2
    )
    return {
        "status": "PASS" if passed else "FAIL",
        "coarse": str(coarse.relative_to(OUT)),
        "fine": str(fine.relative_to(OUT)),
        "N": [ca["n"], cb["n"]],
        "factors": [ca["factor"], cb["factor"]],
        "matched_full_states": count,
        "last_common_s": last_common,
        "metrics": metrics,
        "field_T_target_C": field_T_target,
        "event_delta_s": event_delta,
    }


def case_checks(path: Path) -> dict:
    """Audit full-run balances, independent boundaries, local properties and endpoint state."""
    meta = load_case(path)
    with np.load(path / "summary.npz", allow_pickle=False) as data:
        history = data["history"]
        final = data["final_state"]
    regular = np.abs(history[:, 0] / 60 - np.rint(history[:, 0] / 60)) < 1e-8
    delivery = history[
        (history[:, 0] >= 60) & (regular | (history[:, 0] == meta["end_s"]))
    ]
    maxima = np.max(np.abs(delivery[:, 13:17]), axis=0)
    checks = {
        "water_balance": meta["mass_relative_defect"] <= TARGETS["balance_relative"],
        "heat_balance": meta["heat_relative_defect"] <= TARGETS["balance_relative"],
        "nonlinear_change": meta["max_scaled_change"] <= ITER_TOL,
        "nonlinear_residual": meta["max_scaled_residual"] <= RES_TOL,
        "robin_heat": maxima[0] <= TARGETS["robin_heat_W_m2"],
        "robin_moisture": maxima[1] <= TARGETS["robin_moisture_kg_kg_m_s"],
        "axis_temperature": maxima[2] <= TARGETS["axis_T_C_m"],
        "axis_moisture": maxima[3] <= TARGETS["axis_C_kg_kg_m"],
        "finite_positive_state": bool(
            np.isfinite(final).all() and (final[1] > 0).all()
        ),
        "endpoint": (
            bool(np.all(final[1] < THRESHOLD))
            if meta["event"]
            else meta["end_s"] == meta["identity"]["config"]["horizon"]
        ),
    }
    return {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": {k: "PASS" if v else "FAIL" for k, v in checks.items()},
        "boundary_abs_max": maxima.tolist(),
        "boundary_units": ["W/m2", "(kg/kg)m/s", "C/m", "(kg/kg)/m"],
        "startup_boundaries": history[
            (history[:, 0] < 60) | ((history[:, 0] > 14400) & (history[:, 0] < 14460))
        ][:, [0, 13, 14, 15, 16]].tolist(),
        "mass_relative_defect": meta["mass_relative_defect"],
        "heat_relative_defect": meta["heat_relative_defect"],
        "radial_inversion_kg_kg": meta["radial_inversion_kg_kg"],
        "temporal_increase_kg_kg": meta["temporal_increase_kg_kg"],
        "all_domain_max_above_center": meta["max_C_above_center"],
    }


def event_checks(path: Path, env: np.ndarray, radius: np.ndarray) -> dict:
    """Separate event bisection, local step error and concentration-to-time amplification."""
    meta = load_case(path)
    if meta["event"] is None:
        return {
            "status": "SKIPPED",
            "reason": "No threshold crossing within supplied radius horizon",
        }
    config = Config(**meta["identity"]["config"])
    solver = Solver(config, Boundary(env, radius, config))
    with np.load(path / "summary.npz", allow_pickle=False) as data:
        seed, t0, width = (
            data["event_seed"],
            float(data["event_t0"]),
            float(data["event_width"]),
        )
    tight, _, _ = refine_event(solver, seed, t0, width, 1e-5)
    half, _, _ = refine_event(solver, seed, t0, width, 1e-5, 2)
    quarter, _, _ = refine_event(solver, seed, t0, width, 1e-5, 4)
    ending, _ = local_integrate(solver, seed, t0, meta["event"]["end_s"])
    slope = float(solver.rhs(meta["event"]["end_s"], ending)[1, np.argmax(ending[1])])
    root_delta = tight["critical_s"] - meta["event"]["critical_s"]
    half_delta = half["critical_s"] - tight["critical_s"]
    quarter_delta = quarter["critical_s"] - half["critical_s"]
    passed = (
        abs(root_delta) <= config.event_tol
        and max(abs(half_delta), abs(quarter_delta)) <= TARGETS["event_local_change_s"]
        and np.all(ending[1] < THRESHOLD)
        and slope < 0
    )
    return {
        "status": "PASS" if passed else "FAIL",
        "tight": tight,
        "root_delta_s": root_delta,
        "two_substep_delta_s": half_delta,
        "four_substep_delta_s": quarter_delta,
        "end_slope_kg_kg_s": slope,
        "time_amplification_of_5e_5_s": TARGETS["field_C_kg_kg"] / abs(slope),
    }


def self_tests(env: np.ndarray, radius: np.ndarray) -> dict:
    """Execute physical invariants, rejection probes, source checks and independent BDF."""
    checks = {}
    config = Config(n=80, factor=0.125, horizon=60)
    solver = Solver(config, Boundary(env, radius, config))
    initial = np.stack([np.full(81, 28.0), np.full(81, C0)])
    a, k, D = properties(np.array([28.0, 50.0]), np.array([2.55, 0.15]))
    expected_D = np.array(
        [
            4.2e-4 * math.exp(-0.30 / c) * math.exp(-3850 / (t + 273.15))
            for t, c in ((28.0, 2.55), (50.0, 0.15))
        ]
    )
    checks["appendix4_properties_and_kelvin"] = bool(
        np.allclose(D, expected_D, rtol=1e-14)
        and abs(a[0] - (760 + 90 * 2.55) * (1850 + 2150 * 2.55 / 3.55)) < 1e-8
        and abs(k[0] - (0.12 + 0.20 * 2.55 / 3.55)) < 1e-14
    )
    for label, T, C in [
        ("zero_C_rejected", 28.0, 0.0),
        ("kelvin_rejected", -274.0, 1.0),
        ("nonfinite_rejected", math.nan, 1.0),
        ("underflow_rejected", 28.0, 1e-12),
    ]:
        try:
            properties(np.array([T]), np.array([C]))
            checks[label] = False
        except FloatingPointError:
            checks[label] = True
    closed = Solver(config, Boundary(env, radius, config), h=0.0, hm=0.0)
    uniform, _ = closed.step(initial, 1800.0, 30.0)
    checks["pure_shrink_without_moisture_exchange"] = bool(
        np.max(np.abs(uniform[1] - C0)) < 1e-12
    )
    equilibrium_env = env.copy()
    equilibrium_env[:, 1] = 28.0
    equilibrium = Solver(config, Boundary(equilibrium_env, radius, config), hm=0.0)
    equilibrated, _ = equilibrium.step(initial, 1800.0, 30.0)
    checks["uniform_thermal_equilibrium"] = bool(
        np.max(np.abs(equilibrated[0] - 28)) < 1e-12
    )
    G = solver.conductance(np.linspace(0.1, 0.3, 81))
    arbitrary = 1.0 + 0.1 * solver.x**2
    checks["shared_internal_flux_cancellation"] = (
        abs(float(solver.net(arbitrary, G, 0, 0, R0).sum())) < 1e-12
    )
    checks["control_weights_and_axis"] = bool(
        abs(solver.weights.sum() - 1) < 1e-14
        and abs(solver.weights[0] - (solver.dx / 2) ** 2) < 1e-14
    )
    fixed_config = Config(n=80, shrink=False, horizon=60)
    fixed = Solver(fixed_config, Boundary(env, radius, fixed_config))
    zero_change_radius = np.column_stack([radius[:, 0], np.full(len(radius), R0)])
    equivalent = Solver(config, Boundary(env, zero_change_radius, config))
    state_f, _ = fixed.step(initial, 0, 0.01)
    state_e, _ = equivalent.step(initial, 0, 0.01)
    checks["constant_radius_reduction"] = bool(np.array_equal(state_f, state_e))
    checks["appendix3_is_distinct"] = bool(
        np.max(np.abs(properties(*initial, 3)[2] - properties(*initial, 4)[2])) > 1e-10
    )
    try:
        solver.implicit(initial, initial, 0.1, 0.1, max_iter=0)
        checks["forced_nonlinear_failure"] = False
    except RuntimeError:
        checks["forced_nonlinear_failure"] = True
    try:
        solver.boundary.radius(radius[-1, 0] + 1)
        checks["radius_extrapolation_rejected"] = False
    except ValueError:
        checks["radius_extrapolation_rejected"] = True
    state, t = initial.copy(), 0.0
    while t < 60:
        dt = min(0.5, max(0.00001, 0.003 * (t + 0.002)), 60 - t)
        state, _ = solver.step(state, t, dt)
        t += dt
    size = 162
    offsets = [-82, -81, -80, -1, 0, 1, 80, 81, 82]
    sparsity = diags(
        [np.ones(size - abs(offset)) for offset in offsets],
        offsets,
        shape=(size, size),
        format="csr",
    )
    bdf = solve_ivp(
        lambda t, y: solver.rhs(t, y.reshape(2, 81)).ravel(),
        (0.0, 60.0),
        initial.ravel(),
        method="BDF",
        rtol=2e-10,
        atol=2e-12,
        jac_sparsity=sparsity,
        max_step=0.5,
    )
    difference = np.max(np.abs(state - bdf.y[:, -1].reshape(2, 81)), axis=1)
    checks["independent_BDF_time_integration"] = bool(
        bdf.success
        and difference[0] <= TARGETS["independent_T_C"]
        and difference[1] <= TARGETS["independent_C_kg_kg"]
    )
    tree = ast.parse(SOURCE.read_text(encoding="utf-8"))
    missing_docstrings = [
        n.name
        for n in ast.walk(tree)
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef))
        and not ast.get_docstring(n)
    ]
    checks["function_docstrings"] = not missing_docstrings
    plot_tree = ast.parse(
        (SOURCE.parent / "generate_figures.py").read_text(encoding="utf-8")
    )
    checks["plotting_syntax_and_docstrings"] = all(
        ast.get_docstring(node)
        for node in ast.walk(plot_tree)
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
    )
    probe_config = Config(n=16, horizon=60.0)
    _, _, probe_audit = audit(save=False)
    probe = simulate(probe_config, "resume_probe", env, radius, probe_audit)
    canonical = case_path("resume_probe", case_identity(probe_config, probe_audit))
    if probe == canonical:
        retry = canonical.with_name(canonical.name + "-retry-probe")
        canonical.rename(retry)
        canonical.mkdir()
    else:
        retry = probe
    recovered = simulate(probe_config, "resume_probe", env, radius, probe_audit)
    checks["completed_retry_reuse"] = recovered == retry
    report = {
        "status": "PASS" if all(checks.values()) else "FAIL",
        "checks": {key: "PASS" if value else "FAIL" for key, value in checks.items()},
        "property_samples_a_k_D": [a.tolist(), k.tolist(), D.tolist()],
        "BDF_field_difference_T_C": difference.tolist(),
        "BDF_evaluations": bdf.nfev,
        "independent_check_scope": "N=80, 0..60 s; independent time integrator, shared spatial operator",
        "missing_docstrings": missing_docstrings,
        "source_sha256": sha256(SOURCE),
    }
    json_write(OUT / "review/self_tests.json", report)
    print(json.dumps(report, ensure_ascii=False), flush=True)
    if report["status"] != "PASS":
        raise RuntimeError("Self tests failed")
    return report


def configure_worker() -> None:
    """Keep BLAS single-threaded in each independent numerical worker process."""
    global _BLAS_CONTROLLER
    _BLAS_CONTROLLER = threadpool_limits(limits=1)


def run_batch(
    jobs: list, env: np.ndarray, radius: np.ndarray, audit_report: dict, workers: int
) -> dict:
    """Run independent cases in bounded processes with all implementation in this script."""
    if workers == 1:
        return {
            name: simulate(config, name, env, radius, audit_report)
            for name, config in jobs
        }
    with ProcessPoolExecutor(
        max_workers=workers, initializer=configure_worker
    ) as executor:
        futures = {
            name: executor.submit(simulate, config, name, env, radius, audit_report)
            for name, config in jobs
        }
        return {name: future.result() for name, future in futures.items()}


def verification(
    env: np.ndarray,
    radius: np.ndarray,
    audit_report: dict,
    start_n: int = 1600,
    max_n: int = 51200,
    workers: int = 3,
    prod_tail: str = "fluct",
    prod_seed: int = 0,
) -> dict:
    """Execute separate space/time refinement, local event and scenario validations."""
    targets = {
        "thresholds": TARGETS,
        "source_sha256": sha256(SOURCE),
        "declared_before_runs": True,
        "production_tail": prod_tail,
        "production_seed": prod_seed if prod_tail == "fluct" else None,
        "rationale": "C、T 最细场差小于四位小数半单位；时长差另限 5 s；差值非严格误差界。"
        "独立平衡相对误差 1e-7；热 Robin 1e-3 W/m2；湿 Robin 1e-9 (kg/kg)m/s。"
        "0..60 s 启动场、全部共同 60 s 全场、表面 2% 薄层、终点前 1 h 全部覆盖。"
        "生产族默认采用随机波动环境（与第三问一致），热场加密阈值单独放宽到 1e-3 °C，湿场仍 5e-5 kg/kg。",
    }
    json_write(OUT / "review/acceptance_targets.json", targets)
    tests = self_tests(env, radius)
    prod_kwargs = {"tail": prod_tail}
    if prod_tail == "fluct":
        prod_kwargs["seed"] = prod_seed
    paths, spatial = [], []
    n = start_n
    while n <= max_n:
        current = simulate(
            Config(n=n, **prod_kwargs), f"space_n{n}", env, radius, audit_report
        )
        paths.append(current)
        if len(paths) > 1:
            comparison = compare_cases(paths[-2], paths[-1])
            spatial.append(comparison)
            json_write(OUT / "review/space_progress.json", spatial)
            print("SPACE " + json.dumps(comparison, ensure_ascii=False), flush=True)
            if len(paths) >= 3 and comparison["status"] == "PASS":
                break
        n *= 2
    if len(paths) < 3 or not spatial or spatial[-1]["status"] != "PASS":
        raise RuntimeError(
            "Spatial refinement failed declared targets; see space_progress.json"
        )
    finest = paths[-1]
    time_paths = run_batch(
        [
            ("main", Config(n=n, factor=0.5, **prod_kwargs)),
            ("time_quarter", Config(n=n, factor=0.25, **prod_kwargs)),
        ],
        env,
        radius,
        audit_report,
        min(workers, 2),
    )
    half, quarter = time_paths["main"], time_paths["time_quarter"]
    temporal = [compare_cases(finest, half), compare_cases(half, quarter)]
    main_checks = case_checks(half)
    event = event_checks(half, env, radius)
    checks = {"self_tests": tests, "main": main_checks, "event": event}
    report = {
        "status": "PENDING",
        "source_sha256": sha256(SOURCE),
        "inputs": audit_report["sha256"],
        "targets": targets,
        "main": str(half.relative_to(OUT)),
        "N": n,
        "space": spatial,
        "time": temporal,
        "checks": checks,
        "scenarios": [],
    }
    json_write(OUT / "review/verification_pending.json", report)
    if (
        any(x["status"] != "PASS" for x in temporal)
        or main_checks["status"] != "PASS"
        or event["status"] == "FAIL"
    ):
        raise RuntimeError(
            "Main time, boundary or event checks failed; see verification_pending.json"
        )
    cases = [
        ("appendix3_fixed", 3, False, "last"),
        ("appendix3_shrink", 3, True, "last"),
        ("appendix4_fixed", 4, False, "last"),
        ("appendix4_last", 4, True, "last"),
        ("appendix4_mean30", 4, True, "mean30"),
    ]
    jobs = []
    for label, appendix, shrink, tail in cases:
        for suffix, grid, factor in (
            ("", n, 0.5),
            ("_coarse", n // 2, 0.5),
            ("_time", n, 0.25),
        ):
            jobs.append(
                (
                    label + suffix,
                    Config(
                        n=grid,
                        factor=factor,
                        appendix=appendix,
                        shrink=shrink,
                        tail=tail,
                        seed=0 if tail == "fluct" else None,
                    ),
                )
            )
    scenario_paths = run_batch(jobs, env, radius, audit_report, workers)
    for label, appendix, shrink, tail in cases:
        case = scenario_paths[label]
        coarse = scenario_paths[label + "_coarse"]
        finer_time = scenario_paths[label + "_time"]
        check = {
            "name": label,
            "path": str(case.relative_to(OUT)),
            "space": compare_cases(coarse, case),
            "time": compare_cases(case, finer_time),
            "checks": case_checks(case),
            "event": event_checks(case, env, radius),
        }
        check["status"] = (
            "PASS"
            if all(check[k]["status"] == "PASS" for k in ("space", "time", "checks"))
            and check["event"]["status"] != "FAIL"
            else "FAIL"
        )
        report["scenarios"].append(check)
        json_write(OUT / "review/verification_pending.json", report)
        if check["status"] != "PASS":
            raise RuntimeError(f"Scenario validation failed: {label}")
    report["status"] = "PASS"
    report["case_metadata_sha256"] = {
        str(p.relative_to(OUT)): sha256(p / "metadata.json")
        for p in [
            *paths,
            half,
            quarter,
            *scenario_paths.values(),
        ]
    }
    json_write(OUT / "review/verification.json", report)
    return report


def mapped_moisture(
    state: np.ndarray, x: np.ndarray, R: float, distances_m: np.ndarray
) -> list:
    """Map fixed actual distances strictly inside the current radius, followed by its surface."""
    return [
        float(np.interp(r / R, x, state[1])) if r < R - 1e-12 else None
        for r in distances_m
    ] + [float(state[1, -1])]


def verify_delivery(
    path: Path,
    expected_rows: list,
    headers: list,
    table_rows: list,
    table_distances: np.ndarray,
) -> dict:
    """Read every workbook cell and CSV value back against the same unrounded full states."""
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if book.sheetnames != ["Sheet1"]:
        raise AssertionError("Delivery workbook has unexpected sheets")
    sheet = book["Sheet1"]
    if sheet.max_row != len(expected_rows) + 1 or sheet.max_column != len(headers):
        raise AssertionError("Delivery shape mismatch")
    if [c.value for c in next(sheet.iter_rows())] != headers:
        raise AssertionError("Delivery headers mismatch")
    checked, blanks, max_error = 0, 0, 0.0
    for cells, expected in zip(sheet.iter_rows(min_row=2), expected_rows):
        for j, (cell, value) in enumerate(zip(cells, expected)):
            checked += 1
            if value is None:
                if cell.value is not None:
                    raise AssertionError(
                        f"Outside-domain cell populated: {cell.coordinate}"
                    )
                blanks += 1
            else:
                if not isinstance(cell.value, (int, float)) or cell.data_type == "f":
                    raise AssertionError(
                        f"Non-numeric delivery cell: {cell.coordinate}"
                    )
                error = abs(cell.value - value)
                max_error = max(max_error, error)
                if error > max(1e-12, abs(value) * 2e-15):
                    raise AssertionError(f"Delivery value mismatch: {cell.coordinate}")
                if j > 0 and cell.number_format != "0.0000":
                    raise AssertionError(f"Display format mismatch: {cell.coordinate}")
    book.close()
    with (path.parent / "table6_moisture.csv").open(
        encoding="utf-8-sig", newline=""
    ) as stream:
        rows = list(csv.reader(stream))
    if len(rows) != len(table_rows) + 1:
        raise AssertionError("Table 6 row count mismatch")
    for saved, expected in zip(rows[1:], table_rows):
        if len(saved) != len(expected):
            raise AssertionError("Table 6 column count mismatch")
        for value, source in zip(saved, expected):
            if (source is None and value != "") or (
                source is not None and abs(float(value) - source) > 1e-12
            ):
                raise AssertionError("Table 6 numerical readback mismatch")
    return {
        "status": "PASS",
        "cells_checked": checked,
        "blank_mask_cells": blanks,
        "max_roundtrip_difference": max_error,
        "table6_rows": len(table_rows),
        "table6_distance_cm": (100 * table_distances).tolist(),
    }


def delivery(report: dict, audit_report: dict) -> None:
    """Publish workbook, tables and documentation only after all numerical prerequisites pass."""
    if report["status"] != "PASS" or report["source_sha256"] != sha256(SOURCE):
        raise RuntimeError("Unverified or stale batch cannot publish")
    for relative, digest in report["case_metadata_sha256"].items():
        if sha256(OUT / relative / "metadata.json") != digest:
            raise ValueError("Verified case metadata changed")
    case = OUT / report["main"]
    meta = load_case(case)
    if meta["event"] is None:
        text = (
            "# 第四问数据覆盖内结果\n\n主方案在实测半径覆盖的 72 h 内未达标。"
            f"终点全域最大含水率为 {meta['end_max_C']:.12g} kg/kg。\n\n"
            "已保存完整计算场与验证记录。缺少后续半径输入，未生成声称烘干结束的表格。"
            "后续可选择补充实测半径，或明确采用末半径不变情景后再计算。\n"
        )
        (OUT / "结果与验证.md").write_text(text, encoding="utf-8")
        return
    env, radius, _ = audit(save=False)
    config = Config(**meta["identity"]["config"])
    boundary = Boundary(env, radius, config)
    x = np.linspace(0, 1, config.n + 1)
    distances = np.arange(0.0, boundary.radius(60), 0.001)
    table_distances = np.arange(0.0, boundary.radius(21600), 0.005)
    rows, table_rows, radii = [], [], []
    for t, state in iter_fields(case, meta):
        end = abs(t - meta["end_s"]) < 1e-8
        regular = t >= 60 and abs(t / 60 - round(t / 60)) < 1e-8
        R = boundary.radius(t)
        if regular or end:
            rows.append([t, *mapped_moisture(state, x, R, distances)])
            radii.append([t, R, R * 100])
        if (t >= 21600 and abs(t / 21600 - round(t / 21600)) < 1e-8) or end:
            table_rows.append(
                [t / 3600, *mapped_moisture(state, x, R, table_distances)]
            )
    batch = OUT / "review" / ("delivery_" + case.name)
    batch.mkdir(parents=True, exist_ok=True)
    headers = [
        "时间\\到药材中心的距离",
        *(float(round(r * 100, 10)) for r in distances),
        "药材表面",
    ]
    book = openpyxl.load_workbook(ROOT / "data/附件3/result4.xlsx")
    sheet = book["Sheet1"]
    sheet.delete_rows(1, sheet.max_row)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    sheet.freeze_panes = "B2"
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 30
    for j in range(2, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(j)].width = 12
    sheet.column_dimensions[get_column_letter(len(headers))].width = 16
    for cell in sheet[1]:
        cell.font = Font(name="Microsoft YaHei", size=10, bold=True)
        cell.fill = PatternFill("solid", fgColor="E8EDF0")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = Border(bottom=Side(style="thin", color="657178"))
    sheet.row_dimensions[1].height = 32
    for row in sheet.iter_rows(min_row=2):
        for j, cell in enumerate(row):
            cell.font = Font(name="Microsoft YaHei", size=10)
            cell.number_format = "0.0000" if j else "0.########"
            cell.alignment = Alignment(horizontal="right", vertical="center")
    sheet.print_options.horizontalCentered = True
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    book.save(batch / "result4.xlsx")
    book.close()
    table_headers = ["时间/h", *(f"{r*100:g} cm" for r in table_distances), "药材表面"]
    with (batch / "table6_moisture.csv").open(
        "w", encoding="utf-8-sig", newline=""
    ) as stream:
        writer = csv.writer(stream)
        writer.writerow(table_headers)
        writer.writerows(table_rows)
    lines = [
        "# 表 6 药材烘干过程的水分浓度",
        "",
        "单位：kg/kg；固定列为实际距离，空白表示位于药材外。",
        "",
        "| " + " | ".join(table_headers) + " |",
        "|" + "---:|" * len(table_headers),
    ]
    for j, row in enumerate(table_rows):
        label = f"{row[0]:.8f}" if j == len(table_rows) - 1 else f"{row[0]:g}"
        if j == len(table_rows) - 1:
            label = "结束 " + label
        lines.append(
            "| "
            + label
            + " | "
            + " | ".join("" if v is None else f"{v:.4f}" for v in row[1:])
            + " |"
        )
    lines += [
        "",
        f"严格结束时刻：{meta['end_s']:.12f} s；全域未舍入最大值 {meta['end_max_C']:.15f} kg/kg。",
        "末行显示为 0.1500 时，严格判据仍以未舍入全场为准。",
    ]
    (batch / "table6.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    check = verify_delivery(
        batch / "result4.xlsx", rows, headers, table_rows, table_distances
    )
    json_write(batch / "delivery_checks.json", check)
    if input_hashes() != audit_report["sha256"]:
        raise RuntimeError("Raw inputs changed during export")
    # Only verified artifacts are promoted from this batch to formal paths.
    for filename in ("result4.xlsx", "table6_moisture.csv", "table6.md"):
        target = OUT / filename
        if target.exists() and target.read_bytes() != (batch / filename).read_bytes():
            archive = (
                OUT / "review" / ("previous_" + sha256(target)[:16] + "_" + filename)
            )
            if not archive.exists():
                archive.write_bytes(target.read_bytes())
        target.write_bytes((batch / filename).read_bytes())
    np.savetxt(
        OUT / "data/radius_history.csv",
        radii,
        delimiter=",",
        header="time_s,radius_m,radius_cm",
        comments="",
    )
    json_write(OUT / "review/delivery_checks.json", check)
    write_report(report, meta, check, lines, audit_report)
    manifest = {
        "status": "PASS",
        "main": report["main"],
        "source_sha256": sha256(SOURCE),
        "inputs": audit_report["sha256"],
        "verification_sha256": sha256(OUT / "review/verification.json"),
        "main_metadata_sha256": sha256(case / "metadata.json"),
        "files": {
            f: sha256(OUT / f)
            for f in (
                "result4.xlsx",
                "table6.md",
                "table6_moisture.csv",
                "data/radius_history.csv",
                "data/ambient_observed.csv",
                "data/radius_observed_SI.csv",
                "结果与验证.md",
                "review/delivery_checks.json",
            )
        },
    }
    json_write(OUT / "delivery_manifest.json", manifest)


def write_report(
    report: dict,
    meta: dict,
    workbook_check: dict,
    table_lines: list,
    audit_report: dict,
) -> None:
    """Explain actual results, indicator definitions, convergence evidence and model limits."""
    event = meta["event"]
    spatial, temporal = report["space"][-1], report["time"][-1]
    space_delta = abs(spatial["event_delta_s"] or 0)
    time_delta = abs(temporal["event_delta_s"] or 0)
    empirical_error = math.ceil(
        2 * (space_delta + time_delta) + event["width_s"] + 0.01
    )
    lines = [
        "# A 题第四问结果与验证",
        "",
        f"主模型临界时长 **{event['critical_s']/3600:.8f} h**；严格达标结束时刻 **{event['end_s']/3600:.8f} h**。",
        f"未舍入全域最大含水率 {meta['end_max_C']:.15f} kg/kg，裕度 {event['margin_kg_kg']:.9e} kg/kg。",
        f"当前最细空间、时间差支持约 **{empirical_error} s 的经验数值误差尺度**，非严格上界或工艺安全余量。",
        "",
        "## 输入与假设",
        "",
        "原始题面、附件和模板只读；重新逐行审计未发现需要处理的异常。半径 cm 转 m，分段线性插值，未拟合或平滑。",
        f"附件 1 有 {audit_report['sheets'][0]['count']} 条观测，附件 2 有 {audit_report['sheets'][1]['count']} 条观测。",
        (
            f"主方案 4 h 后环境温度 {meta['environment_tail'][0]:.9g} °C、环境干基水分浓度 {meta['environment_tail'][1]:.9g} kg/kg 保持末值。"
            if meta.get("environment_mode", "last") != "fluct"
            else f"主方案 4 h 后环境取末 30 min 均值 {meta['environment_tail'][0]:.9g} °C、"
            f"{meta['environment_tail'][1]:.9g} kg/kg，并叠加固定种子 {meta['environment_seed']} 的 AR(1) 随机波动（与第三问生产情景一致）。"
        ),
        "长度固定 0.25 m，骨架均匀径向收缩；初值 28 °C、2.55 kg/kg、半径 0.02 m。实测半径适用范围止于 72 h。",
        "主方程、边界、干基守恒消项和密度解释分别对应原推导公式（57）至（61）、（25）至（36）、（89）至（90）。",
        "实现推导与全部指标定义见 [模型与算法说明](../../src/problem-4/问题4_模型与算法说明.md)。",
        "",
        "## 全域事件",
        "",
        "| 指标 | 数值 |",
        "|---|---:|",
        *[
            f"| {key} | {value:.12g} |"
            for key, value in event.items()
            if isinstance(value, (float, int))
        ],
        f"| 全程最大值高于中心的最大差 / kg/kg | {meta['max_C_above_center']:.6e} |",
        f"| 径向反序最大量 / kg/kg | {meta['radial_inversion_kg_kg']:.6e} |",
        "",
        "每个接受状态检查全部计算节点；分段线性重构不会在节点间产生更大值。连续解误差由独立加密研究评估。",
        "临界时间取正负夹逼中点；严格结束取经过重新积分确认的上端点。定位宽度不含累计积分或空间误差。",
        "",
        *table_lines,
        "",
        "## 数值参数与守恒",
        "",
        f"N={meta['identity']['config']['n']}，SDIRK2，时间倍率 {meta['identity']['config']['factor']}。",
        f"实际步长 {meta['step_range_s'][0]:.8g} 至 {meta['step_range_s'][1]:.8g} s；接受 {meta['steps']} 步，拒绝 {meta['rejected_steps']} 步，阶段最多 {meta['max_iterations']} 次迭代。",
        "所有隐式阶段按各自时刻更新半径、环境和局部物性；遇到输入插值节点及保存时刻精确停步。",
        f"水分独立平衡相对缺陷 {meta['mass_relative_defect']:.9e}，最大绝对缺陷 {meta['mass_max_abs_defect_kg_kg']:.9e} kg/kg。",
        f"显热独立平衡相对缺陷 {meta['heat_relative_defect']:.9e}，最大累计差 {meta['heat_max_abs_defect_J']:.9e} J。",
        f"两场最大尺度化迭代变化 {meta['max_scaled_change']:.9e}，原非线性方程缺陷 {meta['max_scaled_residual']:.9e}。",
        "水分边界累计来自实际阶段表面通量；显热分别累计变热容储存率和表面热输入，不使用热容乘温度的首末差。",
        f"60 s 及以后边界梯度独立重构最大绝对残差（热 Robin、湿 Robin、轴线温度梯度、轴线水分梯度）为 {report['checks']['main']['boundary_abs_max']}，单位依次 W/m²、(kg/kg)m/s、°C/m、(kg/kg)/m。",
        "初始角点不相容单列诊断，不以 t=0 的梯度作为已满足瞬时 Robin 的证据。",
        "",
    ]
    for title, comparisons in (
        ("空间加密", report["space"]),
        ("时间加密", report["time"]),
    ):
        lines += [
            f"## {title}",
            "",
            "| N / 时间倍率 | 全场 C 最大差 kg/kg | 全场 T 最大差 °C | 时长差 s | 状态 |",
            "|---|---:|---:|---:|---|",
        ]
        for item in comparisons:
            lines.append(
                f"| {item['N']} / {item['factors']} | {item['metrics']['full']['C']:.9e} | "
                f"{item['metrics']['full']['T']:.9e} | {item['event_delta_s']} | {item['status']} |"
            )
        lines += [
            "",
            "粗级未达到目标的 FAIL 如实保留；最终验收使用最细相邻级。完整比较覆盖全部共同保存的双场、细网格点与实际交付点。",
            "全场误差用粗网格分段线性重构在每个细网格点计算，含节点间重构差；启动、表面薄层和临近终点独立记录。",
            "",
        ]
    lines += [
        "## 事件与独立时间方法",
        "",
        f"事件细化证据：`{json.dumps(report['checks']['event'], ensure_ascii=False)}`",
        "",
        f"独立 BDF 对照误差 T/C：{report['checks']['self_tests']['BDF_field_difference_T_C']}，范围 N=80、0 至 60 s。",
        "BDF 与 SDIRK2 共享空间通量算子，因此仅为独立时间方法核对。",
        "",
        "## 物性、几何与环境对照",
        "",
        "| 情景 | 临界时长 h | 覆盖终点最大 C | 验证 |",
        "|---|---:|---:|---|",
    ]
    lines.append(
        f"| 附录 4 实测收缩（主方案·随机波动 seed {meta.get('environment_seed')}） | "
        f"{event['critical_s']/3600:.8f} | {meta['end_max_C']:.9g} | PASS |"
    )
    for case in report["scenarios"]:
        m = json_read(OUT / case["path"] / "metadata.json")
        label = {
            "appendix3_fixed": "附录 3 固定半径",
            "appendix3_shrink": "附录 3 实测收缩",
            "appendix4_fixed": "附录 4 固定半径",
            "appendix4_last": "附录 4 收缩、末值延拓",
            "appendix4_mean30": "附录 4 收缩、末 30 min 均值",
            "appendix4_fluct": "附录 4 收缩、随机波动（种子 0）",
        }[case["name"]]
        duration = (
            f"{m['event']['critical_s']/3600:.8f}" if m["event"] else "72 h 内未达标"
        )
        lines.append(
            f"| {label} | {duration} | {m['end_max_C']:.9g} | {case['status']} |"
        )
    lines += [
        "",
        f"均值情景窗口及算术均值：{meta['mean30_window']}。观测区间不变，4 h 后立即切换均值，可能有小跳变。",
        "主方案随机波动按观测末段去趋势残差估计标准差与滞后一阶自相关，新息自助重采样，种子固定，事件可复现；对照情景分别取末值和末 30 min 均值。",
        "对照均从相同初值出发，并进行各自空间与时间核查。由原推导公式（91）至（94），附录物性变化与收缩几何共同决定历程，不能仅按末半径缩放第三问时长；非线性效应不具有普遍可加性。",
        "",
        "## 表格与运行证据",
        "",
        f"工作簿逐格重读 {workbook_check['cells_checked']} 格，几何空格 {workbook_check['blank_mask_cells']} 格，最大写回差 {workbook_check['max_roundtrip_difference']:.9e}。",
        "Sheet1 仅包含干基含水率；从 60 s 起每 60 s 加同一严格终点；固定列为实际 cm，表面单独一列。",
        "固定距离达到或超出实际半径时留空。每行实际半径见 data/radius_history.csv。表 6 每 6 h 加结束行，由同一未舍入场提取。",
        f"求解脚本 SHA-256：`{sha256(SOURCE)}`。依赖：`{json.dumps(dependencies())}`。",
        f"主案例：`{report['main']}`。每个完整场分块文件均有 SHA-256，完整验证见 review/verification.json。",
        "",
        "| 原始文件 | SHA-256 |",
        "|---|---|",
        *[
            f"| {name} | `{digest}` |"
            for name, digest in audit_report["sha256"].items()
        ],
        "",
        "## 限制与未完成事项",
        "",
        "经验密度仅用于显热容量，不能同时强制作为均匀收缩下的真实湿密度。模型未计潜热、湿分携带焓、压缩功、端面和内部非均匀变形。",
        "四位小数为显示要求；相邻细化差不是严格误差界，数值严格达标裕度不是实际工艺安全余量。4 h 后环境和内部同缩都是假设，尚无独立实验验证。",
        "绘图由 generate_figures.py 单独执行，图像数据指纹及视觉检查状态见 image/figure_manifest.json；本数值报告不预先宣称图像检查完成。",
        "",
    ]
    (OUT / "结果与验证.md").write_text("\n".join(lines), encoding="utf-8")


def main() -> None:
    """Dispatch audited, resumable single-case and full-delivery command-line operations."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--audit", action="store_true")
    parser.add_argument("--self-test", action="store_true")
    parser.add_argument("--all", action="store_true")
    parser.add_argument("--verify", action="store_true")
    parser.add_argument("--deliver", action="store_true")
    parser.add_argument("--case")
    parser.add_argument("--inspect-case", type=Path)
    parser.add_argument("--compare", nargs=2, type=Path)
    parser.add_argument("--n", type=int, default=1600)
    parser.add_argument("--max-n", type=int, default=51200)
    parser.add_argument("--workers", type=int, choices=(1, 2, 3, 4), default=3)
    parser.add_argument("--factor", type=float, default=1.0)
    parser.add_argument("--appendix", type=int, choices=(3, 4), default=4)
    parser.add_argument("--fixed", action="store_true")
    parser.add_argument("--tail", choices=("last", "mean30", "fluct"), default="last")
    parser.add_argument("--seed", type=int, default=None)
    parser.add_argument("--tau", type=float, default=None)
    parser.add_argument("--sigma-scale", type=float, default=1.0)
    parser.add_argument(
        "--prod-tail", choices=("last", "mean30", "fluct"), default="fluct"
    )
    parser.add_argument("--prod-seed", type=int, default=0)
    parser.add_argument("--stop-s", type=float, default=259200.0)
    parser.add_argument("--method", choices=("sdirk2", "be"), default="sdirk2")
    args = parser.parse_args()
    if not any(
        (
            args.audit,
            args.self_test,
            args.all,
            args.verify,
            args.deliver,
            args.case,
            args.inspect_case,
            args.compare,
        )
    ):
        parser.error(
            "Specify --audit, --self-test, --case NAME, --verify, --deliver or --all"
        )
    with threadpool_limits(limits=1):
        run_commands(args)


def run_commands(args: argparse.Namespace) -> None:
    """Execute selected CLI operations with recorded, bounded BLAS threads."""
    env, radius, audit_report = audit()
    if args.audit:
        print(json.dumps(audit_report, ensure_ascii=False, indent=2))
    if args.self_test:
        self_tests(env, radius)
    if args.case:
        simulate(
            Config(
                n=args.n,
                factor=args.factor,
                appendix=args.appendix,
                shrink=not args.fixed,
                tail=args.tail,
                horizon=args.stop_s,
                method=args.method,
                seed=args.seed,
                tau_s=args.tau,
                sigma_scale=args.sigma_scale,
            ),
            args.case,
            env,
            radius,
            audit_report,
        )
    if args.inspect_case:
        print(
            json.dumps(
                case_checks(args.inspect_case.resolve()), ensure_ascii=False, indent=2
            )
        )
    if args.compare:
        print(
            json.dumps(
                compare_cases(*(path.resolve() for path in args.compare)),
                ensure_ascii=False,
                indent=2,
            )
        )
    if args.all or args.verify:
        before = protected_hashes()
        json_write(OUT / "review/protected_before.json", before)
        report = verification(
            env,
            radius,
            audit_report,
            args.n,
            args.max_n,
            args.workers,
            args.prod_tail,
            args.prod_seed,
        )
        after = protected_hashes()
        json_write(OUT / "review/protected_after.json", after)
        json_write(
            OUT / "review/preservation.json",
            {
                "status": "PASS" if before == after else "FAIL",
                "files_checked": len(before),
                "changed": [k for k in before if before[k] != after.get(k)],
            },
        )
        if before != after:
            raise RuntimeError("Protected files changed")
        if args.all:
            delivery(report, audit_report)
    if args.deliver:
        delivery(json_read(OUT / "review/verification.json"), audit_report)


if __name__ == "__main__":
    main()
