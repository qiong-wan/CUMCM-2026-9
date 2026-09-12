"""Generate Chinese scientific figures from verified Problem 4 data only.

No solver import, PDE integration or mutation of numerical case files occurs.
PNG files are written but never read. JPEG previews come directly from Figures.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
import warnings

import matplotlib

matplotlib.use("Agg")

from matplotlib import font_manager, pyplot as plt
from matplotlib.colors import Normalize
from matplotlib.text import Text
import numpy as np
import openpyxl


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "output/problem-4"
IMAGES = OUT / "image"
PREVIEWS = ROOT.parent / "tmp/problem4_figure_previews"
COLORS = ("#0072B2", "#009E73", "#E69F00", "#D55E00", "#CC79A7", "#56B4E9")
LABELS = {
    "appendix3_fixed": "附录 3 · 固定半径",
    "appendix3_shrink": "附录 3 · 实测收缩",
    "appendix4_fixed": "附录 4 · 固定半径",
    "appendix4_mean30": "附录 4 · 末段均值",
}


def sha256(path: Path) -> str:
    """Read a non-PNG input or SVG file in bounded blocks for provenance."""
    if path.suffix.lower() == ".png":
        raise ValueError("Reading PNG output is forbidden")
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def read_json(path: Path) -> dict:
    """Read a saved numerical or figure manifest."""
    return json.loads(path.read_text(encoding="utf-8"))


def configure_style() -> str:
    """Use a verified CJK font and the established preceding-question publication palette."""
    installed = {item.name for item in font_manager.fontManager.ttflist}
    candidates = ("Microsoft YaHei", "SimHei", "Noto Sans CJK SC", "Source Han Sans SC")
    family = next((name for name in candidates if name in installed), None)
    if family is None:
        raise RuntimeError("No supported Chinese font installed")
    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [family, "DejaVu Sans"],
            "axes.unicode_minus": False,
            "mathtext.fontset": "stix",
            "svg.fonttype": "path",
            "font.size": 10,
            "axes.titlesize": 11,
            "axes.labelsize": 10,
            "legend.fontsize": 9,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "axes.edgecolor": "#3D4652",
            "axes.labelcolor": "#20242A",
            "text.color": "#20242A",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "figure.facecolor": "white",
            "savefig.facecolor": "white",
            "axes.formatter.use_mathtext": True,
        }
    )
    warnings.filterwarnings("error", message="Glyph .* missing from font.*")
    return family


def load_verified() -> tuple:
    """Reject stale sources, unverified batches and altered numerical chunks before plotting."""
    manifest = read_json(OUT / "delivery_manifest.json")
    if manifest["status"] != "PASS":
        raise RuntimeError("Delivery is not verified")
    solver = ROOT / "src/problem-4/solve_problem4.py"
    if sha256(solver) != manifest["source_sha256"]:
        raise RuntimeError("Solver source changed after verification")
    for name, digest in manifest["inputs"].items():
        if sha256(ROOT / name) != digest:
            raise RuntimeError(f"Input changed: {name}")
    for name, digest in manifest["files"].items():
        if sha256(OUT / name) != digest:
            raise RuntimeError(f"Delivery changed: {name}")
    if sha256(OUT / "review/verification.json") != manifest["verification_sha256"]:
        raise RuntimeError("Verification changed")
    report = read_json(OUT / "review/verification.json")
    if report["status"] != "PASS":
        raise RuntimeError("Numerical prerequisite checks not passed")
    case_paths = [report["main"], *(case["path"] for case in report["scenarios"])]
    numerical_hashes = {}
    for relative in case_paths:
        path = OUT / relative
        meta = read_json(path / "metadata.json")
        expected_meta = report["case_metadata_sha256"][relative]
        if sha256(path / "metadata.json") != expected_meta:
            raise RuntimeError("Case metadata altered")
        numerical_hashes[str(path / "metadata.json")] = expected_meta
        for filename, digest in meta["files"].items():
            source = path / filename
            if sha256(source) != digest:
                raise RuntimeError(f"Case data altered: {source}")
            numerical_hashes[str(source)] = digest
    main_path = OUT / report["main"]
    main_meta = read_json(main_path / "metadata.json")
    with np.load(main_path / "summary.npz", allow_pickle=False) as data:
        summary = {name: data[name].copy() for name in data.files}
    env = np.loadtxt(OUT / "data/ambient_observed.csv", delimiter=",", skiprows=1)
    radius = np.loadtxt(OUT / "data/radius_observed_SI.csv", delimiter=",", skiprows=1)
    for original, derived, relative, conversion in (
        ("附件1.xlsx", env, "data/ambient_observed.csv", 1.0),
        ("附件2.xlsx", radius, "data/radius_observed_SI.csv", 0.01),
    ):
        book = openpyxl.load_workbook(
            ROOT / "data" / original, read_only=True, data_only=True
        )
        observed = np.asarray(
            list(book["Sheet1"].iter_rows(min_row=2, values_only=True))
        )
        book.close()
        observed[:, 1:] *= conversion
        if not np.array_equal(observed, derived):
            raise RuntimeError(
                f"Plot input copy differs from audited observations: {relative}"
            )
        numerical_hashes[str(OUT / relative)] = sha256(OUT / relative)
    return report, main_meta, summary, env, radius, numerical_hashes


def finish_axes(ax, xlabel: str, ylabel: str, title: str = "") -> None:
    """Apply quiet scientific axes with visible units and light major grids."""
    ax.set(xlabel=xlabel, ylabel=ylabel, title=title)
    ax.grid(True, color="#D8DEE6", alpha=0.7, linewidth=0.55)
    ax.set_axisbelow(True)


def save_figure(fig, name: str, sources: list, scope: str, manifest: list) -> None:
    """Check canvas and text bounds, then export SVG/PNG and a directly rendered JPEG."""
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    bounds = fig.bbox
    outside = []
    hidden_ticks = set()
    for ax in fig.axes:
        for axis, limits in ((ax.xaxis, ax.get_xlim()), (ax.yaxis, ax.get_ylim())):
            low, high = sorted(limits)
            for tick in (*axis.get_major_ticks(), *axis.get_minor_ticks()):
                if tick.get_loc() < low or tick.get_loc() > high:
                    hidden_ticks.update((id(tick.label1), id(tick.label2)))
    for artist in fig.findobj(match=Text):
        if (
            id(artist) in hidden_ticks
            or not artist.get_visible()
            or not artist.get_text().strip()
        ):
            continue
        box = artist.get_window_extent(renderer)
        if (
            box.width > 0
            and box.height > 0
            and (
                box.x0 < bounds.x0 - 2
                or box.y0 < bounds.y0 - 2
                or box.x1 > bounds.x1 + 2
                or box.y1 > bounds.y1 + 2
            )
        ):
            outside.append(artist.get_text())
    if outside:
        raise RuntimeError(f"Figure text outside canvas: {name}: {outside}")
    pixels = np.asarray(fig.canvas.buffer_rgba())
    fraction = float(np.mean(np.any(pixels[:, :, :3] < 235, axis=2)))
    if fraction < 0.005:
        raise RuntimeError(f"Blank figure: {name}")
    fig.savefig(IMAGES / f"{name}.svg", bbox_inches="tight")
    fig.savefig(IMAGES / f"{name}.png", dpi=320, bbox_inches="tight")
    fig.savefig(
        PREVIEWS / f"{name}.jpg",
        dpi=150,
        bbox_inches="tight",
        pil_kwargs={"quality": 92},
    )
    manifest.append(
        {
            "name": name,
            "files": [f"{name}.svg", f"{name}.png"],
            "svg_sha256": sha256(IMAGES / f"{name}.svg"),
            "preview": str(PREVIEWS / f"{name}.jpg"),
            "sources": sources,
            "scope_and_limit": scope,
            "text_bounds": "PASS",
            "nonblank_fraction": fraction,
            "font_glyphs": "PASS",
            "dpi": 320,
        }
    )
    plt.close(fig)


def plot_inputs(
    env: np.ndarray, radius: np.ndarray, meta: dict, manifest: list
) -> None:
    """Distinguish observed ambient conditions, assumed extension and observed radius."""
    fig, axes = plt.subplots(3, 1, figsize=(8.5, 8.3), layout="constrained")
    end = meta["end_s"] / 3600
    for j, (label, color) in enumerate(
        (("环境温度 / °C", "#C44E52"), ("环境水分浓度 / (kg/kg)", COLORS[0]))
    ):
        ax = axes[j]
        ax.plot(
            env[:, 0] / 3600,
            env[:, j + 1],
            color=color,
            lw=1.3,
            label="附件 1 实测插值",
        )
        ax.plot(
            [4, end],
            [meta["environment_tail"][j]] * 2,
            "--",
            color=color,
            label="4 h 后末值假设",
        )
        ax.axvline(4, color="#72777C", lw=0.8, ls=":")
        ax.set_xlim(0, end)
        finish_axes(ax, "时间 / h", label)
        ax.legend(loc="lower right")
    axes[2].plot(
        radius[:, 0] / 3600,
        100 * radius[:, 1],
        color=COLORS[1],
        lw=1.2,
        marker="o",
        markersize=2,
        markevery=4,
        label="附件 2 半径",
    )
    axes[2].axvline(end, color="#50565C", ls="--", lw=1, label="主方案结束时刻")
    axes[2].set_xlim(0, 72)
    finish_axes(axes[2], "时间 / h", "实际半径 / cm")
    axes[2].legend()
    save_figure(
        fig,
        "01_inputs_and_radius",
        ["data/ambient_observed.csv", "data/radius_observed_SI.csv"],
        "环境在 4 h 后为假设；半径数据仅覆盖 72 h，内部收缩方式不是外半径数据唯一识别的。",
        manifest,
    )


def plot_history(summary: dict, meta: dict, manifest: list) -> None:
    """Show all-node maximum, center, surface and dry-solid-weighted mean moisture."""
    h = summary["history"]
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.4), layout="constrained")
    series = [
        (4, "中心", COLORS[0], "-"),
        (5, "表面", COLORS[3], "-"),
        (8, "干物质加权平均", COLORS[1], "-"),
        (6, "全域最大", "#32383D", ":"),
    ]
    for ax in axes:
        for column, label, color, style in series:
            ax.plot(
                h[:, 0] / 3600,
                h[:, column],
                style,
                color=color,
                linewidth=1.15 if column != 6 else 1.8,
                label=label,
            )
        ax.axhline(0.15, color="#686D72", ls="--", lw=0.9, label="达标阈值 0.15")
        ax.axvline(meta["event"]["critical_s"] / 3600, color=COLORS[4], ls="--", lw=1)
        ax.axvline(meta["end_s"] / 3600, color="#343A40", ls=":", lw=0.8)
        finish_axes(ax, "时间 / h", "干基含水率 / (kg/kg)")
    axes[0].set_xlim(0, meta["end_s"] / 3600)
    axes[0].legend(loc="upper right")
    axes[1].set_xlim(max(0, meta["end_s"] / 3600 - 12), meta["end_s"] / 3600 + 0.3)
    axes[1].set_ylim(0.045, 0.24)
    axes[1].set_title("末期全域达标")
    axes[1].text(
        0.04,
        0.96,
        f"临界 {meta['event']['critical_s']/3600:.8f} h\n"
        f"严格结束 {meta['end_s']/3600:.8f} h",
        transform=axes[1].transAxes,
        ha="left",
        va="top",
        fontsize=9,
    )
    save_figure(
        fig,
        "02_moisture_history",
        ["main/summary.npz", "main/metadata.json"],
        "中心与最大值可重合；严格判据使用未舍入全部节点；事件夹逼精度不等于累计数值精度。",
        manifest,
    )


def plot_profiles(
    summary: dict, meta: dict, radius: np.ndarray, manifest: list
) -> None:
    """Draw each radial profile only as far as its own actual moving surface."""
    times, states, x = summary["check_times"], summary["check_states"], summary["x"]
    selected = np.flatnonzero(
        (times >= 21600) | np.isclose(times, meta["end_s"], atol=1e-8, rtol=0)
    )
    fig, ax = plt.subplots(figsize=(8.2, 5.1), layout="constrained")
    colors = plt.colormaps["viridis"](np.linspace(0.05, 0.9, len(selected)))
    for color, index in zip(colors, selected):
        t = times[index]
        R = float(np.interp(t, radius[:, 0], radius[:, 1]))
        label = f"{t/3600:g} h" if index != selected[-1] else f"结束 {t/3600:.4f} h"
        ax.plot(100 * R * x, states[index, 1], color=color, lw=1.6, label=label)
        ax.scatter([100 * R], [states[index, 1, -1]], s=15, color=color, zorder=4)
    ax.axhline(0.15, color="#565C62", ls="--", lw=0.9)
    finish_axes(ax, "距药材中心的实际距离 / cm", "干基含水率 / (kg/kg)")
    ax.set_xlim(left=0)
    ax.legend(ncol=2, loc="upper right")
    save_figure(
        fig,
        "03_actual_radius_profiles",
        ["main/summary.npz", "data/radius_observed_SI.csv"],
        "曲线止于各自真实表面；每 6 h 加严格结束时刻；横坐标不是归一化半径。",
        manifest,
    )


def plot_heatmap(
    report: dict, meta: dict, summary: dict, radius: np.ndarray, manifest: list
) -> None:
    """Reconstruct moisture at fixed physical radii and mask the exterior explicitly."""
    case = OUT / report["main"]
    distances = np.linspace(0, 2, 401)
    times, values, boundaries = [], [], []
    for chunk in meta["chunks"]:
        with np.load(case / chunk, allow_pickle=False) as data:
            for t, state in zip(data["times"], data["states"]):
                if 0 < t < 60:
                    continue
                R = 100 * float(np.interp(t, radius[:, 0], radius[:, 1]))
                row = np.full(len(distances), np.nan)
                inside = distances <= R
                row[inside] = np.interp(distances[inside] / R, summary["x"], state[1])
                times.append(t / 3600)
                values.append(row)
                boundaries.append(R)
    values = np.asarray(values)
    if np.isfinite(
        values[np.asarray(distances)[None, :] > np.asarray(boundaries)[:, None]]
    ).any():
        raise AssertionError("Exterior region is not masked")
    fig, ax = plt.subplots(figsize=(9.2, 5.1), layout="constrained")
    ax.set_facecolor("#ECEFF0")
    mesh = ax.pcolormesh(
        times,
        distances,
        np.ma.masked_invalid(values.T),
        shading="nearest",
        cmap="viridis",
        norm=Normalize(vmin=0, vmax=2.55),
        rasterized=True,
    )
    ax.plot(times, boundaries, color="#C44E52", lw=1.6, label="实测收缩边界")
    fig.colorbar(mesh, ax=ax, label="干基含水率 / (kg/kg)", pad=0.025)
    ax.set(xlim=(0, meta["end_s"] / 3600), ylim=(0, 2.0))
    finish_axes(ax, "时间 / h", "距药材中心的实际距离 / cm")
    ax.legend(loc="upper right")
    save_figure(
        fig,
        "04_moisture_space_time",
        [report["main"] + "/fields_*.npz", "data/radius_observed_SI.csv"],
        "一维径向时空图；灰色为药材外部。色块来自保存时刻，未求解二维轴向场。",
        manifest,
    )


def plot_temperature(summary: dict, meta: dict, manifest: list) -> None:
    """Show local versus ambient temperature and diffusion dependence on saved local states."""
    h = summary["history"]
    fig, axes = plt.subplots(1, 2, figsize=(11.0, 4.3), layout="constrained")
    for column, label, color, style in [
        (2, "药材中心", COLORS[0], "-"),
        (3, "药材表面", COLORS[3], "-"),
        (17, "对应烘房环境", "#444A50", "--"),
    ]:
        axes[0].plot(
            h[:, 0] / 3600, h[:, column], style, color=color, lw=1.3, label=label
        )
    axes[0].set_xlim(0, 6)
    finish_axes(axes[0], "时间 / h", "温度 / °C", "初期热传递")
    axes[0].legend()
    for T_column, C_column, label, color in [
        (2, 4, "中心局部扩散系数", COLORS[0]),
        (3, 5, "表面局部扩散系数", COLORS[3]),
    ]:
        D = (
            4.2e-4
            * np.exp(-0.30 / h[:, C_column])
            * np.exp(-3850 / (h[:, T_column] + 273.15))
        )
        axes[1].semilogy(h[:, 0] / 3600, D, color=color, lw=1.3, label=label)
    axes[1].set_xlim(0, meta["end_s"] / 3600)
    finish_axes(axes[1], "时间 / h", "局部扩散系数 / (m²/s)", "局部温湿状态决定扩散")
    axes[1].legend()
    save_figure(
        fig,
        "05_temperature_and_diffusion",
        ["main/summary.npz"],
        "扩散系数按保存的局部温度与干基含水率重构，绝对温度进入指数；未重新积分。",
        manifest,
    )


def plot_convergence(report: dict, meta: dict, manifest: list) -> None:
    """Plot executed spatial and temporal differences and drying-event convergence."""
    space, temporal = report["space"], report["time"]
    fig, axes = plt.subplots(2, 3, figsize=(12.5, 7.5), layout="constrained")
    for row, (items, xlabel, title) in enumerate(
        (
            (space, "细一级径向区间数 N", "空间加密"),
            (temporal, "细一级时间倍率", f"时间加密 · N={report['N']}"),
        )
    ):
        abscissa = [item["N" if row == 0 else "factors"][1] for item in items]
        for column, (field, ylabel) in enumerate(
            (("C", "含水率最大差 / (kg/kg)"), ("T", "温度最大差 / °C"))
        ):
            ax = axes[row, column]
            for region, label, marker, color in (
                ("full", "完整场", "o-", COLORS[0]),
                ("delivery", "实际交付点", "s-", COLORS[1]),
            ):
                ax.loglog(
                    abscissa,
                    [item["metrics"][region][field] for item in items],
                    marker,
                    color=color,
                    label=label,
                )
            ax.axhline(5e-5, color="#666C72", ls="--", lw=0.9, label="场差目标")
            finish_axes(ax, xlabel, ylabel, title)
            ax.legend()
        ax = axes[row, 2]
        ax.loglog(
            abscissa,
            [abs(item["event_delta_s"]) for item in items],
            "o-",
            color=COLORS[3],
            label="相邻级时长差",
        )
        ax.axhline(5, color="#666C72", ls="--", lw=0.9, label="时长差目标 5 s")
        finish_axes(ax, xlabel, "临界时长绝对差 / s", title + " · 时长")
        ax.legend()
        if row == 1:
            for ax in axes[row]:
                ax.invert_xaxis()
    save_figure(
        fig,
        "06_numerical_convergence",
        ["review/verification.json", "review/acceptance_targets.json"],
        "数据来自实际执行的相邻加密；差值是经验精度指标，不是严格误差上界。",
        manifest,
    )


def plot_scenarios(report: dict, meta: dict, manifest: list) -> None:
    """Compare validated geometry/property scenarios and the declared ambient sensitivity."""
    cases = {}
    for case in report["scenarios"]:
        if case["status"] != "PASS":
            raise RuntimeError("Unverified scenario cannot be plotted as accepted")
        cases[case["name"]] = read_json(OUT / case["path"] / "metadata.json")
    fig, axes = plt.subplots(1, 2, figsize=(11.2, 4.8), layout="constrained")
    names = ["appendix3_fixed", "appendix3_shrink", "appendix4_fixed", "main"]
    labels = [LABELS.get(name, "附录 4 · 实测收缩") for name in names]
    values = [cases[name] if name != "main" else meta for name in names]
    for j, value in enumerate(values):
        dry = value["event"] is not None
        hours = value["event"]["critical_s"] / 3600 if dry else value["end_s"] / 3600
        axes[0].barh(
            j,
            hours,
            color=COLORS[j],
            height=0.56,
            hatch=None if dry else "///",
            alpha=0.85,
        )
        label = f"{hours:.4f} h" if dry else f"{hours:g} h 内未达标"
        axes[0].text(hours + 0.8, j, label, va="center", fontsize=9)
    axes[0].set_yticks(range(4), labels)
    axes[0].invert_yaxis()
    axes[0].set_xlim(0, 92)
    finish_axes(axes[0], "临界时长或数据覆盖 / h", "", "相同末值环境下的物性与几何")
    sensitivity = cases["appendix4_mean30"]
    for j, (value, label) in enumerate(
        ((meta, "末次实测值延拓"), (sensitivity, "末 30 min 算术均值延拓"))
    ):
        dry = value["event"] is not None
        hours = value["event"]["critical_s"] / 3600 if dry else 72
        axes[1].barh(j, hours, color=[COLORS[0], COLORS[4]][j], height=0.45, alpha=0.85)
        axes[1].text(
            hours + 0.4,
            j,
            f"{hours:.4f} h" if dry else "72 h 内未达标",
            va="center",
            fontsize=9,
        )
    axes[1].set_yticks([0, 1], ["末值", "末段均值"])
    axes[1].invert_yaxis()
    axes[1].set_xlim(
        0,
        max(
            (
                72
                if sensitivity["event"] is None
                else sensitivity["event"]["critical_s"] / 3600
            ),
            meta["event"]["critical_s"] / 3600,
        )
        * 1.3,
    )
    sensitivity_title = "附录 4 实测收缩的环境敏感性"
    if sensitivity["event"] is not None:
        change_min = (
            sensitivity["event"]["critical_s"] - meta["event"]["critical_s"]
        ) / 60
        sensitivity_title += f"\n均值延拓变化 {change_min:+.2f} min"
    finish_axes(axes[1], "临界时长 / h", "", sensitivity_title)
    save_figure(
        fig,
        "07_scenarios_and_sensitivity",
        ["review/verification.json", "cases/*/metadata.json"],
        "4 h 后边界为确定性假设；未达标情景仅报告覆盖内状态，不推断无限烘干时长；效应不普遍可加。",
        manifest,
    )


def workbook_previews(manifest: list) -> None:
    """Render saved workbook samples with its actual number formats for visual inspection."""
    book = openpyxl.load_workbook(OUT / "result4.xlsx", read_only=False, data_only=True)
    sheet = book["Sheet1"]
    row_ids = [1, *range(2, 8), 361, sheet.max_row - 1, sheet.max_row]
    row_ids = sorted(set(row for row in row_ids if row <= sheet.max_row))
    for number, columns in enumerate(
        (
            list(range(1, min(12, sheet.max_column + 1))),
            [1, *range(12, sheet.max_column + 1)],
        ),
        1,
    ):
        text = []
        for row in row_ids:
            row_text = []
            for col in columns:
                cell = sheet.cell(row, col)
                value = cell.value
                if value is None:
                    formatted = ""
                elif row == 1:
                    formatted = str(value).replace(
                        "时间\\到药材中心的距离", "时间/s\\距离/cm"
                    )
                elif col == 1:
                    formatted = f"{value:.8f}".rstrip("0").rstrip(".")
                else:
                    formatted = f"{value:.4f}"
                row_text.append(formatted)
            text.append(row_text)
        fig, ax = plt.subplots(figsize=(12, 3.9), layout="constrained")
        ax.axis("off")
        widths = [0.19] + [(1 - 0.19) / (len(columns) - 1)] * (len(columns) - 1)
        table = ax.table(cellText=text, cellLoc="right", colWidths=widths, loc="center")
        table.auto_set_font_size(False)
        table.set_fontsize(9)
        table.scale(1, 1.7)
        for (row, col), cell in table.get_celld().items():
            cell.set_edgecolor("#DBE0E3")
            cell.set_linewidth(0.35)
            if row == 0:
                cell.set_facecolor("#E8EDF0")
        ax.set_title(
            f"result4.xlsx · 选定行与第 {number} 组实际距离列", loc="left", fontsize=11
        )
        fig.canvas.draw()
        fig.savefig(
            PREVIEWS / f"workbook_part{number}.jpg",
            dpi=150,
            bbox_inches="tight",
            pil_kwargs={"quality": 92},
        )
        plt.close(fig)
    book.close()
    manifest.append(
        {
            "name": "workbook_preview",
            "sources": ["result4.xlsx"],
            "preview": [
                str(PREVIEWS / "workbook_part1.jpg"),
                str(PREVIEWS / "workbook_part2.jpg"),
            ],
            "scope_and_limit": "保存后工作簿选定行的数值与四位小数格式重构预览；数值逐格校验见 delivery_checks.json。",
        }
    )


def main() -> None:
    """Generate all required figures after validating the numerical delivery manifest."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--require-verified",
        action="store_true",
        help="Require verified data (also enforced by default)",
    )
    args = parser.parse_args()
    del args
    font = configure_style()
    report, meta, summary, env, radius, before = load_verified()
    IMAGES.mkdir(parents=True, exist_ok=True)
    PREVIEWS.mkdir(parents=True, exist_ok=True)
    figures = []
    plot_inputs(env, radius, meta, figures)
    plot_history(summary, meta, figures)
    plot_profiles(summary, meta, radius, figures)
    plot_heatmap(report, meta, summary, radius, figures)
    plot_temperature(summary, meta, figures)
    plot_convergence(report, meta, figures)
    plot_scenarios(report, meta, figures)
    workbook_previews(figures)
    for path, digest in before.items():
        if sha256(Path(path)) != digest:
            raise RuntimeError("Numerical data changed during plotting")
    manifest = {
        "status": "PASS",
        "automated_visual_checks": "PASS",
        "visual_review": "PENDING_JPEG_INSPECTION",
        "font": font,
        "plot_dependencies": {
            "matplotlib": matplotlib.__version__,
            "numpy": np.__version__,
            "openpyxl": openpyxl.__version__,
        },
        "plot_code_sha256": sha256(Path(__file__).resolve()),
        "solve_code_sha256": meta["identity"]["solve_code_sha256"],
        "verification_sha256": sha256(OUT / "review/verification.json"),
        "main_case": report["main"],
        "numerical_preservation": "PASS",
        "PNG_outputs_read": False,
        "figures": figures,
    }
    (IMAGES / "figure_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, allow_nan=False),
        encoding="utf-8",
    )
    print(
        json.dumps(
            {
                "status": "PASS",
                "scientific_figure_groups": 7,
                "font": font,
                "preview_directory": str(PREVIEWS),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
