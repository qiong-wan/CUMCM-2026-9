"""Generate publication figures for Problem 1.

The script reads the ambient observations and the existing result1.xlsx in
read-only mode.  Six figures are written to output/problem-1/image by default:

1. cylindrical model and boundary-condition schematic;
2. time-dependent drying-room boundary conditions;
3. temperature and moisture spatiotemporal fields;
4. radial profiles at representative times;
5. time histories at representative radii;
6. spatial-grid convergence.

Run from the Solution directory or any other working directory:

    python src/problem-1/generate_figures.py

Use ``--skip-convergence`` for a fast redraw of Figures 1--5.
"""

from __future__ import annotations

import argparse
from concurrent.futures import ProcessPoolExecutor
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.font_manager as font_manager
import matplotlib.pyplot as plt
from matplotlib.patches import Circle, Ellipse, FancyArrowPatch, Rectangle
from matplotlib.ticker import NullFormatter
import numpy as np
import openpyxl

from solve_problem1 import C_INIT, N_OUT, NSUB, T_INIT, simulate


ROOT = Path(__file__).resolve().parents[2]
AMBIENT_FILE = ROOT / "data" / "附件1.xlsx"
RESULT_FILE = ROOT / "output" / "problem-1" / "result1.xlsx"
DEFAULT_IMAGE_DIR = ROOT / "output" / "problem-1" / "image"

T_END = 1800
SELECTED_TIMES = (100, 600, 1200, 1800)
SELECTED_RADII_CM = (0.0, 1.0, 2.0)
CONVERGENCE_GRIDS = (100, 200, 400, 800, 1600)
CONVERGENCE_TARGET = 5.0e-5

TEMP_COLOR = "#C44E52"
MOISTURE_COLOR = "#277DA1"
ACCENT_COLOR = "#3A7D44"
NEUTRAL_COLOR = "#3D4652"
GRID_COLOR = "#D8DEE6"


def configure_matplotlib() -> str:
    """Configure a deterministic, Chinese-capable publication style."""
    installed = {item.name for item in font_manager.fontManager.ttflist}
    candidates = (
        "Microsoft YaHei",
        "SimHei",
        "Noto Sans CJK SC",
        "Source Han Sans SC",
        "Arial Unicode MS",
    )
    font_name = next((name for name in candidates if name in installed), "DejaVu Sans")

    plt.rcParams.update(
        {
            "font.family": "sans-serif",
            "font.sans-serif": [font_name, "DejaVu Sans"],
            "axes.unicode_minus": False,
            "mathtext.fontset": "stix",
            "font.size": 10,
            "axes.titlesize": 11,
            "axes.labelsize": 10,
            "legend.fontsize": 9,
            "xtick.labelsize": 9,
            "ytick.labelsize": 9,
            "axes.edgecolor": NEUTRAL_COLOR,
            "axes.labelcolor": "#20242A",
            "text.color": "#20242A",
            "xtick.color": "#343A40",
            "ytick.color": "#343A40",
            "savefig.facecolor": "white",
        }
    )
    return font_name


def load_ambient() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Read and validate ambient observations for 0--1800 s."""
    workbook = openpyxl.load_workbook(
        AMBIENT_FILE,
        read_only=True,
        data_only=True,
    )
    try:
        worksheet = workbook.active
        rows = [
            row[:3]
            for row in worksheet.iter_rows(min_row=2, values_only=True)
            if row[0] is not None and float(row[0]) <= T_END
        ]
    finally:
        workbook.close()

    if not rows or any(value is None for row in rows for value in row):
        raise ValueError("附件1在0--1800 s范围内存在空值。")

    data = np.asarray(rows, dtype=float)
    times, temperatures, moistures = data.T
    if times[0] != 0 or times[-1] != T_END:
        raise ValueError("附件1必须完整覆盖0--1800 s。")
    if not np.all(np.diff(times) == 60):
        raise ValueError("附件1在问题1区间内必须采用60 s等间隔记录。")
    if not np.all(np.isfinite(data)):
        raise ValueError("附件1包含非有限数值。")
    return times, temperatures, moistures


def _read_result_sheet(
    workbook: openpyxl.Workbook,
    sheet_name: str,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Read one result sheet and return time, radius, and field arrays."""
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"result1.xlsx缺少工作表：{sheet_name}")

    rows = list(workbook[sheet_name].iter_rows(values_only=True))
    if len(rows) < 2 or len(rows[0]) < 2:
        raise ValueError(f"工作表{sheet_name}没有完整结果。")

    radii_cm = np.asarray(rows[0][1:], dtype=float)
    times = np.asarray([row[0] for row in rows[1:]], dtype=float)
    field = np.asarray([row[1:] for row in rows[1:]], dtype=float)

    if field.shape != (times.size, radii_cm.size):
        raise ValueError(f"工作表{sheet_name}的时间列与数据区域形状不一致。")
    if not np.all(np.isfinite(field)):
        raise ValueError(f"工作表{sheet_name}包含空值或非有限数值。")
    if radii_cm.size != 21 or not np.allclose(radii_cm, np.arange(21) / 10):
        raise ValueError(f"工作表{sheet_name}的半径网格必须为0--2 cm、步长0.1 cm。")
    if not np.allclose(np.diff(times), 1.0):
        raise ValueError(f"工作表{sheet_name}的时间步必须为1 s。")
    if times[-1] != T_END or times[0] not in (0, 1):
        raise ValueError(f"工作表{sheet_name}必须覆盖至1800 s，并从0或1 s开始。")

    return times, radii_cm, field


def load_results() -> tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Read temperature and moisture results, adding the known t=0 state if needed."""
    if not RESULT_FILE.exists():
        raise FileNotFoundError(
            f"未找到{RESULT_FILE}，请先运行solve_problem1.py生成结果。"
        )

    workbook = openpyxl.load_workbook(
        RESULT_FILE,
        read_only=True,
        data_only=True,
    )
    try:
        t_temp, radii_temp, temperature = _read_result_sheet(workbook, "温度")
        t_moist, radii_moist, moisture = _read_result_sheet(workbook, "水分浓度")
    finally:
        workbook.close()

    if not np.array_equal(t_temp, t_moist):
        raise ValueError("温度和水分浓度工作表的时间列不一致。")
    if not np.array_equal(radii_temp, radii_moist):
        raise ValueError("温度和水分浓度工作表的半径列不一致。")

    times = t_temp
    if times[0] == 1:
        times = np.insert(times, 0, 0.0)
        temperature = np.vstack(
            [np.full((1, radii_temp.size), T_INIT), temperature]
        )
        moisture = np.vstack(
            [np.full((1, radii_temp.size), C_INIT), moisture]
        )

    return times, radii_temp, temperature, moisture


def style_axis(axis: plt.Axes, grid_axis: str = "both") -> None:
    """Apply restrained scientific axis styling."""
    axis.spines["top"].set_visible(False)
    axis.spines["right"].set_visible(False)
    axis.grid(
        True,
        axis=grid_axis,
        color=GRID_COLOR,
        linewidth=0.65,
        alpha=0.8,
    )
    axis.set_axisbelow(True)


def save_figure(figure: plt.Figure, output_path: Path, dpi: int) -> Path:
    """Save one figure with consistent publication settings."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output_path, dpi=dpi, bbox_inches="tight", pad_inches=0.08)
    plt.close(figure)
    return output_path


def plot_model_schematic(output_dir: Path, dpi: int) -> Path:
    """Plot cylinder geometry and radial heat/mass boundary conditions."""
    figure, axes = plt.subplots(1, 2, figsize=(11.2, 4.1), layout="constrained")
    side_axis, radial_axis = axes
    figure.suptitle("圆柱药材的一维径向传热传质模型", fontsize=13)

    body = Rectangle(
        (0, -2),
        25,
        4,
        facecolor="#E9ECEF",
        edgecolor=NEUTRAL_COLOR,
        linewidth=1.4,
    )
    left_end = Ellipse(
        (0, 0),
        1.5,
        4,
        facecolor="#DEE2E6",
        edgecolor=NEUTRAL_COLOR,
        linewidth=1.2,
    )
    right_end = Ellipse(
        (25, 0),
        1.5,
        4,
        facecolor="#F1F3F5",
        edgecolor=NEUTRAL_COLOR,
        linewidth=1.2,
    )
    side_axis.add_patch(body)
    side_axis.add_patch(left_end)
    side_axis.add_patch(right_end)
    side_axis.annotate(
        "",
        xy=(25, -3.1),
        xytext=(0, -3.1),
        arrowprops={"arrowstyle": "<->", "color": NEUTRAL_COLOR, "lw": 1.1},
    )
    side_axis.text(12.5, -3.45, "$L=25$ cm", ha="center", va="top")
    side_axis.annotate(
        "",
        xy=(25.9, 2),
        xytext=(25.9, 0),
        arrowprops={"arrowstyle": "<->", "color": NEUTRAL_COLOR, "lw": 1.1},
    )
    side_axis.text(26.25, 1, "$R=2$ cm", rotation=90, ha="left", va="center")
    side_axis.plot([-0.7, 25.7], [0, 0], color="#6C757D", lw=0.8, ls="--")
    side_axis.text(12.5, 0.25, "轴向均匀", ha="center", va="bottom")
    side_axis.set_xlim(-1.8, 27.8)
    side_axis.set_ylim(-4.0, 3.1)
    side_axis.set_aspect("equal")
    side_axis.set_title("(a) 圆柱几何")
    side_axis.axis("off")

    radial_axis.add_patch(
        Circle(
            (0, 0),
            2,
            facecolor="#EEF1F4",
            edgecolor=NEUTRAL_COLOR,
            linewidth=1.5,
        )
    )
    radial_axis.add_patch(
        Circle(
            (0, 0),
            1,
            facecolor="none",
            edgecolor="#ADB5BD",
            linewidth=0.8,
            linestyle="--",
        )
    )
    radial_axis.scatter([0], [0], s=20, color=NEUTRAL_COLOR, zorder=5)
    radial_axis.text(-0.18, -0.3, "$r=0$", ha="right", va="top")
    radial_axis.annotate(
        "",
        xy=(1.95, 0),
        xytext=(0.05, 0),
        arrowprops={"arrowstyle": "->", "color": NEUTRAL_COLOR, "lw": 1.4},
    )
    radial_axis.text(1.0, 0.16, "$r$", ha="center", va="bottom")
    radial_axis.text(0, -0.65, r"$T(r,t),\ C(r,t)$", ha="center", va="center")

    heat_arrow = FancyArrowPatch(
        (3.0, 0.75),
        (1.95, 0.5),
        arrowstyle="-|>",
        mutation_scale=13,
        color=TEMP_COLOR,
        linewidth=1.8,
    )
    mass_arrow = FancyArrowPatch(
        (1.95, -0.5),
        (3.0, -0.75),
        arrowstyle="-|>",
        mutation_scale=13,
        color=MOISTURE_COLOR,
        linewidth=1.8,
    )
    radial_axis.add_patch(heat_arrow)
    radial_axis.add_patch(mass_arrow)
    radial_axis.text(3.08, 0.82, "对流供热  $h$", color=TEMP_COLOR, ha="left")
    radial_axis.text(3.08, -0.9, "对流传质  $h_m$", color=MOISTURE_COLOR, ha="left")
    radial_axis.text(
        0,
        2.55,
        r"$T_\infty(t),\ C_\infty(t)$",
        ha="center",
        va="bottom",
    )
    radial_axis.set_xlim(-3.3, 4.8)
    radial_axis.set_ylim(-3.0, 3.0)
    radial_axis.set_aspect("equal")
    radial_axis.set_title("(b) 径向边界条件")
    radial_axis.axis("off")

    return save_figure(figure, output_dir / "fig1_model_schematic.png", dpi)


def plot_ambient_conditions(
    ambient: tuple[np.ndarray, np.ndarray, np.ndarray],
    output_dir: Path,
    dpi: int,
) -> Path:
    """Plot measured and linearly interpolated drying-room conditions."""
    times, temperatures, moistures = ambient
    dense_times = np.arange(0, T_END + 1, dtype=float)
    dense_temperature = np.interp(dense_times, times, temperatures)
    dense_moisture = np.interp(dense_times, times, moistures)

    figure, axes = plt.subplots(
        2,
        1,
        figsize=(9.2, 6.3),
        sharex=True,
        layout="constrained",
    )
    figure.suptitle("预热平衡阶段的烘房边界条件", fontsize=13)

    axes[0].plot(
        dense_times,
        dense_temperature,
        color=TEMP_COLOR,
        linewidth=1.8,
        label="分段线性插值",
    )
    axes[0].scatter(
        times,
        temperatures,
        s=17,
        color=NEUTRAL_COLOR,
        zorder=3,
        label="附件1观测值",
    )
    axes[0].set_ylabel("烘房温度 / °C")
    axes[0].set_title("(a) 温度边界")
    axes[0].legend(frameon=False, ncol=2, loc="upper left")
    style_axis(axes[0])

    axes[1].plot(
        dense_times,
        dense_moisture,
        color=MOISTURE_COLOR,
        linewidth=1.8,
        label="分段线性插值",
    )
    axes[1].scatter(
        times,
        moistures,
        s=17,
        color=NEUTRAL_COLOR,
        zorder=3,
        label="附件1观测值",
    )
    axes[1].set_xlabel("时间 / s")
    axes[1].set_ylabel("烘房水分浓度 / (kg/kg)")
    axes[1].set_title("(b) 水分浓度边界")
    style_axis(axes[1])
    axes[1].set_xticks(np.arange(0, T_END + 1, 300))

    return save_figure(figure, output_dir / "fig2_ambient_conditions.png", dpi)


def plot_spatiotemporal_fields(
    times: np.ndarray,
    radii_cm: np.ndarray,
    temperature: np.ndarray,
    moisture: np.ndarray,
    output_dir: Path,
    dpi: int,
) -> Path:
    """Plot temperature and moisture fields as aligned contour maps."""
    figure, axes = plt.subplots(1, 2, figsize=(12.6, 4.8), layout="constrained")
    figure.suptitle("30 min 内药材温度与含水率的时空分布", fontsize=13)

    temp_levels = np.linspace(float(temperature.min()), float(temperature.max()), 60)
    temp_map = axes[0].contourf(
        times,
        radii_cm,
        temperature.T,
        levels=temp_levels,
        cmap="plasma",
        extend="both",
    )
    axes[0].contour(
        times,
        radii_cm,
        temperature.T,
        levels=np.linspace(temperature.min(), temperature.max(), 7)[1:-1],
        colors="white",
        linewidths=0.45,
        alpha=0.72,
    )
    axes[0].set_title("(a) 温度场")
    axes[0].set_xlabel("时间 / s")
    axes[0].set_ylabel("到中心的距离 / cm")
    axes[0].set_xticks(np.arange(0, T_END + 1, 300))
    axes[0].set_yticks(np.arange(0, 2.01, 0.5))
    temp_bar = figure.colorbar(temp_map, ax=axes[0], pad=0.02)
    temp_bar.set_label("温度 / °C")

    moisture_levels = np.linspace(float(moisture.min()), float(moisture.max()), 60)
    moisture_map = axes[1].contourf(
        times,
        radii_cm,
        moisture.T,
        levels=moisture_levels,
        cmap="viridis",
        extend="both",
    )
    axes[1].contour(
        times,
        radii_cm,
        moisture.T,
        levels=np.linspace(moisture.min(), moisture.max(), 7)[1:-1],
        colors="white",
        linewidths=0.45,
        alpha=0.72,
    )
    axes[1].set_title("(b) 干基含水率场")
    axes[1].set_xlabel("时间 / s")
    axes[1].set_ylabel("到中心的距离 / cm")
    axes[1].set_xticks(np.arange(0, T_END + 1, 300))
    axes[1].set_yticks(np.arange(0, 2.01, 0.5))
    moisture_bar = figure.colorbar(moisture_map, ax=axes[1], pad=0.02)
    moisture_bar.set_label("干基含水率 / (kg/kg)")

    return save_figure(figure, output_dir / "fig3_spatiotemporal_fields.png", dpi)


def plot_radial_profiles(
    times: np.ndarray,
    radii_cm: np.ndarray,
    temperature: np.ndarray,
    moisture: np.ndarray,
    output_dir: Path,
    dpi: int,
) -> Path:
    """Plot radial profiles at four representative times."""
    figure, axes = plt.subplots(1, 2, figsize=(11.2, 4.5), layout="constrained")
    figure.suptitle("典型时刻的径向温度与含水率分布", fontsize=13)
    colors = plt.get_cmap("cividis")(np.linspace(0.12, 0.88, len(SELECTED_TIMES)))

    for selected_time, color in zip(SELECTED_TIMES, colors, strict=True):
        time_index = int(np.flatnonzero(times == selected_time)[0])
        label = f"{selected_time} s"
        axes[0].plot(
            radii_cm,
            temperature[time_index],
            color=color,
            linewidth=1.8,
            marker="o",
            markersize=3.2,
            markevery=2,
            label=label,
        )
        axes[1].plot(
            radii_cm,
            moisture[time_index],
            color=color,
            linewidth=1.8,
            marker="o",
            markersize=3.2,
            markevery=2,
            label=label,
        )

    axes[0].set_title("(a) 温度径向剖面")
    axes[0].set_xlabel("到中心的距离 / cm")
    axes[0].set_ylabel("温度 / °C")
    style_axis(axes[0])

    axes[1].set_title("(b) 干基含水率径向剖面")
    axes[1].set_xlabel("到中心的距离 / cm")
    axes[1].set_ylabel("干基含水率 / (kg/kg)")
    style_axis(axes[1])

    for axis in axes:
        axis.set_xlim(0, 2)
        axis.set_xticks(np.arange(0, 2.01, 0.5))

    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(
        handles,
        labels,
        loc="outside lower center",
        ncol=len(SELECTED_TIMES),
        frameon=False,
    )

    return save_figure(figure, output_dir / "fig4_radial_profiles.png", dpi)


def plot_time_histories(
    times: np.ndarray,
    radii_cm: np.ndarray,
    temperature: np.ndarray,
    moisture: np.ndarray,
    output_dir: Path,
    dpi: int,
) -> Path:
    """Plot histories at five representative radii."""
    figure, axes = plt.subplots(1, 2, figsize=(11.2, 4.5), layout="constrained")
    figure.suptitle("典型位置的温度与含水率时间历程", fontsize=13)

    selected_radii_cm = (0.0, 0.5, 1.0, 1.5, 2.0)
    colors = ("#355070", "#6D8AAE", ACCENT_COLOR, "#E08A5B", "#D1495B")
    markers = ("s", "o", "^", "D", "v")

    for radius, color, marker in zip(
        selected_radii_cm,
        colors,
        markers,
        strict=True,
    ):
        radius_index = int(np.flatnonzero(np.isclose(radii_cm, radius))[0])
        label = f"r = {radius:g} cm"
        axes[0].plot(
            times,
            temperature[:, radius_index],
            color=color,
            linewidth=1.8,
            marker=marker,
            markersize=3.5,
            markevery=300,
            label=label,
        )
        axes[1].plot(
            times,
            moisture[:, radius_index],
            color=color,
            linewidth=1.8,
            marker=marker,
            markersize=3.5,
            markevery=300,
            label=label,
        )

    axes[0].set_title("(a) 温度响应")
    axes[0].set_xlabel("时间 / s")
    axes[0].set_ylabel("温度 / °C")
    style_axis(axes[0])

    axes[1].set_title("(b) 干基含水率响应")
    axes[1].set_xlabel("时间 / s")
    axes[1].set_ylabel("干基含水率 / (kg/kg)")
    style_axis(axes[1])

    for axis in axes:
        axis.set_xlim(0, T_END)
        axis.set_xticks(np.arange(0, T_END + 1, 300))

    handles, labels = axes[0].get_legend_handles_labels()
    figure.legend(
        handles,
        labels,
        loc="outside lower center",
        ncol=len(selected_radii_cm),
        frameon=False,
    )

    return save_figure(figure, output_dir / "fig5_time_histories.png", dpi)


def _simulate_grid(grid_size: int) -> tuple[int, np.ndarray, np.ndarray]:
    """Return fields sampled on the common 0.1 cm output grid."""
    if grid_size % N_OUT != 0:
        raise ValueError(f"网格数{grid_size}不能被输出网格数{N_OUT}整除。")
    _, temperature, moisture, _ = simulate(
        N=grid_size,
        nsub=NSUB,
        sample_step=grid_size // N_OUT,
    )
    return grid_size, temperature, moisture


def calculate_convergence(
    grid_sizes: tuple[int, ...],
    workers: int,
) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Calculate successive-grid maximum field differences."""
    if len(grid_sizes) < 2:
        raise ValueError("网格收敛分析至少需要两个网格。")

    if workers == 1:
        results = [_simulate_grid(grid_size) for grid_size in grid_sizes]
    else:
        with ProcessPoolExecutor(max_workers=workers) as executor:
            results = list(executor.map(_simulate_grid, grid_sizes))

    results.sort(key=lambda item: item[0])
    fine_grids: list[int] = []
    temperature_errors: list[float] = []
    moisture_errors: list[float] = []

    for previous, current in zip(results[:-1], results[1:], strict=True):
        previous_grid, previous_temperature, previous_moisture = previous
        current_grid, current_temperature, current_moisture = current
        if current_grid <= previous_grid:
            raise ValueError("空间网格必须严格递增。")
        fine_grids.append(current_grid)
        temperature_errors.append(
            float(np.max(np.abs(current_temperature - previous_temperature)))
        )
        moisture_errors.append(
            float(np.max(np.abs(current_moisture - previous_moisture)))
        )

    return (
        np.asarray(fine_grids),
        np.asarray(temperature_errors),
        np.asarray(moisture_errors),
    )


def plot_convergence(
    convergence: tuple[np.ndarray, np.ndarray, np.ndarray],
    output_dir: Path,
    dpi: int,
) -> Path:
    """Plot successive-grid maximum differences on logarithmic scales."""
    fine_grids, temperature_errors, moisture_errors = convergence
    figure, axes = plt.subplots(1, 2, figsize=(11.2, 4.5), layout="constrained")
    figure.suptitle("空间网格收敛性", fontsize=13)

    panels = (
        (axes[0], temperature_errors, TEMP_COLOR, "温度", "最大差 / °C"),
        (
            axes[1],
            moisture_errors,
            MOISTURE_COLOR,
            "干基含水率",
            "最大差 / (kg/kg)",
        ),
    )
    for axis, errors, color, title, ylabel in panels:
        axis.loglog(
            fine_grids,
            errors,
            color=color,
            linewidth=1.8,
            marker="o",
            markersize=5,
        )
        axis.axhline(
            CONVERGENCE_TARGET,
            color="#6C757D",
            linewidth=1.0,
            linestyle="--",
            label="四位小数参考阈值",
        )
        for grid_size, error in zip(fine_grids, errors, strict=True):
            axis.annotate(
                f"{error:.1e}",
                xy=(grid_size, error),
                xytext=(0, 7),
                textcoords="offset points",
                ha="center",
                fontsize=8,
            )
        axis.set_title(title)
        axis.set_xlabel("较细网格的单元数 $N$")
        axis.set_ylabel(ylabel)
        axis.set_xticks(fine_grids)
        axis.set_xticklabels([str(value) for value in fine_grids])
        axis.xaxis.set_minor_formatter(NullFormatter())
        axis.grid(True, which="both", color=GRID_COLOR, linewidth=0.65)
        axis.legend(frameon=False, loc="upper right")
        axis.spines["top"].set_visible(False)
        axis.spines["right"].set_visible(False)

    return save_figure(figure, output_dir / "fig6_grid_convergence.png", dpi)


def generate_figures(
    output_dir: Path,
    dpi: int,
    include_convergence: bool,
    workers: int,
) -> list[Path]:
    """Generate the complete figure set and return written paths."""
    if dpi < 150:
        raise ValueError("论文图像DPI不应低于150。")
    if workers < 1:
        raise ValueError("workers必须至少为1。")

    font_name = configure_matplotlib()
    ambient = load_ambient()
    times, radii_cm, temperature, moisture = load_results()

    outputs = [
        plot_model_schematic(output_dir, dpi),
        plot_ambient_conditions(ambient, output_dir, dpi),
        plot_spatiotemporal_fields(
            times,
            radii_cm,
            temperature,
            moisture,
            output_dir,
            dpi,
        ),
        plot_radial_profiles(
            times,
            radii_cm,
            temperature,
            moisture,
            output_dir,
            dpi,
        ),
        plot_time_histories(
            times,
            radii_cm,
            temperature,
            moisture,
            output_dir,
            dpi,
        ),
    ]

    if include_convergence:
        convergence = calculate_convergence(CONVERGENCE_GRIDS, workers)
        outputs.append(plot_convergence(convergence, output_dir, dpi))

    print(f"Chinese font: {font_name}")
    for output in outputs:
        print(output)
    return outputs


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments."""
    parser = argparse.ArgumentParser(description="Generate Problem 1 paper figures.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=DEFAULT_IMAGE_DIR,
        help="figure output directory",
    )
    parser.add_argument(
        "--dpi",
        type=int,
        default=300,
        help="PNG resolution; default 300",
    )
    parser.add_argument(
        "--workers",
        type=int,
        default=1,
        help="parallel processes for convergence simulations",
    )
    parser.add_argument(
        "--skip-convergence",
        action="store_true",
        help="generate Figures 1--5 without the expensive convergence study",
    )
    return parser.parse_args()


def main() -> None:
    """Command-line entry point."""
    args = parse_args()
    generate_figures(
        output_dir=args.output_dir.resolve(),
        dpi=args.dpi,
        include_convergence=not args.skip_convergence,
        workers=args.workers,
    )


if __name__ == "__main__":
    main()
