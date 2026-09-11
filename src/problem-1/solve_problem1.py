"""Problem 1: radial temperature and moisture concentration of a cylindrical
medicinal material during the preheating/equilibrium stage (0--1800 s).

Model (1D axisymmetric cylinder, radius R, constant length L)
------------------------------------------------------------
Temperature (Fourier, constant properties from Appendix 2):

    rho*cp*dT/dt = (1/r) d/dr ( k r dT/dr ),   0 < r < R,
    dT/dr(0,t) = 0,
    -k dT/dr(R,t) = h [T(R,t) - T_inf(t)],
    T(r,0) = 28 degC.

Moisture (Fick, concentration-dependent diffusivity):

    dC/dt = (1/r) d/dr ( r D(C) dC/dr ),   0 < r < R,
    dC/dr(0,t) = 0,
    -D(C) dC/dr(R,t) = h_m [C(R,t) - C_inf(t)],
    C(r,0) = 2.55 kg/kg,
    D(C) = 7e-9 * exp(-0.89 / C)  [m^2/s].

The ambient temperature T_inf(t) and moisture concentration C_inf(t) are read
from data/附件1.xlsx (every 60 s) and linearly interpolated in time.

Numerical method
----------------
Node-centered finite-volume discretisation on the uniform grid
r_i = i*dr (i = 0..N, dr = R/N).  Backward-Euler in time with sub-step
dt = 1/nsub s and integer-second sampling.  The resulting tridiagonal systems
are solved with scipy.linalg.solve_banded; the nonlinear moisture system is
handled by Picard iteration on the face diffusivities (harmonic mean of D(C)).
See src/problem-1/问题1_解题框架与公式推导.md.

The production run uses a refined mesh (N_REF cells) and samples the fields at
the required 0.1 cm output nodes, which removes the near-surface discretisation
bias; a coarse mesh is retained for the convergence check.

Run:
    python src/problem-1/solve_problem1.py            # results + tables
    python src/problem-1/solve_problem1.py --verify   # also convergence study
"""

from __future__ import annotations

import argparse
from pathlib import Path

import numpy as np
import openpyxl
from scipy.linalg import solve_banded

# --------------------------------------------------------------------------
# Paths
# --------------------------------------------------------------------------
ROOT = Path(__file__).resolve().parents[2]
DATA_FILE = ROOT / "data" / "附件1.xlsx"
OUT_DIR = ROOT / "output" / "problem-1"
RESULT_FILE = OUT_DIR / "result1.xlsx"

# --------------------------------------------------------------------------
# Physical parameters (Appendix 2)
# --------------------------------------------------------------------------
R = 0.02            # cylinder radius [m]
L = 0.25            # cylinder length [m] (unused in the 1D radial model)
RHO = 820.0         # density [kg/m^3]
CP = 2600.0         # specific heat [J/(kg K)]
K = 0.36            # thermal conductivity [W/(m K)]
H = 25.0            # convective heat transfer coefficient [W/(m^2 K)]
HM = 8.0e-7         # convective mass transfer coefficient [m/s]
ALPHA = K / (RHO * CP)

T_INIT = 28.0       # initial temperature [degC]
C_INIT = 2.55       # initial dry-basis moisture [kg/kg]

T_END = 1800.0      # simulation horizon [s]
N_OUT = 20          # output cells -> 0.1 cm radial spacing
N_REF = 800         # production mesh (0.0025 cm cells, sampled onto output grid)
NSUB = 40           # implicit sub-steps per output second (dt = 0.025 s)

TABLE_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800])
TABLE_DIST_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])


def D_of_C(C: np.ndarray) -> np.ndarray:
    """Moisture diffusivity D(C) = 7e-9 exp(-0.89/C) [m^2/s]."""
    C_safe = np.maximum(C, 1.0e-8)
    return 7.0e-9 * np.exp(-0.89 / C_safe)


def harmonic_mean(a: np.ndarray, b: np.ndarray) -> np.ndarray:
    """Harmonic mean 2ab/(a+b), element-wise, safe when a+b == 0."""
    denom = a + b
    out = np.zeros_like(denom)
    nz = denom > 0.0
    out[nz] = 2.0 * a[nz] * b[nz] / denom[nz]
    return out


def load_ambient() -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """Return (t, T_inf, C_inf) observations for 0 <= t <= 1800 s."""
    wb = openpyxl.load_workbook(DATA_FILE, data_only=True)
    ws = wb.active
    rows = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None]
    t = np.array([float(r[0]) for r in rows])
    t_inf = np.array([float(r[1]) for r in rows])
    c_inf = np.array([float(r[2]) for r in rows])
    keep = t <= T_END + 1e-9
    return t[keep], t_inf[keep], c_inf[keep]


def build_geometry(N: int):
    """Build finite-volume geometry on r_i = i*dr (areas scaled by 1/(pi*L))."""
    dr = R / N
    r = np.arange(N + 1) * dr
    vol = np.empty(N + 1)
    vol[0] = (dr / 2.0) ** 2
    vol[1:N] = (r[1:N] + dr / 2.0) ** 2 - (r[1:N] - dr / 2.0) ** 2
    vol[N] = R**2 - (R - dr / 2.0) ** 2
    # east face area of node i (between i and i+1), i = 0..N-1
    ge = 2.0 * (r[:N] + dr / 2.0)
    # west face area of node i (between i-1 and i), i = 1..N (index 0 unused)
    gw = np.empty(N + 1)
    gw[1:] = 2.0 * (r[1:] - dr / 2.0)
    a_r = 2.0 * R  # outer surface area scaled by 1/(pi*L)
    return dr, r, vol, ge, gw, a_r


def build_tridiag(vol, ge, gw, a_r, N, dr, dt, s, h, ge_cond, gw_cond):
    """Assemble lower/diagonal/upper of the implicit finite-volume system."""
    lower = np.zeros(N + 1)
    diag = s * vol / dt
    upper = np.zeros(N + 1)

    diag[:N] += ge * ge_cond / dr
    diag[1:] += gw[1:] * gw_cond[1:] / dr
    diag[N] += a_r * h

    lower[1:] = -gw[1:] * gw_cond[1:] / dr
    upper[:N] = -ge * ge_cond / dr
    return lower, diag, upper


def to_banded(lower, diag, upper):
    """Pack a tridiagonal system into scipy banded storage (1 upper, 1 lower)."""
    n = diag.size
    ab = np.zeros((3, n))
    ab[0, 1:] = upper[:-1]
    ab[1, :] = diag
    ab[2, :-1] = lower[1:]
    return ab


def solve_tridiag(lower, diag, upper, rhs):
    """Solve a tridiagonal system via scipy's banded solver."""
    return solve_banded((1, 1), to_banded(lower, diag, upper), rhs)


def build_rhs(vol, u_old, s, dt, a_r, h, u_inf):
    """Right-hand side of the implicit finite-volume system."""
    rhs = s * vol / dt * u_old
    rhs[-1] += a_r * h * u_inf
    return rhs


def simulate(N: int = N_REF, nsub: int = NSUB, sample_step: int | None = None):
    """Run the heat/moisture simulation.

    Parameters
    ----------
    N : number of radial cells (grid nodes = N+1).
    nsub : implicit sub-steps per 1 s output interval.
    sample_step : radial stride used when returning the fields; default keeps
        the full fine mesh.  Use N // N_OUT to sample onto the 0.1 cm grid.

    Returns
    -------
    r : node radii [m] (sampled).
    T_hist, C_hist : (1801, n_sample) arrays sampled at integer seconds.
    """
    dr, r, vol, ge, gw, a_r = build_geometry(N)
    t_obs, t_inf_obs, c_inf_obs = load_ambient()

    dt = 1.0 / nsub
    nsec = int(round(T_END))
    nsteps = nsec * nsub

    T = np.full(N + 1, T_INIT)
    C = np.full(N + 1, C_INIT)

    if sample_step is None:
        sample_step = 1
    idx = np.arange(0, N + 1, sample_step)
    T_hist = np.empty((nsec + 1, idx.size))
    C_hist = np.empty((nsec + 1, idx.size))
    T_hist[0] = T[idx]
    C_hist[0] = C[idx]

    # --- temperature: constant tridiagonal matrix ---
    s_t = RHO * CP
    k_e = np.full(N, K)
    k_w = np.full(N + 1, K)
    low_t, diag_t, up_t = build_tridiag(vol, ge, gw, a_r, N, dr, dt, s_t, H, k_e, k_w)
    ab_t = to_banded(low_t, diag_t, up_t)

    # --- moisture: nonlinear, Picard iteration each sub-step ---
    s_c = 1.0
    tol = 1.0e-12
    max_it = 60

    for sec in range(1, nsec + 1):
        for sub in range(nsub):
            t_new = ((sec - 1) * nsub + sub + 1) * dt
            t_inf = np.interp(t_new, t_obs, t_inf_obs)
            c_inf = np.interp(t_new, t_obs, c_inf_obs)

            # temperature update (linear)
            rhs = build_rhs(vol, T, s_t, dt, a_r, H, t_inf)
            T = solve_banded((1, 1), ab_t, rhs)

            # moisture update (Picard on D(C))
            rhs_c = build_rhs(vol, C, s_c, dt, a_r, HM, c_inf)
            C_new = C.copy()
            for _ in range(max_it):
                D = D_of_C(C_new)
                ge_cond = harmonic_mean(D[:N], D[1:])
                gw_cond = np.empty(N + 1)
                gw_cond[1:] = ge_cond
                low_c, diag_c, up_c = build_tridiag(
                    vol, ge, gw, a_r, N, dr, dt, s_c, HM, ge_cond, gw_cond
                )
                C_sol = solve_tridiag(low_c, diag_c, up_c, rhs_c)
                if np.max(np.abs(C_sol - C_new)) < tol:
                    C_new = C_sol
                    break
                C_new = C_sol
            C = C_new

        T_hist[sec] = T[idx]
        C_hist[sec] = C[idx]

    return r[idx], T_hist, C_hist


def check_balances(r, T_hist, C_hist, N):
    """Verify global energy and moisture balances over the whole run."""
    dr, _, vol, _, _, a_r = build_geometry(N)
    t_obs, t_inf_obs, c_inf_obs = load_ambient()

    energy0 = np.sum(RHO * CP * vol * T_hist[0])
    energy1 = np.sum(RHO * CP * vol * T_hist[-1])
    moist0 = np.sum(vol * C_hist[0])
    moist1 = np.sum(vol * C_hist[-1])

    t = np.arange(T_hist.shape[0], dtype=float)
    t_inf = np.interp(t, t_obs, t_inf_obs)
    c_inf = np.interp(t, t_obs, c_inf_obs)
    q_heat = a_r * H * (t_inf - T_hist[:, -1])
    q_mass = a_r * HM * (c_inf - C_hist[:, -1])

    heat_in = np.sum(0.5 * (q_heat[1:] + q_heat[:-1]) * np.diff(t))
    mass_in = np.sum(0.5 * (q_mass[1:] + q_mass[:-1]) * np.diff(t))

    e_err = abs((energy1 - energy0) - heat_in) / max(abs(energy1 - energy0), 1e-30)
    m_err = abs((moist1 - moist0) - mass_in) / max(abs(moist1 - moist0), 1e-30)
    return e_err, m_err


def extract_tables(r, T_hist, C_hist):
    """Extract Table 1 (temperature) and Table 2 (moisture) for the paper."""
    t_idx = TABLE_TIMES.astype(int)
    dr = r[1] - r[0]
    r_idx = np.round(TABLE_DIST_CM / 100.0 / dr).astype(int)
    tab_t = T_hist[t_idx][:, r_idx]
    tab_c = C_hist[t_idx][:, r_idx]
    return tab_t, tab_c


def save_result(r, T_hist, C_hist) -> None:
    """Write the full temperature and moisture fields to result1.xlsx."""
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    dist_cm = np.round(r * 100.0, 10)
    times = np.arange(T_hist.shape[0], dtype=float)

    wb = openpyxl.Workbook()
    for sheet_name, data in (("温度", T_hist), ("水分浓度", C_hist)):
        ws = wb.create_sheet(sheet_name)
        ws.cell(row=1, column=1, value="时间/s")
        for j, d in enumerate(dist_cm, start=2):
            ws.cell(row=1, column=j, value=float(d))
        for i, t in enumerate(times, start=2):
            ws.cell(row=i, column=1, value=float(t))
            for j in range(data.shape[1]):
                ws.cell(row=i, column=j + 2, value=round(float(data[i - 2, j]), 4))
    del wb["Sheet"]
    wb.save(RESULT_FILE)


def write_tables(tab_t, tab_c) -> None:
    """Save Table 1 and Table 2 as CSV files (for the paper)."""
    for fname, tab in (("table1_temperature.csv", tab_t), ("table2_moisture.csv", tab_c)):
        lines = ["时间/s," + ",".join(f"{d:g}" for d in TABLE_DIST_CM)]
        for i, t in enumerate(TABLE_TIMES):
            lines.append(f"{t}," + ",".join(f"{v:.4f}" for v in tab[i]))
        (OUT_DIR / fname).write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def run_verification(T_prod, C_prod) -> None:
    """Report global balances and grid/time convergence of the solution."""
    # global conservation on the required 0.1 cm mesh
    r0, T0, C0 = simulate(N=N_OUT, nsub=NSUB, sample_step=1)
    e_err, m_err = check_balances(r0, T0, C0, N=N_OUT)
    print(f"\nGlobal energy balance relative error  : {e_err:.3e}")
    print(f"Global moisture balance relative error : {m_err:.3e}")

    # spatial convergence at fixed dt (coarse 0.1 cm mesh as reference)
    print("\nSpatial convergence (sampled on the 0.1 cm grid):")
    prev = (T0, C0)
    for N in (40, 100, 200, 400):
        _, TT, CC = simulate(N=N, nsub=NSUB, sample_step=N // N_OUT)
        d_t = np.max(np.abs(TT - prev[0]))
        d_c = np.max(np.abs(CC - prev[1]))
        print(f"  N={N:4d} vs previous: dT = {d_t:.3e} degC, dC = {d_c:.3e} kg/kg")
        prev = (TT, CC)
    d_t = np.max(np.abs(T_prod - prev[0]))
    d_c = np.max(np.abs(C_prod - prev[1]))
    print(f"  N={N_REF:4d} vs previous: dT = {d_t:.3e} degC, dC = {d_c:.3e} kg/kg")

    # temporal convergence at fixed mesh
    print(f"\nTemporal convergence (N={N_REF}, sampled on the 0.1 cm grid):")
    prev = (T_prod, C_prod)
    for nsub in (80, 160):
        _, TT, CC = simulate(N=N_REF, nsub=nsub, sample_step=N_REF // N_OUT)
        d_t = np.max(np.abs(TT - prev[0]))
        d_c = np.max(np.abs(CC - prev[1]))
        print(f"  nsub={nsub:3d} vs previous: dT = {d_t:.3e} degC, dC = {d_c:.3e} kg/kg")
        prev = (TT, CC)


def main() -> None:
    parser = argparse.ArgumentParser(description="Solve Problem 1 (preheating stage).")
    parser.add_argument(
        "--verify",
        action="store_true",
        help="run the grid/time convergence study (slower)",
    )
    args = parser.parse_args()

    OUT_DIR.mkdir(parents=True, exist_ok=True)

    # production run: fine mesh, sampled onto the required 0.1 cm grid
    r, T_hist, C_hist = simulate(N=N_REF, nsub=NSUB, sample_step=N_REF // N_OUT)
    save_result(r, T_hist, C_hist)

    tab_t, tab_c = extract_tables(r, T_hist, C_hist)
    write_tables(tab_t, tab_c)

    print("Table 1  Temperature [degC]")
    print("t/s   " + "".join(f"{d:>10g}" for d in TABLE_DIST_CM))
    for i, t in enumerate(TABLE_TIMES):
        print(f"{t:<5d}" + "".join(f"{v:>10.4f}" for v in tab_t[i]))

    print("\nTable 2  Moisture concentration [kg/kg]")
    print("t/s   " + "".join(f"{d:>10g}" for d in TABLE_DIST_CM))
    for i, t in enumerate(TABLE_TIMES):
        print(f"{t:<5d}" + "".join(f"{v:>10.4f}" for v in tab_c[i]))

    if args.verify:
        run_verification(T_hist, C_hist)


if __name__ == "__main__":
    main()
