"""Shared finite-volume utilities for the CUMCM 2026 A drying problems.

The routines here implement the node-centred finite-volume framework first
written for ``src/problem-1/solve_problem1.py`` and reused by the coupled
Problem 2 solver.  The geometry is a 1D axisymmetric cylinder of radius ``R``;
all control volumes and face areas are scaled by ``1/(pi*L)`` so that the
length ``L`` never enters the algebra.
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import openpyxl
from scipy.linalg import solve_banded


def harmonic_mean(a, b, name: str = "quantity") -> np.ndarray:
    """Element-wise harmonic mean ``2ab/(a+b)`` of two face-neighbour arrays.

    Parameters
    ----------
    a, b : array_like
        Non-negative values on the two sides of a control-volume face.
    name : str
        Label used in the error message.

    Returns
    -------
    numpy.ndarray
        Harmonic mean with the same shape as the broadcast inputs.

    Raises
    ------
    FloatingPointError
        If two strictly positive inputs produce a zero harmonic mean, i.e. the
        diffusivity has underflowed to zero and the face conductance is not
        representable.
    """
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    denom = a + b
    out = np.zeros_like(denom)
    nz = denom > 0.0
    out[nz] = 2.0 * a[nz] * b[nz] / denom[nz]
    if np.any((out == 0.0) & (a > 0.0) & (b > 0.0)):
        raise FloatingPointError(
            f"{name}: harmonic mean underflowed to zero (diffusivity too small)"
        )
    return out


def build_geometry(N: int, R: float):
    """Build the node-centred finite-volume geometry.

    Parameters
    ----------
    N : int
        Number of control volumes; nodes are ``r_i = i*R/N`` for ``i=0..N``.
    R : float
        Cylinder radius [m].

    Returns
    -------
    dr : float
        Uniform node spacing [m].
    r : numpy.ndarray
        Node radii, shape ``(N+1,)``.
    vol : numpy.ndarray
        Control-volume areas (``V/(pi*L)``), shape ``(N+1,)``.
    ge : numpy.ndarray
        East face areas (``A/(pi*L)``) for nodes ``0..N-1``, shape ``(N,)``.
    gw : numpy.ndarray
        West face areas for nodes ``1..N`` (index 0 unused), shape ``(N+1,)``.
    a_r : float
        Outer surface area (``A_R/(pi*L)``).
    """
    dr = R / N
    r = np.arange(N + 1) * dr
    vol = np.zeros(N + 1)
    vol[0] = (dr / 2.0) ** 2
    vol[1:N] = (r[1:N] + dr / 2.0) ** 2 - (r[1:N] - dr / 2.0) ** 2
    vol[N] = R**2 - (R - dr / 2.0) ** 2
    ge = 2.0 * (r[:N] + dr / 2.0)
    gw = np.zeros(N + 1)
    gw[1:] = 2.0 * (r[1:] - dr / 2.0)
    a_r = 2.0 * R
    return dr, r, vol, ge, gw, a_r


def build_tridiag(vol, ge, gw, a_r, N, dr, dt, s, h, ge_cond, gw_cond):
    """Assemble lower/diagonal/upper of the implicit finite-volume system.

    Parameters
    ----------
    vol, ge, gw, a_r : numpy.ndarray
        Geometry arrays from :func:`build_geometry`.
    N : int
        Number of control volumes.
    dr : float
        Node spacing [m].
    dt : float
        Time step [s].
    s : float or numpy.ndarray
        Volumetric capacity (``rho*cp`` for heat, ``1`` for moisture).
    h : float
        Convective transfer coefficient at the surface.
    ge_cond, gw_cond : numpy.ndarray
        Face conductivities (``k`` or ``D``); ``gw_cond[0]`` is unused.

    Returns
    -------
    lower, diag, upper : numpy.ndarray
        Tridiagonal coefficients.
    """
    lower = np.zeros(N + 1)
    diag = np.asarray(s, dtype=float) * vol / dt
    upper = np.zeros(N + 1)
    diag[:N] += ge * ge_cond / dr
    diag[1:] += gw[1:] * gw_cond[1:] / dr
    diag[N] += a_r * h
    lower[1:] = -gw[1:] * gw_cond[1:] / dr
    upper[:N] = -ge * ge_cond / dr
    return lower, diag, upper


def to_banded(lower, diag, upper) -> np.ndarray:
    """Pack a tridiagonal system into scipy banded storage (1 upper, 1 lower)."""
    n = diag.size
    ab = np.zeros((3, n))
    ab[0, 1:] = upper[:-1]
    ab[1, :] = diag
    ab[2, :-1] = lower[1:]
    return ab


def solve_tridiag(lower, diag, upper, rhs) -> np.ndarray:
    """Solve a tridiagonal system via scipy's banded solver."""
    return solve_banded((1, 1), to_banded(lower, diag, upper), rhs)


def build_rhs(vol, u_old, s, dt, a_r, h, u_inf) -> np.ndarray:
    """Right-hand side of the implicit finite-volume system."""
    rhs = np.asarray(s, dtype=float) * vol / dt * u_old
    rhs[-1] += a_r * h * u_inf
    return rhs


def load_ambient(path, t_required_end: float, t_required_start: float = 0.0):
    """Load and audit the ambient data for ``[t_required_start, t_required_end]``.

    Parameters
    ----------
    path : str or pathlib.Path
        Workbook with columns ``time``, ``temperature``, ``moisture``.
    t_required_end : float
        Latest time [s] that must be covered by the record.
    t_required_start : float
        Earliest time [s] that must be covered by the record.

    Returns
    -------
    t, t_inf, c_inf : numpy.ndarray
        Strictly increasing time stamps [s] and the ambient temperature [degC]
        and moisture concentration [kg/kg].

    Raises
    ------
    ValueError
        If a column is missing, a value is absent/non-numeric, the times are
        not strictly increasing, or the record does not cover the required
        interval.
    """
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    ws = wb.active
    if ws.max_column < 3:
        raise ValueError(f"{path} must contain time, temperature and moisture columns")

    times: list[float] = []
    temps: list[float] = []
    moist: list[float] = []
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is None and row[1] is None and row[2] is None:
            continue
        if row[0] is None or row[1] is None or row[2] is None:
            raise ValueError(f"{path} has a missing value at row {idx}")
        try:
            times.append(float(row[0]))
            temps.append(float(row[1]))
            moist.append(float(row[2]))
        except (TypeError, ValueError) as exc:
            raise ValueError(f"{path} has a non-numeric value at row {idx}") from exc

    if not times:
        raise ValueError(f"{path} contains no data rows")
    t = np.asarray(times, dtype=float)
    t_inf = np.asarray(temps, dtype=float)
    c_inf = np.asarray(moist, dtype=float)
    if np.any(np.diff(t) <= 0.0):
        raise ValueError(f"{path} times must be strictly increasing (no duplicates)")
    if t[0] > t_required_start + 1e-9:
        raise ValueError(
            f"{path} starts at t={t[0]:g} s but t={t_required_start:g} s is required"
        )
    if t[-1] < t_required_end - 1e-9:
        raise ValueError(
            f"{path} ends at t={t[-1]:g} s but t={t_required_end:g} s is required"
        )
    return t, t_inf, c_inf


def write_result_xlsx(path, r, T_hist, C_hist, header: str = "时间\\到药材中心的距离"):
    """Write a Problem-1/2 style workbook with ``温度`` and ``水分浓度`` sheets.

    Parameters
    ----------
    path : str or pathlib.Path
        Destination workbook.
    r : numpy.ndarray
        Sampled node radii [m]; the header row stores ``100*r`` in cm.
    T_hist, C_hist : numpy.ndarray
        Arrays of shape ``(n+1, len(r))`` holding the fields at t=0..n seconds;
        row 0 is dropped so that the workbook holds t=1..n.
    header : str
        Exact A1 text.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    dist_cm = np.round(r * 100.0, 10)

    wb = openpyxl.Workbook()
    for sheet_name, data in (("温度", T_hist), ("水分浓度", C_hist)):
        ws = wb.create_sheet(sheet_name)
        ws.append([header] + [float(d) for d in dist_cm])
        for i in range(1, data.shape[0]):
            ws.append([float(i)] + [round(float(v), 4) for v in data[i]])
    del wb["Sheet"]
    wb.save(path)
